import openpyxl, datetime as dt, sys
from dateutil.relativedelta import relativedelta
fn = sys.argv[1] if len(sys.argv)>1 else 'calc.xlsx'
wb=openpyxl.load_workbook(fn,data_only=True)
def pmt(i,n,pv,fv,t):
    if i==0: return -(pv+fv)/n
    return -(pv*(1+i)**n+fv)*i/((1+i*t)*((1+i)**n-1))
def xirr(cfs):
    d0=cfs[0][0]
    def f(r): return sum(v/(1+r)**((d-d0).days/365) for d,v in cfs)
    lo,hi=-0.99,10
    for _ in range(300):
        mid=(lo+hi)/2
        if f(lo)*f(mid)<=0: hi=mid
        else: lo=mid
    return mid
cot=wb['Cotizador']
for k in range(1,5):
    ws=wb['Corrida %d'%k]; C=lambda r: ws['C%d'%r].value; I=lambda r: ws['I%d'%r].value
    V=758450; eng=0.05*V; M=V-eng; VR=0.05*V; i=0.17/12; tau=1; n=[24,36,48,60][k-1]; m=1
    f=0.205/12; ff=dt.datetime(2026,9,30); fp=dt.datetime(2026,10,1)
    R=pmt(i,n,-M,VR,tau); Rs=pmt(i,m,-VR,0,tau)
    com=0.008*V; dep=R*1.16; days=(fp-ff).days; rp=R/30*days
    DF=(1+0.205/360)**(-days)
    flows=[(t, R) for t in range(0,n)] + [(n+j, Rs) for j in range(m)]
    N=len(flows)
    net0=-V+eng+com+rp+dep
    cm=net0+sum((v-(dep if idx==N-1 else 0))*DF*(1+f)**(-t) for idx,(t,v) in enumerate(flows))
    cms=net0+sum(((v if t<n else 0)-(dep if idx==N-1 else 0))*DF*(1+f)**(-t) for idx,(t,v) in enumerate(flows))
    cfs=[(ff,net0)]+[(fp+relativedelta(months=t),v-(dep if idx==N-1 else 0)) for idx,(t,v) in enumerate(flows)]
    tir=xirr(cfs)
    cat=xirr([(ff,M-com-rp)]+[(fp+relativedelta(months=t),-v) for t,v in flows])
    # tasa minima: buscar por biseccion la tasa que da CM=8%
    def cm_at(ia):
        ii=ia/12; RR=pmt(ii,n,-M,VR,tau); RRs=pmt(ii,m,-VR,0,tau); dd=RR*1.16; rpp=RR/30*days
        n0=-V+eng+com+rpp+dd
        fl=[(t,RR) for t in range(n)]+[(n+j,RRs) for j in range(m)]
        return (n0+sum((v-(dd if idx==len(fl)-1 else 0))*DF*(1+f)**(-t) for idx,(t,v) in enumerate(fl)))/M
    lo,hi=0.0,1.0
    for _ in range(100):
        mid=(lo+hi)/2
        if cm_at(mid)<0.08: lo=mid
        else: hi=mid
    print(f"Esc{k} n={n}: R xl={C(24):.4f} py={R:.4f} | Rs {C(25):.2f}/{Rs:.2f} | CM {I(4):.2f}/{cm:.2f} | CM% {I(5):.4%} | CMsr {I(6):.2f}/{cms:.2f} | TIR {I(8):.4%}/{tir:.4%} | CAT {I(10):.4%}/{cat:.4%} | Saldo {I(16):.6f} | TasaMin xl={I(15):.5%} biseccion={mid:.5%} | PI {I(20):.2f}")
print('Alertas', [cot.cell(row=r,column=c).value for r in range(1,120) for c in range(3,7) if cot.cell(row=r,column=2).value=='Alertas por escenario'])
