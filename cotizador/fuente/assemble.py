"""Ensambla el .xlsm final: valores en caché + proyecto VBA + pestaña de cinta."""
import os, re, zipfile, datetime as dt, shutil, subprocess, sys, html
import openpyxl
from vbabuild import build_vba_project

SRC = 'Cotizador_Arrendamiento.xlsx'
CALC = 'calc.xlsx'
OUT = sys.argv[1] if len(sys.argv) > 1 else 'Cotizador_Arrendamiento_SOFOPLUS.xlsm'
RECALC = os.environ.get('RECALC_SCRIPT', 'recalc.py')  # script de recálculo con LibreOffice

shutil.copy(SRC, CALC)
res = subprocess.run(['python3', RECALC, CALC, '180'], capture_output=True, text=True)
print(res.stdout.strip())
assert '"total_errors": 0' in res.stdout

vals = openpyxl.load_workbook(CALC, data_only=True)
src_wb = openpyxl.load_workbook(SRC)
EPOCH = dt.datetime(1899, 12, 30)


def enc(v):
    if isinstance(v, bool):
        return 'b', '1' if v else '0'
    if isinstance(v, (int, float)):
        return None, repr(float(v)) if isinstance(v, float) else str(v)
    if isinstance(v, dt.datetime):
        d = v - EPOCH
        return None, repr(d.days + d.seconds / 86400)
    if isinstance(v, dt.date):
        return None, str((dt.datetime(v.year, v.month, v.day) - EPOCH).days)
    if v is None:
        return 'str', ''
    s = str(v)
    if s.startswith('#'):
        return 'e', s
    return 'str', html.escape(s, quote=False)


zin = zipfile.ZipFile(SRC)
files = {n: zin.read(n) for n in zin.namelist()}
wbxml = files['xl/workbook.xml'].decode('utf-8')
rels = files['xl/_rels/workbook.xml.rels'].decode('utf-8')
sheets = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml)
relmap = dict((i, t) for t, i in re.findall(r'Target="([^"]+)" Id="([^"]+)"', rels))
n_inj = 0
for sname, rid in sheets:
    sname = html.unescape(sname)
    path = relmap[rid].lstrip('/')
    xml = files[path].decode('utf-8')
    ws = vals[sname]

    def repl(m):
        global n_inj
        ref, attrs, f = m.group(1), m.group(2), m.group(3)
        t, v = enc(ws[ref].value)
        attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
        n_inj += 1
        tt = ' t="%s"' % t if t else ''
        return '<c r="%s"%s%s><f>%s</f><v>%s</v></c>' % (ref, attrs, tt, f, v)
    xml = re.sub(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v ?/></c>', repl, xml)
    files[path] = xml.encode('utf-8')
print('valores inyectados:', n_inj)

# ---- libro con macros
wbxml = wbxml.replace('<workbookPr />', '<workbookPr codeName="ThisWorkbook" />')
assert 'codeName="ThisWorkbook"' in wbxml
files['xl/workbook.xml'] = wbxml.encode('utf-8')
ct = files['[Content_Types].xml'].decode('utf-8')
ct = ct.replace('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
ct = ct.replace('<Default Extension="xml"', '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject" /><Default Extension="xml"', 1)
assert 'macroEnabled' in ct and 'vbaProject' in ct
files['[Content_Types].xml'] = ct.encode('utf-8')
rels = rels.replace('</Relationships>', '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="/xl/vbaProject.bin" Id="rIdVBA1" /></Relationships>')
files['xl/_rels/workbook.xml.rels'] = rels.encode('utf-8')

# ---- módulos VBA
mods = [dict(name='ThisWorkbook', document=True, base='0{00020819-0000-0000-C000-000000000046}',
             code=open('../macros/ThisWorkbook.cls', encoding='utf-8').read())]
for ws in src_wb.worksheets:
    mods.append(dict(name=ws.sheet_properties.codeName, document=True,
                     base='0{00020820-0000-0000-C000-000000000046}', code=''))
mods.append(dict(name='modCotizador', document=False, code=open('../macros/modCotizador.bas', encoding='utf-8').read()))
mods.append(dict(name='modRibbon', document=False, code=open('../macros/modRibbon.bas', encoding='utf-8').read()))
files['xl/vbaProject.bin'] = build_vba_project('orig_vbaProject.bin', mods, 'vbaProject.bin')

# ---- cinta personalizada
def btn(id_, label, action, img=None, size='large', tip=''):
    a = ' imageMso="%s"' % img if img else ''
    s = ' size="%s"' % size if size else ''
    return '<button id="%s" label="%s"%s%s onAction="%s" screentip="%s"/>' % (id_, html.escape(label), s, a, action, html.escape(tip or label))

ui = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
      '<customUI xmlns="http://schemas.microsoft.com/office/2006/01/customui"><ribbon><tabs>'
      '<tab id="tabCotizador" label="Cotizador Arrendamiento" insertAfterMso="TabHome">'
      '<group id="gCot" label="Cotización">'
      + btn('bNueva', 'Nueva cotización', 'rbNueva', 'FileNew', tip='Limpia la captura y asigna folio (Ctrl+Shift+N)')
      + btn('bGuardar', 'Guardar en historial', 'rbGuardar', 'FileSave', tip='Guarda la cotización en Historial (Ctrl+Shift+G)')
      + btn('bCargar', 'Cargar del historial', 'rbCargar', 'FileOpen')
      + '</group><group id="gEnt" label="Entregar al cliente">'
      + btn('bCarta', 'Carta en PDF', 'rbCartaPDF', 'FileSaveAsPdfOrXps', tip='Genera la carta en PDF (Ctrl+Shift+P)')
      + btn('bTabla', 'Tabla de pagos PDF', 'rbTablaPDF', 'TableInsert')
      + btn('bCorreo', 'Enviar por correo', 'rbCorreo', 'FileSendAsAttachment')
      + '</group><group id="gFin" label="Ajustes financieros">'
      + btn('bTasaMin', 'Aplicar tasa mínima', 'rbTasaMin', 'AutoSum', tip='Tasa que cumple el cash margin objetivo')
      + btn('bTasaRenta', 'Tasa para renta deseada', 'rbTasaRenta', 'PercentStyle')
      + btn('bCopiar', 'Copiar escenario 1', 'rbCopiar', 'Copy')
      + '</group><group id="gIr" label="Ir a">'
      + btn('bIrCot', 'Cotizador', 'rbIrCotizador', None, 'normal')
      + btn('bIrCarta', 'Carta', 'rbIrCarta', None, 'normal')
      + btn('bIrTabla', 'Tabla de pagos', 'rbIrTabla', None, 'normal')
      + btn('bIrHist', 'Historial', 'rbIrHistorial', None, 'normal')
      + btn('bIrSen', 'Sensibilidad', 'rbIrSensibilidad', None, 'normal')
      + btn('bIrCorr', 'Corridas', 'rbIrCorrida', None, 'normal')
      + btn('bIrCfg', 'Configuración', 'rbIrConfig', None, 'normal')
      + btn('bIrIni', 'Inicio', 'rbIrInicio', None, 'normal')
      + '</group><group id="gAdm" label="Administración">'
      + btn('bProt', 'Proteger hojas', 'rbProteger', None, 'normal')
      + btn('bDesp', 'Desproteger hojas', 'rbDesproteger', None, 'normal')
      + '</group></tab></tabs></ribbon></customUI>')
files['customUI/customUI.xml'] = ui.encode('utf-8')
rr = files['_rels/.rels'].decode('utf-8')
rr = rr.replace('</Relationships>', '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/ui/extensibility" Target="customUI/customUI.xml" Id="rIdUI1" /></Relationships>')
files['_rels/.rels'] = rr.encode('utf-8')

order = ['[Content_Types].xml', '_rels/.rels'] + [n for n in files if n not in ('[Content_Types].xml', '_rels/.rels')]
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for n in order:
        z.writestr(n, files[n])
print('escrito', OUT)
