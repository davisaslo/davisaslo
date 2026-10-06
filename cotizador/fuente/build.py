import datetime as dt
import copy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
import openpyxl.worksheet.page as page

OUT = 'Cotizador_Arrendamiento.xlsx'
import os
# Promotores: si existe la variable PROMOTORES (ruta al Excel "Relación de colaboradores"), se cargan
# los datos reales; si no, se usan datos de ejemplo (versión pública sin datos personales).
PROMOTORES = []
if os.environ.get('PROMOTORES'):
    import openpyxl as _ox
    _ws = _ox.load_workbook(os.environ['PROMOTORES'], data_only=True).worksheets[0]
    for _row in _ws.iter_rows(min_row=3, values_only=True):
        if _row[1]:
            _nom = ' '.join(str(x).strip() for x in _row[1:4] if x)
            _p = ' '.join(w if w in ('de', 'del', 'la', 'y') else w.capitalize() for w in str(_row[6] or '').lower().split())
            PROMOTORES.append((str(_row[0]), _nom.upper(), _p, _row[5] or '', _row[4] or ''))
    PROM_DEFAULT = next((p[1] for p in PROMOTORES if 'DORYAN' in p[1]), PROMOTORES[0][1])
else:
    PROMOTORES = [('001', 'NOMBRE DEL PROMOTOR', 'Promotor', 'CDMX', 'promotor@empresa.com.mx'),
                  ('002', 'NOMBRE DEL GERENTE', 'Gerente De Promoción', 'CDMX', 'gerente@empresa.com.mx')]
    PROM_DEFAULT = 'NOMBRE DEL PROMOTOR'
LOGO = 'logo.jpeg'

# ------------------------------------------------------------------ estilo
MORADO = '4B2A7B'
MORADO_CL = 'E9E2F3'
GRIS = 'F2F2F2'
AMARILLO = 'FFF8D6'
FONT = 'Arial'

f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_input = Font(name=FONT, size=10, color='0000FF')
f_link = Font(name=FONT, size=10, color='008000')
f_hdr = Font(name=FONT, size=11, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=18, bold=True, color=MORADO)
f_sub = Font(name=FONT, size=10, italic=True, color='666666')
f_note = Font(name=FONT, size=9, italic=True, color='666666')
fill_hdr = PatternFill('solid', fgColor=MORADO)
fill_sub = PatternFill('solid', fgColor=MORADO_CL)
fill_in = PatternFill('solid', fgColor=AMARILLO)
fill_gris = PatternFill('solid', fgColor=GRIS)
thin = Side(style='thin', color='BFBFBF')
med = Side(style='medium', color=MORADO)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
topline = Border(top=med)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
right = Alignment(horizontal='right', vertical='center')
left_c = Alignment(horizontal='left', vertical='center', wrap_text=True)

MON = '$#,##0.00;[Red]-$#,##0.00;"-"'
MON0 = '$#,##0;[Red]-$#,##0;"-"'
PCT = '0.00%;[Red]-0.00%;"-"'
PCT4 = '0.0000%'
FECHA = 'dd/mm/yyyy'
NUM = '#,##0.00'

wb = Workbook()
wb.remove(wb.active)
names = {}


def name(nm, ref):
    names[nm] = ref
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def absref(sheet, rng):
    parts = rng.split(':')
    def a(c):
        col = ''.join(ch for ch in c if ch.isalpha())
        row = ''.join(ch for ch in c if ch.isdigit())
        return '$%s$%s' % (col, row)
    q = "'%s'" % sheet if ' ' in sheet else sheet
    return q + '!' + ':'.join(a(p) for p in parts)


def style_range(ws, rng, font=None, fill=None, border=None, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = border
            if align: c.alignment = align
            if fmt: c.number_format = fmt


def header_bar(ws, row, c1, c2, text):
    ws.cell(row=row, column=c1, value=text)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_hdr
        cell.font = f_hdr
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[row].height = 20


def inp(cell, value, fmt=None, note=None):
    cell.value = value
    cell.font = f_input
    cell.fill = fill_in
    cell.border = box
    cell.protection = Protection(locked=False)
    if fmt:
        cell.number_format = fmt
    if note:
        cell.comment = Comment(note, 'Cotizador')


def calc(cell, formula, fmt=None, bold=False):
    cell.value = formula
    cell.font = f_bold if bold else f_base
    cell.border = box
    if fmt:
        cell.number_format = fmt


def logo(ws, anchor, h=60):
    img = XLImage(LOGO)
    ratio = img.width / img.height
    img.height = h
    img.width = int(h * ratio)
    ws.add_image(img, anchor)


def base_sheet(title, code, tab=None):
    ws = wb.create_sheet(title)
    ws.sheet_properties.codeName = code
    if tab:
        ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    return ws


def dv_list(ws, formula, cells, prompt=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.error = 'Seleccione un valor de la lista.'
    dv.errorTitle = 'Valor no válido'
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cells)


def dv_num(ws, cells, lo, hi, kind='decimal', msg=''):
    dv = DataValidation(type=kind, operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.error = msg or ('Capture un valor entre %s y %s.' % (lo, hi))
    dv.errorTitle = 'Valor fuera de rango'
    ws.add_data_validation(dv)
    dv.add(cells)


# ==================================================================== HOJAS
ws_ini = base_sheet('Inicio', 'Hoja1', MORADO)
ws_cot = base_sheet('Cotizador', 'Hoja2', 'FFC000')
ws_pro = base_sheet('Propuesta', 'Hoja3', '70AD47')
ws_car = base_sheet('Carta Cotizacion', 'Hoja4', '70AD47')
ws_pin = base_sheet('Pago Inicial', 'Hoja5', '70AD47')
ws_tab = base_sheet('Tabla de Pagos', 'Hoja6', '70AD47')
ws_ven = base_sheet('Promesa de Venta', 'Hoja7', '70AD47')
ws_com = base_sheet('Resumen Comite', 'Hoja8', 'C00000')
ws_sen = base_sheet('Sensibilidad', 'Hoja9', '5B9BD5')
ws_his = base_sheet('Historial', 'Hoja10', '5B9BD5')
ws_cr = [base_sheet('Corrida %d' % k, 'Hoja%d' % (10 + k), '808080') for k in range(1, 5)]
ws_cfg = base_sheet('Configuracion', 'Hoja15', '808080')

# ==================================================================== CONFIGURACION
ws = ws_cfg
for col, w in zip('ABCDEFGHIJKLMNOP', [2, 46, 22, 14, 14, 14, 3, 30, 24, 20, 14, 10, 14, 16, 18, 3]):
    ws.column_dimensions[col].width = w
ws['B1'] = 'CONFIGURACIÓN DEL COTIZADOR'
ws['B1'].font = f_title
ws['B2'] = 'Solo el administrador debe modificar esta hoja. Las celdas amarillas son editables.'
ws['B2'].font = f_sub
header_bar(ws, 3, 2, 6, 'DATOS DE LA EMPRESA Y PARÁMETROS')
cfg = [
    ('Razón social del arrendador', 'SOFOPLUS, S.A.P.I. DE C.V., E.R.', 'par_Arrendador', None),
    ('Nombre comercial', 'SOFOPLUS', 'par_Comercial', None),
    ('Domicilio (línea 1)', 'Paseo de los Tamarindos No. 90 Piso 24 Torre 1', 'par_Dir1', None),
    ('Domicilio (línea 2)', 'Col. Bosques de las Lomas', 'par_Dir2', None),
    ('Domicilio (línea 3)', 'CDMX, C.P. 05120', 'par_Dir3', None),
    ('Ciudad para la fecha de la carta', 'Ciudad de México', 'par_Ciudad', None),
    ('Teléfono / correo / web (pie de carta)', '', 'par_Contacto', None),
    ('Tasa de IVA', 0.16, 'par_IVA', PCT),
    ('Tasa de fondeo anual (valor por defecto)', 0.205, 'par_Fondeo', PCT),
    ('Cash margin mínimo objetivo (% del monto financiado)', 0.08, 'par_CMmin', PCT),
    ('Vigencia de la propuesta (días hábiles)', 5, 'par_Vigencia', '0'),
    ('Base de días para costo de fondeo diario', 360, 'par_BaseDias', '0'),
    ('IVA en arrendamiento financiero', 'Sobre renta', 'par_BaseIVAFin', None),
    ('Prefijo del folio', 'COT-', 'par_Prefijo', None),
    ('Último consecutivo utilizado (lo actualiza la macro)', 0, 'par_Consecutivo', '0'),
    ('Carpeta para PDF (vacío = carpeta del archivo)', '', 'par_CarpetaPDF', None),
    ('Correo con copia (CC) al enviar cotizaciones', '', 'par_CorreoCC', None),
    ('TIIE 28 días vigente (valor por defecto)', 0.086, 'par_TIIE', PCT4),
    ('Gastos de investigación / validación de cuenta (sin IVA)', 0, 'par_GastosInv', MON),
    ('Comisión del promotor por defecto (% del monto financiado)', 0, 'par_ComProm', PCT),
    ('Límite deducible renta automóvil ($ diarios, LISR art. 28-XIII)', 200, 'par_LimAuto', MON),
    ('Límite deducible renta auto eléctrico / híbrido ($ diarios)', 285, 'par_LimAutoEV', MON),
    ('Banco para el pago inicial', '', 'par_Banco', None),
    ('Beneficiario de la cuenta', 'SOFOPLUS, S.A.P.I. DE C.V., E.R.', 'par_Beneficiario', None),
    ('Número de cuenta', '', 'par_Cuenta', None),
    ('CLABE interbancaria', '', 'par_CLABE', None),
    ('Convenio / referencia bancaria', '', 'par_Convenio', None),
]
r = 4
for label, val, nm, fmt in cfg:
    ws.cell(row=r, column=2, value=label).font = f_base
    inp(ws.cell(row=r, column=3), val, fmt)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    name(nm, absref(ws.title, 'C%d' % r))
    r += 1
ws['C4'].comment = Comment('Fuente: carta de cotización del archivo original.', 'Cotizador')
ws['C12'].comment = Comment('Dato del archivo original (CM 24m!B17 = 20.5%). Actualícelo con su costo de fondeo vigente.', 'Cotizador')
ws['C13'].comment = Comment('El archivo original usaba 8% como meta por defecto en la macro AjustarCash. El rate card (derecha) define el mínimo por plazo.', 'Cotizador')
ws['C21'].comment = Comment('Referencia: el cotizador de ABC Leasing usaba 8.60%. Actualice con la TIIE publicada por Banxico.', 'Cotizador')
ws['C24'].comment = Comment('Ley del ISR art. 28 fr. XIII: renta de automóviles deducible hasta $200 diarios ($285 eléctricos/híbridos). Verifique con su área fiscal.', 'Cotizador')
dv_list(ws, '=lst_BaseIVA', 'C16')

r += 1
DEF_ROW = r + 1
header_bar(ws, r, 2, 6, 'VALORES POR DEFECTO PARA UNA NUEVA COTIZACIÓN')
for k in range(4):
    c = ws.cell(row=r, column=3 + k, value='Escenario %d' % (k + 1))
    c.alignment = center
r += 1
defaults = [
    ('Plazo básico (meses)', [24, 36, 48, 60], 'def_Plazo', '0'),
    ('Plazo sucesivo (meses, 0 = opción de compra)', [1, 1, 1, 1], 'def_Sucesivo', '0'),
    ('Tasa anual plazo básico', [0.17] * 4, 'def_Tasa', PCT),
    ('Tasa anual plazo sucesivo', [0.17] * 4, 'def_TasaSuc', PCT),
    ('Margen sobre TIIE (modo TIIE + margen)', [0.084] * 4, 'def_Margen', PCT),
    ('Enganche (% del valor)', [0.05] * 4, 'def_Enganche', PCT),
    ('Valor residual (% del valor)', [0.05] * 4, 'def_Residual', PCT),
    ('Comisión por apertura (% del valor)', [0.008] * 4, 'def_Comision', PCT),
    ('Depósito en garantía (# rentas con IVA)', [1] * 4, 'def_Deposito', '0.00'),
]
for label, vals, nm, fmt in defaults:
    ws.cell(row=r, column=2, value=label).font = f_base
    for k, v in enumerate(vals):
        inp(ws.cell(row=r, column=3 + k), v, fmt)
    name(nm, absref(ws.title, 'C%d:F%d' % (r, r)))
    r += 1
singles = [
    ('Tipo de arrendamiento', 'Arrendamiento Puro', 'def_Tipo', '=lst_Tipo'),
    ('Moneda', 'Moneda Nacional', 'def_Moneda', '=lst_Moneda'),
    ('Modalidad de pago', 'Anticipado', 'def_Modalidad', '=lst_Modalidad'),
]
for label, v, nm, lst in singles:
    ws.cell(row=r, column=2, value=label).font = f_base
    inp(ws.cell(row=r, column=3), v)
    name(nm, absref(ws.title, 'C%d' % r))
    dv_list(ws, lst, 'C%d' % r)
    r += 1

# ---- textos de la carta
r += 1
header_bar(ws, r, 2, 6, 'TEXTOS DE LA CARTA (use {ARRENDADOR} y {VIGENCIA} como comodines)')
r += 1
TEXTOS = [
    ('Gastos de Ratificación:', 'Dependiendo de la ciudad, pagaderos (MXN) por cada Anexo de Arrendamiento.'),
    ('Opciones al Término del Contrato:', 'Al concluir el Plazo Básico, o cualquier extensión del mismo, de cada Anexo de Arrendamiento, y siempre y cuando el Arrendatario se encuentre al corriente en el cumplimiento de todas y cada una de las obligaciones que del Arrendamiento Maestro, así como cualesquiera Anexo de Arrendamiento le imponen, este último tendrá derecho preferente para adquirir la totalidad (pero no menos que la totalidad) del Equipo materia del arrendamiento instrumentado en el Anexo de arrendamiento de que se trate. Si el Arrendatario no ejerce su derecho de preferencia mencionado en el párrafo anterior, el Arrendatario deberá, a su costo, desinstalar, re empaquetar y devolver el equipo (no menos que todo) al Arrendador en su domicilio, o el lugar y fecha que este último le indique por escrito para tal efecto, en el estado y condiciones en que el Arrendatario lo hubiese recibido, sin mayor desgaste que el normalmente derivado de su uso o deterioro normal. En caso de que el Equipo no fuese devuelto en la fecha a que se refiere el párrafo anterior, el plazo del arrendamiento respectivo se prorrogará automáticamente trimestralmente hasta la devolución respectiva, en la inteligencia de que las prórrogas correspondientes se considerarán forzosas para efectos del arrendamiento y durante la vigencia de cada una de ellas el Arrendatario se obliga a pagar al Arrendador una cantidad igual al Pago Periódico correspondiente. La totalidad de impuestos, derechos, contribuciones o gastos de cualquier naturaleza que se ocasionen con motivo de la adquisición del Equipo, serán a cargo del Arrendatario.'),
    ('1.- Compra del Equipo:', 'Anticipamos que el Arrendatario enviará al proveedor del equipo su orden de compra. El Arrendador asumirá el pago de la orden de compra del Arrendatario, lo anterior condicionado a que el Arrendatario rente el equipo del Arrendador.'),
    ('2.- Beneficios Fiscales:', 'El Arrendador deberá ser considerado como el dueño del equipo para los efectos fiscales correspondientes.'),
    ('3.- Arrendamiento Neto:', 'El Arrendatario conviene en que su obligación de pagar la Renta, mediante los Pagos Periódicos, y otras cantidades que sean pagaderas conforme al presente, será absoluta e incondicional, y no estará sujeta a reducciones de ninguna clase, incluyendo sin limitación alguna, reducciones por causa de cualquier demanda pasada, presente o futura que sea consecuencia de este Arrendamiento, cualquier Anexo de Arrendamiento o por cualquier otra razón, o en contra del fabricante, proveedor o constructor del Equipo o en contra de cualquier otra persona física o persona moral. El Arrendatario es responsable por todos los gastos de mantenimiento, seguro e impuestos relativos a la compra, renta, posesión y uso del equipo excepto por el Impuesto al Valor Agregado y gastos de importación correspondientes a la compra e importación del Equipo, siendo estos últimos parte del costo del Equipo. El Arrendatario también es responsable del Impuesto al Valor Agregado de las rentas, excluyendo el impuesto sobre la renta a cargo del Arrendador.'),
    ('4.- Moneda:', 'Todos los pagos por concepto de arrendamiento están expresados en la moneda descrita en el tipo de operación así como cualquier otra cantidad que deba ser pagada al Arrendador de acuerdo a la documentación (incluyendo sin limitación alguna costos y gastos).'),
    ('5.- Mantenimiento y Seguro:', 'Todos los mantenimientos y seguros de los equipos serán responsabilidad del Arrendatario durante todo el tiempo que dure el Plazo Básico, o cualquier extensión del mismo establecido en cualesquiera de los Anexos de arrendamiento. El Arrendatario deberá contratar el seguro por los montos y compañía de seguros aceptable al Arrendador y de acuerdo a disponibilidad de coberturas ofrecidas en México. El Arrendatario deberá enviar al Arrendador copia de las pólizas o certificados de seguro como evidencia de la cobertura.'),
    ('6.- Caducidad:', 'La presente propuesta, sus términos y condiciones expiran {VIGENCIA} días hábiles posteriores a la fecha aquí marcada, si {ARRENDADOR} no recibe aceptación por escrito de la misma.'),
    ('7.- Cambio Material Adverso:', 'La presente propuesta está sujeta a que no suceda ningún Cambio Material Adverso que pueda afectar negativamente los términos y condiciones presentados, entendiéndose por cambio material adverso cualquier alteración en las condiciones macroeconómicas, microeconómicas y sociopolíticas del País, liberando por tanto el Arrendatario a el Arrendador de cualquier responsabilidad a este efecto.'),
    ('8.- Confidencialidad:', 'La presente propuesta solo podrá ser vista y analizada por el personal de la Arrendataria y aquellas personas que formen parte del proceso de autorización de esta propuesta.'),
    ('9.- Riesgo Cambiario:', '{ARRENDADOR} no asume ninguna responsabilidad derivada por cualquier variación al tipo de cambio al momento del pago al proveedor, obligándose el Arrendatario a pagar al Arrendador previo pago al proveedor, cualquier diferencia que pudiera existir con relación al tipo de cambio estimado para el cierre de la operación de arrendamiento.'),
    ('10.- Renta Proporcional:', 'El Arrendatario se obliga a pagar una contraprestación por concepto de renta proporcional, misma que se genere por los días que transcurran entre la fecha de pago al proveedor y la fecha de inicio del plazo del arrendamiento.'),
    ('11.- Autorización:', 'La autorización de la presente propuesta está sujeta a que el departamento de crédito y legal de {ARRENDADOR} la aprueben mediante los procesos internos establecidos y la celebración de la documentación correspondiente, por lo que ninguna de las partes involucradas estará legalmente comprometida con la otra por causa de esta propuesta, ni se generará ningún derecho, responsabilidad u obligación como resultado de la misma.'),
    ('Cierre:', 'Agradecemos de antemano el favorecernos con su decisión, esperando tener noticias de usted a la brevedad posible.'),
]
TXT_ROW0 = r
for label, txt in TEXTOS:
    ws.cell(row=r, column=2, value=label).font = f_bold
    ws.cell(row=r, column=2).alignment = wrap
    inp(ws.cell(row=r, column=3), txt)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ws.cell(row=r, column=3).alignment = wrap
    ws.row_dimensions[r].height = max(15, 12.5 * (len(txt) // 72 + 1))
    r += 1
ws.cell(row=TXT_ROW0 - 1, column=2).comment = Comment('Textos tomados de la hoja "Carta Cotizacion" del archivo original.', 'Cotizador')

# ---- listas
header_bar(ws, 3, 8, 21, 'LISTAS DESPLEGABLES')
listas = [
    ('H', 'Tipo de arrendamiento', ['Arrendamiento Puro', 'Arrendamiento Financiero'], 'lst_Tipo'),
    ('I', 'Moneda', ['Moneda Nacional', 'Dólares Americanos'], 'lst_Moneda'),
    ('J', 'Modalidad', ['Anticipado', 'Vencido'], 'lst_Modalidad'),
    ('K', 'Sí / No', ['Sí', 'No'], 'lst_SiNo'),
    ('L', 'Forma de pago', ['Financiado', 'Contado'], 'lst_Forma'),
    ('M', 'Base IVA financiero', ['Sobre renta', 'Sobre intereses'], 'lst_BaseIVA'),
    ('N', 'Estatus', ['Enviada', 'En seguimiento', 'Aceptada', 'Rechazada', 'Vencida'], 'lst_Estatus'),
    ('O', 'Producto', ['Arrendamiento SOFOPLUS', 'FASTPLUS'], 'lst_Producto'),
    ('P', 'Tipo de activo', ['Equipo / maquinaria', 'Equipo médico', 'Equipo de cómputo / tecnología',
                             'Automóvil', 'Automóvil eléctrico / híbrido', 'Vehículo de carga / utilitario'], 'lst_TipoActivo'),
    ('Q', 'Modo de tasa', ['Tasa fija', 'TIIE + margen'], 'lst_ModoTasa'),
    ('R', 'Anticipo se factura como', ['Enganche', 'Renta extraordinaria'], 'lst_UsoAnticipo'),
    ('S', 'Escenario', [1, 2, 3, 4], 'lst_Esc'),
]
for col, title, vals, nm in listas:
    c = ws['%s4' % col]
    c.value = title
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, v in enumerate(vals):
        cell = ws['%s%d' % (col, 5 + i)]
        cell.value = v
        cell.font = f_base
        cell.border = box
    name(nm, absref(ws.title, '%s5:%s%d' % (col, col, 4 + len(vals))))
for col in 'HIJKLMNOPQRS':
    ws.column_dimensions[col].width = 18
ws.row_dimensions[4].height = 30

# ---- rate card
RC0 = 13
header_bar(ws, RC0, 8, 12, 'RATE CARD – MÍNIMOS POR PLAZO (política de precio)')
for j, h in enumerate(['Plazo desde (meses)', 'Cash margin mínimo (neto)', 'TIR mínima', 'Residual máximo', 'Notas']):
    c = ws.cell(row=RC0 + 1, column=8 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = Alignment(wrap_text=True, horizontal='center')
ws.row_dimensions[RC0 + 1].height = 30
rc_rows = [(1, 0.08, 0.23, 0.40), (24, 0.08, 0.23, 0.38), (36, 0.08, 0.23, 0.35), (48, 0.08, 0.23, 0.30),
           (60, 0.08, 0.23, 0.25), (72, 0.08, 0.23, 0.20)]
for i, (pl, cm, tir, res) in enumerate(rc_rows):
    rr = RC0 + 2 + i
    inp(ws.cell(row=rr, column=8), pl, '0')
    inp(ws.cell(row=rr, column=9), cm, PCT)
    inp(ws.cell(row=rr, column=10), tir, PCT)
    inp(ws.cell(row=rr, column=11), res, PCT)
name('rc_Tabla', absref(ws.title, 'H%d:K%d' % (RC0 + 2, RC0 + 1 + len(rc_rows))))
ws.cell(row=RC0 + 2, column=12, value='Se usa la fila cuyo "plazo desde" sea el mayor que no exceda el plazo cotizado.').font = f_note
ws.cell(row=RC0 + 2, column=12).alignment = wrap
ws.merge_cells(start_row=RC0 + 2, start_column=12, end_row=RC0 + 5, end_column=14)
ws.cell(row=RC0 + 1, column=10).comment = Comment('Valores iniciales sugeridos; ajústelos a la política de SOFOPLUS. Como referencia, el rate card de ABC Leasing pedía TIR mínima de 23%.', 'Cotizador')

# ---- oficinas
OF0 = RC0 + 10
header_bar(ws, OF0, 8, 12, 'OFICINAS POR REGIÓN (encabezado de la propuesta)')
for j, h in enumerate(['Región', 'Domicilio (línea 1)', 'Domicilio (línea 2)', 'Domicilio (línea 3)', 'Teléfono']):
    c = ws.cell(row=OF0 + 1, column=8 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
oficinas = [('CDMX', 'Paseo de los Tamarindos No. 90 Piso 24 Torre 1', 'Col. Bosques de las Lomas', 'CDMX, C.P. 05120', ''),
            ('GUADALAJARA', '', '', '', ''), ('CANCÚN', '', '', '', ''), ('QUERÉTARO', '', '', '', ''),
            ('', '', '', '', ''), ('', '', '', '', '')]
for i, row in enumerate(oficinas):
    for j, v in enumerate(row):
        inp(ws.cell(row=OF0 + 2 + i, column=8 + j), v if v else None)
name('ofi_Tabla', absref(ws.title, 'H%d:L%d' % (OF0 + 2, OF0 + 1 + len(oficinas))))
name('ofi_Region', absref(ws.title, 'H%d:H%d' % (OF0 + 2, OF0 + 1 + len(oficinas))))
ws.cell(row=OF0 + 2 + len(oficinas), column=8, value='Si una región no tiene domicilio se usa el domicilio corporativo (columna izquierda).').font = f_note

# ---- promotores
PR0 = OF0 + 10
NPROM = 60
header_bar(ws, PR0, 8, 12, 'PROMOTORES (lista desplegable del Cotizador)')
for j, h in enumerate(['Clave', 'Nombre completo', 'Puesto', 'Región', 'Correo electrónico']):
    c = ws.cell(row=PR0 + 1, column=8 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
for i in range(NPROM):
    row = PROMOTORES[i] if i < len(PROMOTORES) else (None,) * 5
    for j, v in enumerate(row):
        inp(ws.cell(row=PR0 + 2 + i, column=8 + j), v)
name('prom_Tabla', absref(ws.title, 'H%d:L%d' % (PR0 + 2, PR0 + 1 + NPROM)))
name('lst_Promotores', absref(ws.title, 'I%d:I%d' % (PR0 + 2, PR0 + 1 + NPROM)))
ws.column_dimensions['I'].width = 34
ws.column_dimensions['J'].width = 30
ws.column_dimensions['L'].width = 28
ws.column_dimensions['K'].width = 26

# ---- datos derivados del promotor seleccionado (automático)
SEL0 = r + 2
header_bar(ws, SEL0, 2, 6, 'PROMOTOR SELECCIONADO (automático, no editar)')
sel = [
    ('Nombre (formato carta)', '=PROPER(inp_Promotor)', 'sel_PromNombre'),
    ('Puesto', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),3)&"","")', 'sel_PromPuesto'),
    ('Región', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),4)&"","")', 'sel_PromRegion'),
    ('Correo', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),5)&"","")', 'sel_PromCorreo'),
    ('Domicilio línea 1', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir1,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""),par_Dir1)', 'sel_Dir1'),
    ('Domicilio línea 2', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir2,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),3)&""),par_Dir2)', 'sel_Dir2'),
    ('Domicilio línea 3', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir3,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),4)&""),par_Dir3)', 'sel_Dir3'),
    ('Teléfono', '=IFERROR(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),5)&"","")', 'sel_Tel'),
]
for i, (lab, f, nm) in enumerate(sel):
    rr = SEL0 + 1 + i
    ws.cell(row=rr, column=2, value=lab).font = f_base
    calc(ws.cell(row=rr, column=3), f)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    name(nm, absref(ws.title, 'C%d' % rr))
ws.freeze_panes = 'A4'

# ==================================================================== COTIZADOR
ws = ws_cot
for col, w in zip('ABCDEFGH', [2, 46, 18, 18, 18, 18, 2, 60]):
    ws.column_dimensions[col].width = w
logo(ws, 'B1', 58)
ws['C1'] = 'COTIZADOR DE ARRENDAMIENTO'
ws['C1'].font = f_title
ws.merge_cells('C1:F1')
ws['C2'] = 'Corridas financieras, cash margin, TIR y carta para el cliente'
ws['C2'].font = f_sub
ws.merge_cells('C2:F2')
ws['C3'] = '="Folio: "&inp_Folio&"   |   "&inp_Cliente'
ws['C3'].font = Font(name=FONT, size=11, bold=True, color=MORADO)
ws.merge_cells('C3:F3')
for rr in (1, 2, 3):
    ws.row_dimensions[rr].height = 22
ws['H1'] = 'CÓMO USAR'
ws['H1'].font = f_bold
ws['H2'] = '1) Capture las celdas AMARILLAS (texto azul). 2) Revise resultados y semáforo. 3) Use la pestaña "Cotizador Arrendamiento" de la cinta (o Alt+F8) para guardar, generar PDF o enviar.'
ws['H2'].alignment = wrap
ws['H2'].font = f_note
ws.merge_cells('H2:H4')

header_bar(ws, 5, 2, 8, '1. DATOS DEL CLIENTE')
hoy = dt.datetime(2026, 10, 5)
datos = [
    ('Folio de la cotización', 'COT-0001', 'inp_Folio', None, 'La macro "Nueva cotización" asigna el siguiente folio automáticamente.'),
    ('Fecha de la cotización', hoy, 'inp_Fecha', FECHA, 'Fecha que aparece en la carta.'),
    ('Cliente / Arrendatario (razón social)', 'CLIENTE DE EJEMPLO, S.A. DE C.V.', 'inp_Cliente', None, 'Razón social completa del arrendatario.'),
    ('RFC del cliente', '', 'inp_RFC', None, 'Opcional.'),
    ('Atención a (nombre del contacto)', '', 'inp_Contacto', None, 'Opcional. Aparece en la carta como "At\'n".'),
    ('Correo electrónico del cliente', '', 'inp_Correo', None, 'Se usa para el botón "Enviar por correo".'),
    ('Proveedor del equipo', '', 'inp_Proveedor', None, ''),
    ('Descripción del equipo', 'Auto BYD M9 2026', 'inp_Equipo', None, 'Marca, modelo, año, cantidad.'),
    ('Promotor', PROM_DEFAULT, 'inp_Promotor', None, 'Lista editable en la hoja Configuracion (puesto, región y correo salen de ahí).'),
    ('Producto / línea de negocio', 'Arrendamiento SOFOPLUS', 'inp_Producto', None, 'SOFOPLUS o FASTPLUS.'),
    ('Obligado solidario', 'Por definir', 'inp_Obligado', None, ''),
]
r = 6
for label, val, nm, fmt, note in datos:
    ws.cell(row=r, column=2, value=label).font = f_base
    inp(ws.cell(row=r, column=3), val, fmt)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    name(nm, absref(ws.title, 'C%d' % r))
    ws.cell(row=r, column=8, value=note).font = f_note
    r += 1
dv_list(ws, '=lst_Promotores', names['inp_Promotor'].split('!')[1].replace('$', ''))
dv_list(ws, '=lst_Producto', names['inp_Producto'].split('!')[1].replace('$', ''))
_pr = int(names['inp_Promotor'].split('$')[-1])
ws.cell(row=_pr, column=8, value='=IF(sel_PromPuesto="","Promotor no encontrado en la lista",sel_PromPuesto&" · "&sel_PromRegion&" · "&sel_PromCorreo)').font = f_note

r += 1
header_bar(ws, r, 2, 8, '2. CONDICIONES GENERALES DE LA OPERACIÓN')
r += 1
gen = [
    ('Tipo de arrendamiento', 'Arrendamiento Puro', 'inp_Tipo', None, '=lst_Tipo', 'Puro: IVA sobre la renta completa. Financiero: según Configuracion.'),
    ('Moneda', 'Moneda Nacional', 'inp_Moneda', None, '=lst_Moneda', 'Todos los montos se capturan en la moneda de la operación.'),
    ('Tipo de cambio (informativo)', 1, 'inp_TC', '#,##0.0000', None, 'Solo informativo; capture 1 si es Moneda Nacional.'),
    ('Modalidad de pago de las rentas', 'Anticipado', 'inp_Modalidad', None, '=lst_Modalidad', 'Anticipado: la renta se paga al inicio de cada mes. Vencido: al final. Cambia la fórmula de la renta (PMT tipo 1 ó 0).'),
    ('Tipo de activo', 'Automóvil eléctrico / híbrido', 'inp_TipoActivo', None, '=lst_TipoActivo', 'Para automóviles se estima la deducibilidad de la renta (límite ISR por día).'),
    ('Precio del equipo capturado', 758450, 'inp_Precio', MON, None, 'Precio de factura del proveedor (ejemplo: archivo original).'),
    ('¿El precio capturado incluye IVA?', 'No', 'inp_PrecioIVA', None, '=lst_SiNo', 'Si captura el precio con IVA (como en el cotizador de ABC) elija Sí; se desglosa solo.'),
    ('Fecha de firma / pago al proveedor', dt.datetime(2026, 9, 30), 'inp_FechaFirma', FECHA, None, 'Fecha del desembolso; aquí se cobran enganche, comisión, depósito y renta proporcional.'),
    ('Fecha de la primera renta', dt.datetime(2026, 10, 1), 'inp_FechaPrimera', FECHA, None, 'Si hay días entre la firma y el inicio del plazo se cobra renta proporcional (renta/30 x días).'),
    ('Comisión banco / alianza ($, uso interno)', 0, 'inp_ComBanco', MON, None, 'Costo pagado a un tercero; reduce el cash margin.'),
    ('Modo de tasa', 'Tasa fija', 'inp_ModoTasa', None, '=lst_ModoTasa', 'Tasa fija: use la fila "Tasa anual". TIIE + margen: tasa = TIIE + margen de cada escenario (como ABC).'),
    ('TIIE 28 días', 0.086, 'inp_TIIE', PCT4, None, 'Solo se usa en modo TIIE + margen. Valor inicial desde Configuracion.'),
    ('El anticipo se factura como', 'Enganche', 'inp_UsoAnticipo', None, '=lst_UsoAnticipo', 'Solo cambia el nombre en los documentos; el cálculo es el mismo.'),
    ('Gastos de investigación / validación (sin IVA)', 0, 'inp_GastosInv', MON, None, 'Se cobran en el pago inicial (más IVA).'),
    ('Tasa de fondeo anual (uso interno)', 0.205, 'inp_Fondeo', PCT, None, 'Costo del dinero para calcular cash margin. Valor inicial desde Configuracion.'),
    ('Comisión del promotor (% del monto financiado, interna)', 0, 'inp_ComProm', PCT, None, 'Se resta del cash margin para obtener el margen neto.'),
    ('Cash margin objetivo manual (vacío = rate card)', None, 'inp_CMobj', PCT, None, 'Déjelo vacío para usar el mínimo por plazo de la hoja Configuracion.'),
]
for label, val, nm, fmt, lst, note in gen:
    ws.cell(row=r, column=2, value=label).font = f_base
    inp(ws.cell(row=r, column=3), val, fmt)
    name(nm, absref(ws.title, 'C%d' % r))
    if lst:
        dv_list(ws, lst, 'C%d' % r)
    ws.cell(row=r, column=8, value=note).font = f_note
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical='center')
    r += 1
    if nm == 'inp_PrecioIVA':
        ws.cell(row=r, column=2, value='   Valor del equipo sin IVA').font = f_bold
        calc(ws.cell(row=r, column=3), '=IF(inp_PrecioIVA="Sí",inp_Precio/(1+par_IVA),inp_Precio)', MON, bold=True)
        name('inp_Valor', absref(ws.title, 'C%d' % r))
        r += 1
        ws.cell(row=r, column=2, value='   IVA del equipo').font = f_base
        calc(ws.cell(row=r, column=3), '=inp_Valor*par_IVA', MON)
        r += 1
        ws.cell(row=r, column=2, value='   Valor del equipo con IVA').font = f_base
        calc(ws.cell(row=r, column=3), '=inp_Valor+C%d' % (r - 1), MON)
        ws.cell(row=r, column=4, value='=IF(inp_TC<>1,"Equivalente MXN: "&TEXT(C%d*inp_TC,"$#,##0.00"),"")' % r).font = f_note
        r += 1
dv_num(ws, names['inp_Precio'].split('!')[1].replace('$', ''), 0, 1e12)
dv_num(ws, names['inp_ComProm'].split('!')[1].replace('$', ''), 0, 1)
dv_num(ws, names['inp_Fondeo'].split('!')[1].replace('$', ''), 0, 1)
dv_num(ws, names['inp_CMobj'].split('!')[1].replace('$', ''), -1, 1)

r += 1
header_bar(ws, r, 2, 8, '3. ESCENARIOS A COTIZAR (hasta 4 plazos en la misma carta)')
r += 1
ws.cell(row=r, column=2, value='Escenario elegido para Propuesta, Pago inicial, Tabla, Promesa y Comité').font = f_bold
inp(ws.cell(row=r, column=3), 1, '0')
ws.cell(row=r, column=3).alignment = Alignment(horizontal='center')
name('inp_EscTabla', absref(ws.title, 'C%d' % r))
dv_list(ws, '"1,2,3,4"', 'C%d' % r)
ws.cell(row=r, column=8, value='La Carta Cotizacion compara los 4 escenarios; los demás documentos usan este.').font = f_note
r += 1
ESC_HDR = r
for k in range(4):
    c = ws.cell(row=r, column=3 + k, value='Escenario %d' % (k + 1))
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
    c.border = box
ws.cell(row=r, column=2, value='Concepto').font = f_bold
ws.cell(row=r, column=2).fill = fill_sub
r += 1
esc = [
    ('Incluir en la carta', ['Sí'] * 4, 'esc_Incluir', None, '=lst_SiNo', 'Elija "No" para ocultar el escenario en la carta.'),
    ('Plazo básico (meses)', [24, 36, 48, 60], 'esc_Plazo', '0', ('whole', 1, 84), '1 a 84 meses.'),
    ('Plazo sucesivo (meses; 0 = opción de compra)', [1] * 4, 'esc_Sucesivo', '0', ('whole', 0, 12), 'El valor residual se cobra en rentas del plazo sucesivo. Con 0 se cobra como opción de compra única.'),
    ('Tasa anual plazo básico', [0.17] * 4, 'esc_Tasa', PCT, ('decimal', 0, 2), 'Tasa nominal anual; se divide entre 12 (pagos mensuales).'),
    ('Margen sobre TIIE (modo TIIE + margen)', [0.084] * 4, 'esc_Margen', PCT, ('decimal', -1, 2), 'Solo aplica en modo TIIE + margen.'),
    ('Tasa anual plazo sucesivo (modo tasa fija)', [0.17] * 4, 'esc_TasaSuc', PCT, ('decimal', 0, 2), 'En modo TIIE + margen el sucesivo usa la misma tasa del básico.'),
    ('Enganche / downpayment (% del valor)', [0.05] * 4, 'esc_Enganche', PCT, ('decimal', 0, 0.9), 'Archivo original: 5%.'),
    ('Valor residual (% del valor)', [0.05] * 4, 'esc_Residual', PCT, ('decimal', 0, 0.9), 'Saldo que queda al final del plazo básico (valor futuro en la fórmula de renta).'),
    ('Pago final distinto al residual (% del valor, opcional)', [None] * 4, 'esc_PagoFinal', PCT, ('decimal', 0, 0.9), 'Déjelo VACÍO para que el pago final sea igual al residual (corrida que cuadra en cero).'),
    ('Comisión por apertura (% del valor)', [0.008] * 4, 'esc_Comision', PCT, ('decimal', 0, 0.2), 'Archivo original: 0.8% sobre el valor del equipo.'),
    ('Depósito en garantía (# de rentas con IVA)', [1] * 4, 'esc_Deposito', '0.00', ('decimal', 0, 12), 'Se devuelve/aplica en el último pago.'),
    ('Seguro del equipo (monto total sin IVA)', [0] * 4, 'esc_Seguro', MON, ('decimal', 0, 1e12), 'Monto de la póliza por todo el plazo.'),
    ('Seguro: forma de pago', ['Financiado'] * 4, 'esc_SeguroForma', None, '=lst_Forma', 'Financiado: se suma al monto a financiar. Contado: se cobra en el pago inicial.'),
    ('GPS (monto total sin IVA)', [0] * 4, 'esc_GPS', MON, ('decimal', 0, 1e12), ''),
    ('GPS: forma de pago', ['Financiado'] * 4, 'esc_GPSForma', None, '=lst_Forma', ''),
    ('Mantenimiento / garantía extendida / otros (sin IVA)', [0] * 4, 'esc_Otros', MON, ('decimal', 0, 1e12), 'Como en ABC: mantenimiento preventivo, garantía extendida u otros gastos.'),
    ('Otros: forma de pago', ['Financiado'] * 4, 'esc_OtrosForma', None, '=lst_Forma', ''),
]
ESC = {}
for label, vals, nm, fmt, val, note in esc:
    ws.cell(row=r, column=2, value=label).font = f_base
    for k, v in enumerate(vals):
        inp(ws.cell(row=r, column=3 + k), v, fmt)
        ws.cell(row=r, column=3 + k).alignment = Alignment(horizontal='center')
    name(nm, absref(ws.title, 'C%d:F%d' % (r, r)))
    ESC[nm] = r
    if isinstance(val, str):
        dv_list(ws, val, 'C%d:F%d' % (r, r))
    elif val:
        dv_num(ws, 'C%d:F%d' % (r, r), val[1], val[2], val[0])
    ws.cell(row=r, column=8, value=note).font = f_note
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical='center')
    r += 1

# ---- resultados (se llenan después de definir celdas de las corridas)
r += 1
RES_CLIENTE = r
header_bar(ws, r, 2, 8, '4. RESULTADOS PARA EL CLIENTE')
r += 1

# ==================================================================== CORRIDAS
FIRST = 60          # primera fila de pagos
NROWS = 98          # filas de pagos (84 + 12 + holgura)
LAST = FIRST + NROWS - 1
T0 = FIRST - 1

P = {}  # etiquetas -> celda de parámetro (columna C)


def build_corrida(ws, k):
    for col, w in zip('ABCDEFGHIJKLMNOPQRS', [7, 17, 7, 12, 16, 15, 14, 15, 16, 13, 16, 16, 11, 16, 16, 12, 16, 16, 16]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['B'].width = 44
    ws['A1'] = '="CORRIDA FINANCIERA – ESCENARIO "&$C$3&IF($C$4="Sí",""," (no incluido en la carta)")'
    ws['A1'].font = f_title
    ws['A2'] = 'Hoja de cálculo interna (protegida). Todos los datos vienen de la hoja Cotizador.'
    ws['A2'].font = f_sub
    params = [
        ('Escenario', k, None, 'esc'),
        ('Incluido en la carta', '=INDEX(esc_Incluir,1,$C$3)', None, 'incl'),
        ('Plazo básico (n, meses)', '=INDEX(esc_Plazo,1,$C$3)', '0', 'n'),
        ('Plazo sucesivo (m, meses)', '=INDEX(esc_Sucesivo,1,$C$3)', '0', 'm'),
        ('Tipo de pago (1 = anticipado, 0 = vencido)', '=IF(inp_Modalidad="Anticipado",1,0)', '0', 'tau'),
        ('Tasa anual plazo básico', '=IF(inp_ModoTasa="TIIE + margen",inp_TIIE+INDEX(esc_Margen,1,$C$3),INDEX(esc_Tasa,1,$C$3))', PCT4, 'ia'),
        ('Tasa mensual plazo básico (i)', '=C8/12', PCT4, 'i'),
        ('Tasa anual plazo sucesivo', '=IF(inp_ModoTasa="TIIE + margen",C8,INDEX(esc_TasaSuc,1,$C$3))', PCT4, 'i2a'),
        ('Tasa mensual plazo sucesivo (i2)', '=C10/12', PCT4, 'i2'),
        ('Tasa de fondeo anual', '=inp_Fondeo', PCT4, 'fa'),
        ('Tasa de fondeo mensual (f)', '=C12/12', PCT4, 'f'),
        ('Tasa de IVA', '=par_IVA', PCT, 'iva'),
        ('Valor del equipo sin IVA (V)', '=inp_Valor', MON, 'V'),
        ('Enganche', '=C15*INDEX(esc_Enganche,1,$C$3)', MON, 'eng'),
        ('Seguro financiado', '=IF(INDEX(esc_SeguroForma,1,$C$3)="Financiado",INDEX(esc_Seguro,1,$C$3),0)', MON, 'segf'),
        ('GPS financiado', '=IF(INDEX(esc_GPSForma,1,$C$3)="Financiado",INDEX(esc_GPS,1,$C$3),0)', MON, 'gpsf'),
        ('Monto a financiar (M = V − anticipo + financiados)', '=C15-C16+C17+C18+C42', MON, 'M'),
        ('Valor residual (% del valor)', '=INDEX(esc_Residual,1,$C$3)', PCT, 'resp'),
        ('Valor residual (VR)', '=C15*C20', MON, 'VR'),
        ('Pago final (% del valor)', '=IF(INDEX(esc_PagoFinal,1,$C$3)="",C20,INDEX(esc_PagoFinal,1,$C$3))', PCT, 'pfp'),
        ('Pago final ($)', '=C15*C22', MON, 'PF'),
        ('RENTA BÁSICA  R = PMT(i, n, −M, VR, tipo)', '=PMT(C9,C5,-C19,C21,C7)', MON, 'R'),
        ('Renta plazo sucesivo  Rs = PMT(i2, m, −PF, 0, tipo)', '=IF(C6>0,PMT(C11,C6,-C23,0,C7),0)', MON, 'Rs'),
        ('Opción de compra (si m = 0)', '=IF(C6=0,C23,0)', MON, 'OC'),
        ('Comisión por apertura (%)', '=INDEX(esc_Comision,1,$C$3)', PCT, 'comp'),
        ('Comisión por apertura ($, sobre el valor)', '=C15*C27', MON, 'com'),
        ('Depósito en garantía (# rentas con IVA)', '=INDEX(esc_Deposito,1,$C$3)', '0.00', 'depn'),
        ('Depósito en garantía ($)', '=C29*C24*(1+C14)', MON, 'dep'),
        ('Fecha de firma / desembolso', '=inp_FechaFirma', FECHA, 'ff'),
        ('Fecha de la primera renta', '=inp_FechaPrimera', FECHA, 'fp'),
        ('Inicio del plazo (t = 0)', '=IF(C7=1,C32,EDATE(C32,-1))', FECHA, 'fi'),
        ('Días entre firma e inicio del plazo', '=MAX(0,C33-C31)', '0', 'd'),
        ('Renta proporcional (R/30 × días)', '=C24/30*C34', MON, 'rp'),
        ('Comisión banco / alianza', '=inp_ComBanco', MON, 'cb'),
        ('Factor de fondeo por días  (1+fa/base)^(−días)', '=(1+C12/par_BaseDias)^(-C34)', '0.000000', 'DF'),
        ('Número de pagos en la tabla', '=C5+IF(C6>0,C6,IF(C23>0,1,0))', '0', 'N'),
        ('Desembolso del arrendador en la firma', '=-(C15+C17+C18+C42)', MON, 'out0'),
        ('Cobros en la firma (anticipo+comisión+renta prop.+depósito+gastos−com. banco)', '=C16+C28+C35+C30-C36+C43', MON, 'in0'),
        ('Flujo neto del arrendador en la firma', '=C39+C40', MON, 'net0'),
        ('Mantenimiento / garantía / otros financiados', '=IF(INDEX(esc_OtrosForma,1,$C$3)="Financiado",INDEX(esc_Otros,1,$C$3),0)', MON, 'otrf'),
        ('Gastos de investigación / validación', '=inp_GastosInv', MON, 'gi'),
        ('Comisión del promotor ($)', '=inp_ComProm*C19', MON, 'cp'),
        ('Cash margin objetivo (manual o rate card)', '=IF(inp_CMobj="",VLOOKUP(C5,rc_Tabla,2,TRUE),inp_CMobj)', PCT, 'obj'),
        ('TIR mínima (rate card)', '=VLOOKUP(C5,rc_Tabla,3,TRUE)', PCT, 'tirmin'),
        ('Residual máximo (rate card)', '=VLOOKUP(C5,rc_Tabla,4,TRUE)', PCT, 'resmax'),
        ('Renta: parte del equipo  PMT(i, n, −(V−anticipo), VR, tipo)', '=PMT(C9,C5,-(C15-C16),C21,C7)', MON, 'Req'),
        ('Renta: parte del seguro financiado', '=PMT(C9,C5,-C17,0,C7)', MON, 'Rseg'),
        ('Renta: parte de GPS y otros financiados', '=PMT(C9,C5,-(C18+C42),0,C7)', MON, 'Rotr'),
        ('Deducibilidad estimada de la renta (ISR)', '=IF(inp_TipoActivo="Automóvil",MIN(1,par_LimAuto*30/C24),IF(inp_TipoActivo="Automóvil eléctrico / híbrido",MIN(1,par_LimAutoEV*30/C24),1))', PCT, 'ded'),
        ('Fecha del último pago', '=IFERROR(INDEX(D%d:D%d,C38),"")' % (FIRST, LAST), FECHA, 'fult'),
    ]
    rr = 3
    for label, f, fmt, key in params:
        ws.cell(row=rr, column=2, value=label).font = f_base
        c = ws.cell(row=rr, column=3)
        calc(c, f, fmt, bold=key in ('R', 'M'))
        if key == 'esc':
            c.font = f_input
        P[key] = 'C%d' % rr
        rr += 1
    assert P['net0'] == 'C41', P['net0']
    assert P['otrf'] == 'C42' and P['gi'] == 'C43' and P['cp'] == 'C44' and P['obj'] == 'C45' and P['fult'] == 'C52', P
    assert P['R'] == 'C24' and P['N'] == 'C38' and P['DF'] == 'C37'

    # ---- resultados
    ws['E3'] = 'RESULTADOS DEL ESCENARIO'
    for c in range(5, 10):
        ws.cell(row=3, column=c).fill = fill_hdr
        ws.cell(row=3, column=c).font = f_hdr
    rows = [
        ('Cash margin C/R ($) = VP al fondeo de todos los flujos', '=C41+SUM(N%d:N%d)' % (FIRST, LAST), MON, 'CM'),
        ('Cash margin C/R (% del monto financiado)', '=IFERROR(I4/C19,0)', PCT, 'CMp'),
        ('Cash margin S/R ($) = sin pago final / residual', '=C41+SUM(S%d:S%d)' % (FIRST, LAST), MON, 'CMs'),
        ('Cash margin S/R (%)', '=IFERROR(I6/C19,0)', PCT, 'CMsp'),
        ('TIR anual efectiva del arrendador (XIRR)', '=IFERROR(XIRR(Q%d:Q%d,P%d:P%d),"n/d")' % (T0, LAST, T0, LAST), PCT, 'TIR'),
        ('TIR nominal anual (capitalizable mensual)', '=IFERROR(12*((1+I8)^(1/12)-1),"n/d")', PCT, 'TIRn'),
        ('CAT informativo sin IVA (flujos del cliente)', '=IFERROR(XIRR(R%d:R%d,P%d:P%d),"n/d")' % (T0, LAST, T0, LAST), PCT, 'CAT'),
        ('Σ factores de descuento de rentas básicas', '=SUMIF(B%d:B%d,"Básico",M%d:M%d)' % (FIRST, LAST, FIRST, LAST), '0.000000', 'SB'),
        ('Factor de descuento del último pago', '=IFERROR(INDEX(M%d:M%d,C38),0)' % (FIRST, LAST), '0.000000', 'ML'),
        ('Sensibilidad del CM a la renta (∂CM/∂R)', '=I11+C29*(1+C14)*(1-I12)+C34/30', '0.000000', 'B'),
        ('Renta mínima para el CM objetivo (neto de comisión)', '=IFERROR((C45*C19+C44-(I4-C24*I13))/I13,0)', MON, 'Rmin'),
        ('Tasa anual mínima para el CM objetivo', '=IFERROR(RATE(C5,I14,-C19,C21,C7)*12,"n/d")', PCT, 'TasaMin'),
        ('Saldo final de la tabla (debe ser 0)', '=IFERROR(INDEX(I%d:I%d,C38),0)' % (FIRST, LAST), MON, 'Saldo'),
        ('Pago inicial: base sin IVA (anticipo+comisión+renta prop.+contado+gastos)', '=C16+C28+C35+C43+I29', MON, 'PIb'),
        ('Pago inicial: IVA', '=I17*C14', MON, 'PIi'),
        ('Pago inicial: depósito en garantía', '=C30', MON, 'PId'),
        ('PAGO INICIAL TOTAL (con IVA)', '=I17+I18+I19', MON, 'PI'),
        ('IVA de la primera renta', '=IFERROR(J%d,0)' % FIRST, MON, 'IVA1'),
        ('Total de pagos con IVA (rentas + pago final)', '=SUM(K%d:K%d)' % (FIRST, LAST), MON, 'TotK'),
        ('Spread (tasa cliente − tasa fondeo)', '=C8-C12', PCT, 'Spread'),
        ('Tasa anual aplicada', '=C8', PCT, 'Tasa'),
        ('Cash margin neto de comisión del promotor ($)', '=I4-C44', MON, 'CMn'),
        ('Cash margin neto (% del monto financiado)', '=IFERROR(I25/C19,0)', PCT, 'CMnp'),
        ('Dictamen vs. rate card', '=IF(AND(I26>=C45,N(I8)>=C46),"CUMPLE",IF(I26>=C45,"CM OK / TIR BAJA","NO CUMPLE"))', None, 'Dict'),
        ('Margen mínimo sobre TIIE (si la tasa mínima aplica)', '=IFERROR(I15-inp_TIIE,"n/d")', PCT, 'MargMin'),
        ('Seguro, GPS y otros de contado (sin IVA)', '=IF(INDEX(esc_SeguroForma,1,$C$3)="Contado",INDEX(esc_Seguro,1,$C$3),0)+IF(INDEX(esc_GPSForma,1,$C$3)="Contado",INDEX(esc_GPS,1,$C$3),0)+IF(INDEX(esc_OtrosForma,1,$C$3)="Contado",INDEX(esc_Otros,1,$C$3),0)', MON, 'Cont'),
    ]
    rr = 4
    for label, f, fmt, key in rows:
        ws.cell(row=rr, column=5, value=label).font = f_base
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=8)
        calc(ws.cell(row=rr, column=9), f, fmt, bold=key in ('CM', 'CMp', 'TIR', 'PI'))
        P[key] = 'I%d' % rr
        rr += 1
    assert P['CM'] == 'I4' and P['CMp'] == 'I5' and P['CMs'] == 'I6' and P['TIR'] == 'I8'
    assert P['SB'] == 'I11' and P['ML'] == 'I12' and P['B'] == 'I13' and P['Rmin'] == 'I14' and P['PIb'] == 'I17'
    assert P['CMn'] == 'I25' and P['CMnp'] == 'I26' and P['Cont'] == 'I29', P

    ws['E34'] = ('Método: CM = flujo en firma + Σ flujo_t × (1+fa/base)^(−días) × (1+f)^(−t). '
                 'Como el CM es lineal en la renta, la renta mínima se despeja en forma cerrada y la tasa mínima '
                 'se obtiene con RATE. La tabla usa interés = (saldo − tipo × pago) × i, de modo que el saldo al término '
                 'del plazo básico es exactamente el valor residual.')
    ws['E34'].font = f_note
    ws['E34'].alignment = wrap
    ws.merge_cells('E34:I40')

    # ---- tabla
    hdr = ['#', 'Concepto', 't', 'Fecha', 'Saldo inicial', 'Pago s/IVA', 'Interés', 'Capital', 'Saldo final',
           'IVA', 'Total c/IVA', 'Flujo arrendador', 'Factor desc.', 'VP flujo', 'Flujo S/R', 'Fecha (XIRR)',
           'Flujo XIRR', 'Flujo cliente (CAT)', 'VP flujo S/R']
    for j, h in enumerate(hdr):
        c = ws.cell(row=T0 - 1, column=1 + j, value=h)
        c.font = f_hdr
        c.fill = fill_hdr
        c.alignment = center
    ws.row_dimensions[T0 - 1].height = 30
    # fila 0 (firma)
    t0 = {1: 0, 2: 'Firma / desembolso', 3: 0, 4: '=C31', 12: '=C41', 16: '=C31', 17: '=C41',
          18: '=C19-C28-C35-C43'}
    for col, v in t0.items():
        c = ws.cell(row=T0, column=col, value=v)
        c.font = f_bold
        c.fill = fill_gris
    for col in range(1, 20):
        ws.cell(row=T0, column=col).fill = fill_gris
    for kk in range(1, NROWS + 1):
        r_ = FIRST + kk - 1
        prev = r_ - 1
        f = {
            'A': kk,
            'B': '=IF(A{r}>$C$38,"",IF(A{r}<=$C$5,"Básico",IF($C$6>0,"Sucesivo","Opción de compra")))',
            'C': '=IF(B{r}="","",IF(B{r}="Opción de compra",$C$5,A{r}-$C$7))',
            'D': '=IF(B{r}="","",EDATE($C$33,C{r}))',
            'E': ('=IF(B{r}="","",$C$19)' if kk == 1 else '=IF(B{r}="","",I{p})'),
            'F': '=IF(B{r}="","",IF(B{r}="Básico",$C$24,IF(B{r}="Sucesivo",$C$25,$C$26)))',
            'G': '=IF(B{r}="","",IF(B{r}="Opción de compra",0,(E{r}-$C$7*F{r})*IF(B{r}="Sucesivo",$C$11,$C$9)))',
            'H': '=IF(B{r}="","",F{r}-G{r})',
            'I': '=IF(B{r}="","",E{r}-H{r})',
            'J': '=IF(B{r}="","",IF(OR(inp_Tipo="Arrendamiento Puro",par_BaseIVAFin="Sobre renta",B{r}="Opción de compra"),F{r},G{r})*$C$14)',
            'K': '=IF(B{r}="","",F{r}+J{r})',
            'L': '=IF(B{r}="","",F{r}-IF(A{r}=$C$38,$C$30,0))',
            'M': '=IF(B{r}="","",$C$37*(1+$C$13)^(-C{r}))',
            'N': '=IF(B{r}="","",L{r}*M{r})',
            'O': '=IF(B{r}="","",IF(B{r}="Básico",F{r},0)-IF(A{r}=$C$38,$C$30,0))',
            'P': '=IF(B{r}="",$C$31,D{r})',
            'Q': '=IF(B{r}="",0,L{r})',
            'R': '=IF(B{r}="",0,-F{r})',
            'S': '=IF(B{r}="","",O{r}*M{r})',
        }
        for col, form in f.items():
            c = ws['%s%d' % (col, r_)]
            c.value = form.format(r=r_, p=prev) if isinstance(form, str) else form
            c.font = f_base
            if col in 'EFGHIJKLNOQRS':
                c.number_format = MON
            elif col in 'DP':
                c.number_format = FECHA
            elif col == 'M':
                c.number_format = '0.000000'
    ws.freeze_panes = 'C%d' % T0
    ws.conditional_formatting.add('A%d:S%d' % (FIRST, LAST),
                                  FormulaRule(formula=['$B%d="Sucesivo"' % FIRST], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.conditional_formatting.add('A%d:S%d' % (FIRST, LAST),
                                  FormulaRule(formula=['$B%d="Opción de compra"' % FIRST], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '%d:%d' % (T0 - 1, T0 - 1)


for k, w in enumerate(ws_cr, start=1):
    build_corrida(w, k)


def cr(key, k):
    return "'Corrida %d'!%s" % (k, P[key].replace('C', '$C$', 1) if P[key][0] == 'C' else P[key].replace('I', '$I$', 1))


# ==================================================================== COTIZADOR: resultados
ws = ws_cot
for k in range(4):
    c = ws.cell(row=RES_CLIENTE, column=3 + k, value='Escenario %d' % (k + 1))
    c.font = f_hdr
    c.alignment = center
res = [
    ('Plazo', lambda k: '=INDEX(esc_Plazo,1,%d)&" + "&INDEX(esc_Sucesivo,1,%d)&" meses"' % (k, k), None, None, False),
    ('Monto a financiar', lambda k: '=%s' % cr('M', k), MON, 'res_Monto', False),
    ('Renta mensual (sin IVA)', lambda k: '=%s' % cr('R', k), MON, 'res_Renta', True),
    ('IVA de la renta', lambda k: '=%s' % cr('IVA1', k), MON, 'res_IVA', False),
    ('Renta mensual con IVA', lambda k: '=%s+%s' % (cr('R', k), cr('IVA1', k)), MON, 'res_RentaIVA', True),
    ('Renta plazo sucesivo / opción de compra (sin IVA)', lambda k: '=%s+%s' % (cr('Rs', k), cr('OC', k)), MON, 'res_Final', False),
    ('Pago inicial total (con IVA)', lambda k: '=%s' % cr('PI', k), MON, 'res_PagoInicial', True),
    ('Total de pagos del contrato (con IVA)', lambda k: '=%s' % cr('TotK', k), MON, 'res_Total', False),
    ('CAT informativo sin IVA', lambda k: '=%s' % cr('CAT', k), PCT, 'res_CAT', False),
    ('Deducibilidad estimada de la renta (ISR)', lambda k: '=%s' % cr('ded', k), PCT, 'res_Ded', False),
]
r = RES_CLIENTE + 1
for label, fn, fmt, nm, bold in res:
    ws.cell(row=r, column=2, value=label).font = f_bold if bold else f_base
    for k in range(1, 5):
        c = ws.cell(row=r, column=2 + k)
        calc(c, fn(k), fmt, bold)
        c.font = Font(name=FONT, size=10, bold=bold, color='008000')
        c.alignment = Alignment(horizontal='center')
    if nm:
        name(nm, absref(ws.title, 'C%d:F%d' % (r, r)))
    r += 1
ws.cell(row=r - 1, column=8, value='Automóviles: límite de deducción de la renta por día (Configuracion). Estimado; confirme con su área fiscal.').font = f_note
ws.cell(row=r - 1, column=8).alignment = Alignment(wrap_text=True, vertical='center')
ws.cell(row=r - 2, column=8, value='CAT calculado con XIRR sobre los flujos del cliente; informativo, no sustituye el cálculo oficial de Banxico.').font = f_note
ws.cell(row=r - 2, column=8).alignment = Alignment(wrap_text=True, vertical='center')

r += 1
RES_INT = r
header_bar(ws, r, 2, 8, '5. ANÁLISIS INTERNO – RENTABILIDAD (no aparece en la carta)')
for k in range(4):
    c = ws.cell(row=r, column=3 + k, value='Escenario %d' % (k + 1))
    c.alignment = center
r += 1
intr = [
    ('Tasa anual aplicada', lambda k: '=%s' % cr('Tasa', k), PCT, 'res_Tasa'),
    ('Spread (tasa cliente − fondeo)', lambda k: '=%s' % cr('Spread', k), PCT, None),
    ('TIR anual efectiva del arrendador', lambda k: '=%s' % cr('TIR', k), PCT, 'res_TIR'),
    ('TIR nominal anual (cap. mensual)', lambda k: '=%s' % cr('TIRn', k), PCT, None),
    ('Cash margin C/R ($)', lambda k: '=%s' % cr('CM', k), MON, 'res_CM'),
    ('Cash margin C/R (% del monto financiado)', lambda k: '=%s' % cr('CMp', k), PCT, 'res_CMpct'),
    ('Cash margin S/R ($)', lambda k: '=%s' % cr('CMs', k), MON, None),
    ('Cash margin S/R (%)', lambda k: '=%s' % cr('CMsp', k), PCT, None),
    ('Comisión del promotor ($)', lambda k: '=%s' % cr('cp', k), MON, None),
    ('Cash margin NETO ($)', lambda k: '=%s' % cr('CMn', k), MON, 'res_CMneto'),
    ('Cash margin NETO (%)', lambda k: '=%s' % cr('CMnp', k), PCT, 'res_CMnetoPct'),
    ('CM mínimo / TIR mínima (rate card)', lambda k: '=TEXT(%s,"0.00%%")&" / "&TEXT(%s,"0.00%%")' % (cr('obj', k), cr('tirmin', k)), None, None),
    ('Dictamen vs. rate card', lambda k: '=%s' % cr('Dict', k), None, 'res_Dictamen'),
    ('Renta mínima para el CM objetivo (sin IVA)', lambda k: '=%s' % cr('Rmin', k), MON, None),
    ('Tasa anual mínima para el CM objetivo', lambda k: '=%s' % cr('TasaMin', k), PCT, 'res_TasaMin'),
    ('Margen mínimo sobre TIIE', lambda k: '=%s' % cr('MargMin', k), PCT, None),
    ('Cuadre de la corrida (saldo final)', lambda k: '=IF(ABS(%s)<0.01,"OK","Dif. "&TEXT(%s,"#,##0.00"))' % (cr('Saldo', k), cr('Saldo', k)), None, None),
]
for label, fn, fmt, nm in intr:
    ws.cell(row=r, column=2, value=label).font = f_base
    for k in range(1, 5):
        c = ws.cell(row=r, column=2 + k)
        calc(c, fn(k), fmt)
        c.font = Font(name=FONT, size=10, color='008000')
        c.alignment = Alignment(horizontal='center')
    if nm:
        name(nm, absref(ws.title, 'C%d:F%d' % (r, r)))
    if label.startswith('Dictamen'):
        DICT_ROW = r
    r += 1
ws.conditional_formatting.add('C%d:F%d' % (DICT_ROW, DICT_ROW), CellIsRule(operator='equal', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
ws.conditional_formatting.add('C%d:F%d' % (DICT_ROW, DICT_ROW), CellIsRule(operator='equal', formula=['"NO CUMPLE"'], fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(name=FONT, bold=True, color='9C0006')))
ws.conditional_formatting.add('C%d:F%d' % (DICT_ROW, DICT_ROW), CellIsRule(operator='equal', formula=['"CM OK / TIR BAJA"'], fill=PatternFill('solid', fgColor='FFEB9C'), font=Font(name=FONT, bold=True, color='9C5700')))
ws.cell(row=RES_INT + 1, column=8, value='C/R = con residual (incluye pago final). S/R = sin residual. Descontado a la tasa de fondeo, igual que el archivo original.').font = f_note
ws.cell(row=RES_INT + 1, column=8).alignment = Alignment(wrap_text=True, vertical='center')
ws.cell(row=DICT_ROW + 2, column=8, value='Macro "Aplicar tasa mínima": copia esta tasa a la fila de tasas del escenario.').font = f_note

r += 1
header_bar(ws, r, 2, 8, '6. VALIDACIONES Y ALERTAS')
r += 1
ALERT_ROW = r
for k in range(1, 5):
    E = lambda nm: 'INDEX(%s,1,%d)' % (nm, k)
    f = ('=IF(inp_Valor<=0,"Capture el valor del equipo. ","")'
         '&IF(inp_FechaPrimera<inp_FechaFirma,"La 1a renta es anterior a la firma. ","")'
         '&IF({tasa}<inp_Fondeo,"Spread negativo (tasa < fondeo). ","")'
         '&IF({tasa}=0,"Tasa en cero. ","")'
         '&IF({res}+{eng}>=1,"Enganche + residual >= 100%. ","")'
         '&IF(AND({pf}<>"",{pf}<>{res}),"Pago final distinto al residual: la tasa real difiere de la pactada. ","")'
         '&IF({cmp}<{obj},"CM neto por debajo del mínimo. ","")'
         '&IF(N({tir})<{tirmin},"TIR por debajo de la mínima. ","")'
         '&IF({res}>{resmax},"Residual mayor al máximo de política. ","")'
         '&IF(ABS({saldo})>=0.01,"La corrida no cuadra. ","")').format(
        tasa=E('esc_Tasa'), res=E('esc_Residual'), eng=E('esc_Enganche'), pf=E('esc_PagoFinal'),
        cmp=cr('CMnp', k), saldo=cr('Saldo', k), obj=cr('obj', k), tir=cr('TIR', k), tirmin=cr('tirmin', k), resmax=cr('resmax', k))
    f = '=IF(%s="","Sin alertas",%s)' % (f[1:], f[1:])
    c = ws.cell(row=r, column=2 + k, value=f)
    c.font = Font(name=FONT, size=9, color='C00000')
    c.alignment = wrap
    c.border = box
ws.cell(row=r, column=2, value='Alertas por escenario').font = f_bold
ws.row_dimensions[r].height = 75
ws.conditional_formatting.add('C%d:F%d' % (r, r), CellIsRule(operator='equal', formula=['"Sin alertas"'], font=Font(name=FONT, color='006100')))
name('res_Alertas', absref(ws.title, 'C%d:F%d' % (r, r)))
ws.freeze_panes = 'C5'
ws.print_area = 'A1:H%d' % r
ws.page_setup.orientation = 'portrait'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ==================================================================== CARTA
ws = ws_car
for col, w in zip('ABCDEFG', [2, 40, 17, 17, 17, 17, 2]):
    ws.column_dimensions[col].width = w
logo(ws, 'B1', 62)
for i, nm in enumerate(['par_Arrendador', 'sel_Dir1', 'sel_Dir2', 'sel_Dir3'], start=1):
    c = ws.cell(row=i, column=4, value='=' + nm)
    c.font = Font(name=FONT, size=9, bold=(i == 1), color='404040')
    c.alignment = Alignment(horizontal='right')
    ws.merge_cells(start_row=i, start_column=4, end_row=i, end_column=6)
MESES = '"enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"'
ws['B6'] = '=par_Ciudad&", a "&DAY(inp_Fecha)&" de "&CHOOSE(MONTH(inp_Fecha),%s)&" de "&YEAR(inp_Fecha)' % MESES
ws['B6'].font = f_base
ws['F6'] = '="Folio: "&inp_Folio'
ws['F6'].font = f_bold
ws['F6'].alignment = Alignment(horizontal='right')
ws['B8'] = '=UPPER(inp_Cliente)'
ws['B8'].font = Font(name=FONT, size=11, bold=True)
ws['B9'] = '=IF(inp_Contacto="","","At\'n: "&inp_Contacto)'
ws['B9'].font = f_base
ws['B11'] = '="Por medio de la presente nos permitimos poner a su consideración la cotización de "&inp_Tipo&", con base en la siguiente información:"'
ws['B11'].font = f_base
ws['B11'].alignment = wrap
ws.merge_cells('B11:F11')
ws.row_dimensions[11].height = 28
header_bar(ws, 13, 2, 6, 'TÉRMINOS Y CONDICIONES')
info = [
    ('Tipo de operación:', '=inp_Tipo&" en "&inp_Moneda'),
    ('Arrendatario:', '=inp_Cliente'),
    ('Arrendador:', '=par_Arrendador'),
    ('Obligado solidario:', '=inp_Obligado'),
    ('Proveedor:', '=IF(inp_Proveedor="","Por definir",inp_Proveedor)'),
    ('Equipo:', '=inp_Equipo'),
    ('Valor del equipo:', '=TEXT(inp_Valor,"$#,##0.00")&" más IVA"'),
    ('Tipo de cambio (informativo):', '=IF(inp_Moneda="Moneda Nacional","No aplica",TEXT(inp_TC,"#,##0.0000"))'),
    ('Forma de pago de las rentas:', '="Mensual, "&LOWER(inp_Modalidad)&IF(inp_Modalidad="Anticipado"," (al inicio de cada mes)"," (al final de cada mes)")'),
]
r = 14
for label, f in info:
    ws.cell(row=r, column=2, value=label).font = f_bold
    c = ws.cell(row=r, column=3, value=f)
    c.font = f_base
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    r += 1
r += 1
TBL = r
ws.cell(row=r, column=2, value='CONCEPTO').font = f_hdr
for col in range(2, 7):
    ws.cell(row=r, column=col).fill = fill_hdr
    ws.cell(row=r, column=col).alignment = center
for k in range(1, 5):
    c = ws.cell(row=r, column=2 + k, value='=IF(INDEX(esc_Incluir,1,%d)="Sí","OPCIÓN "&%d,"")' % (k, k))
    c.font = f_hdr
ws.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[r].height = 20
r += 1


def carta_row(label, fn, fmt=MON, bold=False, sub=False, top=False):
    global r
    c0 = ws.cell(row=r, column=2, value=label)
    c0.font = f_bold if (bold or sub) else f_base
    if sub:
        for col in range(2, 7):
            ws.cell(row=r, column=col).fill = fill_sub
    for k in range(1, 5):
        c = ws.cell(row=r, column=2 + k)
        if fn:
            c.value = '=IF(INDEX(esc_Incluir,1,%d)="Sí",%s,"")' % (k, fn(k))
        c.number_format = fmt or 'General'
        c.font = f_bold if bold else f_base
        c.alignment = Alignment(horizontal='center')
        if top:
            c.border = Border(top=Side(style='thin', color='000000'))
    if top:
        c0.border = Border(top=Side(style='thin', color='000000'))
    r += 1


carta_row('Plazo básico', lambda k: 'INDEX(esc_Plazo,1,%d)&" meses"' % k, None, bold=True)
carta_row('Plazo sucesivo', lambda k: 'IF(INDEX(esc_Sucesivo,1,%d)>0,INDEX(esc_Sucesivo,1,%d)&IF(INDEX(esc_Sucesivo,1,%d)=1," mes"," meses"),"Opción de compra")' % (k, k, k), None)
carta_row('Monto a financiar (más IVA)', lambda k: cr('M', k))
carta_row('PAGO INICIAL (IVA incluido)', None, sub=True)
carta_row('Anticipo (enganche / renta extraordinaria)', lambda k: '%s*(1+par_IVA)' % cr('eng', k))
carta_row('Comisión por apertura', lambda k: '%s*(1+par_IVA)' % cr('com', k))
carta_row('Depósito en garantía', lambda k: cr('dep', k))
carta_row('Renta proporcional', lambda k: '%s*(1+par_IVA)' % cr('rp', k))
carta_row('Seguro, GPS y otros de contado', lambda k: '%s*(1+par_IVA)' % cr('Cont', k))
carta_row('Gastos de investigación / validación', lambda k: '%s*(1+par_IVA)' % cr('gi', k))
carta_row('Total pago inicial', lambda k: cr('PI', k), bold=True, top=True)
carta_row('RENTAS MENSUALES', None, sub=True)
carta_row('Renta mensual (más IVA)', lambda k: cr('R', k))
carta_row('IVA', lambda k: cr('IVA1', k))
carta_row('Renta mensual con IVA', lambda k: '%s+%s' % (cr('R', k), cr('IVA1', k)), bold=True, top=True)
carta_row('Renta plazo sucesivo / opción de compra (más IVA)', lambda k: '%s+%s' % (cr('Rs', k), cr('OC', k)))
carta_row('Seguro y GPS financiados (incluidos en la renta)', lambda k: 'IF(%s+%s=0,"No aplica",TEXT(%s+%s,"$#,##0.00")&" + IVA")' % (cr('segf', k), cr('gpsf', k), cr('segf', k), cr('gpsf', k)), None)
for rr_ in range(TBL + 1, r):
    for col in range(2, 7):
        cc = ws.cell(row=rr_, column=col)
        if not cc.border.top.style:
            cc.border = Border(bottom=Side(style='hair', color='BFBFBF'))
r += 1
ws.cell(row=r, column=2, value='Cifras en la moneda de la operación. Las rentas causan IVA a la tasa vigente. Sujeto a aprobación de crédito.').font = f_note
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
r += 2


def texto(label_row_cfg, est_chars):
    global r
    cfgrow = TXT_ROW0 + label_row_cfg
    lab = ws.cell(row=r, column=2, value="=Configuracion!$B$%d" % cfgrow)
    lab.font = f_bold
    lab.alignment = wrap
    c = ws.cell(row=r, column=3, value='=SUBSTITUTE(SUBSTITUTE(Configuracion!$C$%d,"{ARRENDADOR}",par_Arrendador),"{VIGENCIA}",par_Vigencia)' % cfgrow)
    c.font = Font(name=FONT, size=9)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='justify')
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    lines = est_chars // 95 + 1
    ws.row_dimensions[r].height = max(14, 12 * lines + 4)
    r += 1


for i, (label, txt) in enumerate(TEXTOS[:-1]):
    if i == 2:
        ws.cell(row=r, column=2, value='Esta propuesta se basa en los siguientes términos y condiciones adicionales:').font = f_bold
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
    texto(i, len(txt) + 20)
r += 1
c = ws.cell(row=r, column=2, value="=Configuracion!$C$%d" % (TXT_ROW0 + len(TEXTOS) - 1))
c.font = f_base
c.alignment = wrap
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
ws.row_dimensions[r].height = 28
r += 3
ws.cell(row=r, column=2, value='A T E N T A M E N T E').font = f_bold
ws.cell(row=r, column=4, value='PROPUESTA ACEPTADA POR:').font = f_bold
r += 3
SIG = r
ws.cell(row=r, column=2).border = Border(bottom=Side(style='thin', color='000000'))
for lab in ['Nombre:', 'Puesto:', 'Fecha:', 'Firma:']:
    ws.cell(row=r, column=4, value=lab).font = f_base
    for col in (5, 6):
        ws.cell(row=r, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    r += 1
ws.cell(row=SIG + 1, column=2, value='=sel_PromNombre').font = f_bold
ws.cell(row=SIG + 2, column=2, value='=IF(sel_PromPuesto="","Promoción",sel_PromPuesto)&" – "&par_Comercial').font = f_base
ws.cell(row=SIG + 3, column=2, value='=sel_PromCorreo&IF(par_Contacto="",""," · "&par_Contacto)').font = f_note
LAST_CARTA = r
ws.print_area = 'A1:G%d' % LAST_CARTA
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = page.PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.6)
ws.oddFooter.center.text = 'Página &P de &N'
ws.oddFooter.center.size = 8

# ==================================================================== TABLA DE PAGOS
ws = ws_tab
for col, w in zip('ABCDEFGHI', [2, 7, 13, 22, 17, 15, 17, 18, 2]):
    ws.column_dimensions[col].width = w
logo(ws, 'B1', 55)
ws['D1'] = 'TABLA DE PAGOS'
ws['D1'].font = f_title
ws.merge_cells('D1:H1')
ws['D2'] = '=par_Arrendador'
ws['D2'].font = f_sub
ws.merge_cells('D2:H2')
lab = [('Cliente:', '=inp_Cliente'), ('Equipo:', '=inp_Equipo'), ('Folio:', '=inp_Folio')]
for i, (a, b) in enumerate(lab):
    ws.cell(row=4 + i, column=2, value=a).font = f_bold
    ws.cell(row=4 + i, column=4, value=b).font = f_base
    ws.merge_cells(start_row=4 + i, start_column=4, end_row=4 + i, end_column=8)
ws['B7'] = 'Escenario a mostrar:'
ws['B7'].font = f_bold
ws['F7'] = 'Se elige en el Cotizador'
ws['F7'].font = f_note
ws.merge_cells('B7:C7')
ws['D7'] = '=inp_EscTabla'
ws['D7'].font = f_link
ws['D7'].alignment = Alignment(horizontal='left')
CH = lambda ref: 'CHOOSE(inp_EscTabla,%s)' % ','.join("'Corrida %d'!%s" % (k, ref) for k in range(1, 5))
ws['B8'] = '="Plazo: "&%s&" meses + "&%s&" sucesivo  |  Tasa anual: "&TEXT(%s,"0.00%%")&"  |  Pago "&LOWER(inp_Modalidad)&"  |  "&inp_Moneda' % (CH('$C$5'), CH('$C$6'), CH('$C$8'))
ws['B8'].font = f_base
ws.merge_cells('B8:H8')
hdr = ['No.', 'Fecha', 'Concepto', 'Importe sin IVA', 'IVA', 'Total con IVA', 'Acumulado con IVA']
for j, h in enumerate(hdr):
    c = ws.cell(row=10, column=2 + j, value=h)
    c.font = f_hdr
    c.fill = fill_hdr
    c.alignment = center
# pago inicial
r = 11
ws.cell(row=r, column=2, value=0)
ws.cell(row=r, column=3, value='=inp_FechaFirma').number_format = FECHA
ws.cell(row=r, column=4, value='Pago inicial')
ws.cell(row=r, column=5, value='=%s+%s' % (CH('$I$17'), CH('$I$19'))).number_format = MON
ws.cell(row=r, column=6, value='=%s' % CH('$I$18')).number_format = MON
ws.cell(row=r, column=7, value='=E11+F11').number_format = MON
ws.cell(row=r, column=8, value='=G11').number_format = MON
for col in range(2, 9):
    ws.cell(row=r, column=col).fill = fill_gris
    ws.cell(row=r, column=col).font = f_base
for kk in range(1, NROWS + 1):
    r = 11 + kk
    src = FIRST + kk - 1
    ws.cell(row=r, column=2, value='=IF(%s="","",%d)' % (CH('$B$%d' % src), kk))
    ws.cell(row=r, column=3, value='=IF(B%d="","",%s)' % (r, CH('$D$%d' % src))).number_format = FECHA
    ws.cell(row=r, column=4, value='=IF(B%d="","",IF(%s="Básico","Renta mensual",IF(%s="Sucesivo","Renta plazo sucesivo","Opción de compra")))' % (r, CH('$B$%d' % src), CH('$B$%d' % src)))
    ws.cell(row=r, column=5, value='=IF(B%d="","",%s)' % (r, CH('$F$%d' % src))).number_format = MON
    ws.cell(row=r, column=6, value='=IF(B%d="","",%s)' % (r, CH('$J$%d' % src))).number_format = MON
    ws.cell(row=r, column=7, value='=IF(B%d="","",E%d+F%d)' % (r, r, r)).number_format = MON
    ws.cell(row=r, column=8, value='=IF(B%d="","",H%d+G%d)' % (r, r - 1, r)).number_format = MON
    for col in range(2, 9):
        ws.cell(row=r, column=col).font = f_base
        ws.cell(row=r, column=col).alignment = Alignment(horizontal='center') if col <= 4 else Alignment()
TOTR = 11 + NROWS + 1
ws.cell(row=TOTR, column=4, value='TOTAL').font = f_bold
for col, L_ in ((5, 'E'), (6, 'F'), (7, 'G')):
    c = ws.cell(row=TOTR, column=col, value='=SUM(%s11:%s%d)' % (L_, L_, TOTR - 1))
    c.number_format = MON
    c.font = f_bold
    c.border = Border(top=Side(style='thin', color='000000'))
ws.cell(row=TOTR + 1, column=2, value='El depósito en garantía se incluye en el pago inicial y se aplica en el último pago. Cifras sujetas a aprobación de crédito.').font = f_note
ws.merge_cells(start_row=TOTR + 1, start_column=2, end_row=TOTR + 1, end_column=8)
ws.conditional_formatting.add('B12:H%d' % (TOTR - 1), FormulaRule(formula=['AND($B12<>"",MOD($B12,2)=0)'], fill=PatternFill('solid', fgColor='F7F4FB')))
ws.freeze_panes = 'B11'
ws.print_area = 'A1:I%d' % (TOTR + 1)
ws.print_title_rows = '10:10'
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddFooter.center.text = 'Página &P de &N'

# ==================================================================== DOCUMENTOS (Propuesta, Pago inicial, Promesa, Comité)
# Todos usan el escenario elegido en el Cotizador (inp_EscTabla) a través de CH().
fill_box = PatternFill('solid', fgColor=MORADO_CL)
f_box = Font(name=FONT, size=10, bold=True, color=MORADO)
f_it = Font(name=FONT, size=10, italic=True)
f_big = Font(name=FONT, size=12, bold=True, color=MORADO)
thin_b = Border(bottom=Side(style='hair', color='BFBFBF'))
total_b = Border(top=Side(style='thin', color='000000'))


def doc_sheet(ws, widths, last_row, landscape=False):
    for i, w in enumerate(widths):
        ws.column_dimensions[L(i + 1)].width = w
    ws.print_area = 'A1:%s%d' % (L(len(widths)), last_row)
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = page.PageMargins(left=0.45, right=0.45, top=0.45, bottom=0.5)


def membrete(ws, c_right1, c_right2):
    logo(ws, 'B1', 60)
    lines = [('=par_Arrendador', True), ('=sel_Dir1', False), ('=sel_Dir2', False), ('=sel_Dir3', False),
             ('=IF(sel_Tel="","",("Tel. "&sel_Tel&"  ·  "))&sel_PromCorreo', False)]
    for i, (f, b) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=c_right1, value=f)
        c.font = Font(name=FONT, size=9, bold=b, color=MORADO if b else '404040')
        c.alignment = Alignment(horizontal='right')
        ws.merge_cells(start_row=i, start_column=c_right1, end_row=i, end_column=c_right2)
    for col in range(2, c_right2 + 1):
        ws.cell(row=6, column=col).border = Border(bottom=Side(style='medium', color=MORADO))


def box_title(ws, row, c1, c2, text):
    c = ws.cell(row=row, column=c1, value=text)
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = fill_hdr
        ws.cell(row=row, column=col).font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 18


def kv(ws, row, cl, cv, label, formula, fmt=MON, bold=False, top=False, fill=None):
    a = ws.cell(row=row, column=cl, value=label)
    b = ws.cell(row=row, column=cv, value=formula)
    a.font = Font(name=FONT, size=10, bold=bold)
    b.font = Font(name=FONT, size=10, bold=bold)
    b.alignment = Alignment(horizontal='right')
    if fmt:
        b.number_format = fmt
    for c in (a, b):
        c.border = total_b if top else thin_b
        if fill:
            c.fill = fill


def parrafo(ws, row, c1, c2, formula, height, font=None, align='justify'):
    c = ws.cell(row=row, column=c1, value=formula)
    c.font = font or Font(name=FONT, size=9)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal=align)
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.row_dimensions[row].height = height


IVA1 = '(1+par_IVA)'
FECHA_TXT = 'DAY({d})&" de "&CHOOSE(MONTH({d}),%s)&" de "&YEAR({d})' % MESES


# ----------------------------------------------------------------- PROPUESTA
ws = ws_pro
membrete(ws, 4, 6)
ws['B8'] = '="Apreciable: "&IF(inp_Contacto="",inp_Cliente,inp_Contacto)'
ws['B8'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E8'] = '=par_Ciudad&", a "&' + FECHA_TXT.format(d='inp_Fecha')
ws['E8'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E8'].alignment = Alignment(horizontal='right')
ws.merge_cells('E8:F8')
ws['B9'] = '=UPPER(inp_Cliente)'
ws['B9'].font = Font(name=FONT, size=10, bold=True)
ws['F9'] = '="Folio: "&inp_Folio'
ws['F9'].font = f_bold
ws['F9'].alignment = Alignment(horizontal='right')
ws['B10'] = 'P r e s e n t e'
ws['B10'].font = Font(name=FONT, size=10, bold=True, italic=True)
parrafo(ws, 12, 2, 6, '=par_Arrendador&" se complace en poner a su amable consideración la presente propuesta para celebrar una operación de "&LOWER(inp_Tipo)&IF(inp_Producto="FASTPLUS"," (esquema FASTPLUS)","")&" sobre el equipo que a continuación se describe:"', 28, Font(name=FONT, size=10))
box_title(ws, 14, 2, 6, 'DESCRIPCIÓN DEL EQUIPO')
for i, (lab, f) in enumerate([('Tipo de activo', '=inp_TipoActivo'), ('Descripción', '=inp_Equipo'),
                               ('Proveedor', '=IF(inp_Proveedor="","Por definir",inp_Proveedor)'),
                               ('Tipo de operación', '=inp_Tipo&" en "&inp_Moneda')]):
    rr = 15 + i
    ws.cell(row=rr, column=2, value=lab).font = f_it
    c = ws.cell(row=rr, column=3, value=f)
    c.font = Font(name=FONT, size=10)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
# bloque izquierdo: valor y financiados
box_title(ws, 20, 2, 3, 'VALOR DEL EQUIPO Y CONCEPTOS FINANCIADOS (IVA incluido)')
kv(ws, 21, 2, 3, 'Valor del equipo', '=inp_Valor*' + IVA1)
kv(ws, 22, 2, 3, 'Seguro financiado ¹', '=%s*%s' % (CH('$C$17'), IVA1))
kv(ws, 23, 2, 3, 'GPS, mantenimiento y otros financiados ¹', '=(%s+%s)*%s' % (CH('$C$18'), CH('$C$42'), IVA1))
kv(ws, 24, 2, 3, '=IF(inp_UsoAnticipo="Renta extraordinaria","(−) Renta extraordinaria","(−) Anticipo / enganche")', '=-%s*%s' % (CH('$C$16'), IVA1))
kv(ws, 25, 2, 3, 'Monto total financiado', '=SUM(C21:C24)', bold=True, top=True, fill=fill_box)
# bloque derecho: plazo
box_title(ws, 20, 5, 6, 'PLAZO Y FORMA DE PAGO')
kv(ws, 21, 5, 6, 'Plazo del arrendamiento', '=%s&" meses"' % CH('$C$5'), None)
kv(ws, 22, 5, 6, 'Plazo sucesivo', '=IF(%s>0,%s&IF(%s=1," mes"," meses"),"Opción de compra")' % (CH('$C$6'), CH('$C$6'), CH('$C$6')), None)
kv(ws, 23, 5, 6, 'Rentas', '="Mensuales, "&LOWER(inp_Modalidad)&"s"', None)
kv(ws, 24, 5, 6, 'Primera renta', '=inp_FechaPrimera', FECHA)
kv(ws, 25, 5, 6, 'Último pago', '=%s' % CH('$C$52'), FECHA)
# pagos mensuales
box_title(ws, 27, 2, 3, 'PAGOS MENSUALES (no incluyen IVA)')
kv(ws, 28, 2, 3, 'Renta básica del equipo', '=%s' % CH('$C$48'))
kv(ws, 29, 2, 3, 'Pago por seguro financiado', '=%s' % CH('$C$49'))
kv(ws, 30, 2, 3, 'GPS, mantenimiento y otros', '=%s' % CH('$C$50'))
kv(ws, 31, 2, 3, 'Total pago mensual', '=%s' % CH('$C$24'), bold=True, top=True)
kv(ws, 32, 2, 3, 'IVA', '=%s' % CH('$I$21'))
kv(ws, 33, 2, 3, 'Total pago mensual con IVA', '=C31+C32', bold=True, top=True, fill=fill_box)
# pago inicial
box_title(ws, 27, 5, 6, 'PAGO INICIAL')
kv(ws, 28, 5, 6, '=IF(inp_UsoAnticipo="Renta extraordinaria","Renta extraordinaria","Anticipo / enganche")', '=%s' % CH('$C$16'))
kv(ws, 29, 5, 6, 'Comisión por apertura', '=%s' % CH('$C$28'))
kv(ws, 30, 5, 6, 'Renta proporcional', '=%s' % CH('$C$35'))
kv(ws, 31, 5, 6, 'Seguro, GPS y otros de contado', '=%s' % CH('$I$29'))
kv(ws, 32, 5, 6, 'Gastos de investigación / validación ²', '=%s' % CH('$C$43'))
kv(ws, 33, 5, 6, 'Subtotal pago inicial', '=SUM(F28:F32)', bold=True, top=True)
kv(ws, 34, 5, 6, 'Impuesto al Valor Agregado', '=F33*par_IVA')
kv(ws, 35, 5, 6, 'Depósito en garantía (no causa IVA)', '=%s' % CH('$C$30'))
kv(ws, 36, 5, 6, 'TOTAL PAGO INICIAL', '=F33+F34+F35', bold=True, top=True, fill=fill_box)
# residual y vigencia
box_title(ws, 35, 2, 3, 'AL TÉRMINO DEL PLAZO (no incluye IVA)')
kv(ws, 36, 2, 3, '=IF(%s>0,"Renta plazo sucesivo ("&%s&" mes(es))","Opción de compra")' % (CH('$C$6'), CH('$C$6')), '=IF(%s>0,%s,%s)' % (CH('$C$6'), CH('$C$25'), CH('$C$26')))
kv(ws, 37, 2, 3, 'Valor residual', '=%s' % CH('$C$21'))
kv(ws, 38, 2, 3, 'Vigencia de la propuesta', '=WORKDAY(inp_Fecha,par_Vigencia)', FECHA, bold=True)
kv(ws, 39, 2, 3, '=IF(%s<1,"Renta deducible estimada (ISR)","")' % CH('$C$51'), '=IF(%s<1,%s,"")' % (CH('$C$51'), CH('$C$51')), PCT)
parrafo(ws, 41, 2, 6, '="¹ Montos financiados por el plazo total del arrendamiento.  ² Gastos de investigación / validación de cuenta."&IF(%s<1,"  Deducibilidad estimada conforme al límite diario de la Ley del ISR para automóviles; confirme con su asesor fiscal.","")' % CH('$C$51'), 24, Font(name=FONT, size=8, italic=True))
parrafo(ws, 43, 2, 6, 'Cualquier cambio respecto a las condiciones establecidas dentro de la aprobación de crédito correspondiente invalida la presente propuesta, misma que únicamente será posible ejercer durante su vigencia. Sujeto a aprobación de crédito.', 26, Font(name=FONT, size=9, italic=True))
parrafo(ws, 44, 2, 6, '="Para formalizar la presente operación será necesario: 1) Realizar el pago inicial por la cantidad de "&TEXT(F36,"$#,##0.00")&" y 2) Firmar el contrato de arrendamiento correspondiente."', 26, Font(name=FONT, size=9, italic=True, bold=True))
parrafo(ws, 45, 2, 6, '="En "&par_Comercial&" nos esforzamos por otorgarle un excelente servicio. Si tiene cualquier duda o comentario respecto a la presente propuesta, no dude en contactar a su ejecutivo "&sel_PromNombre&IF(sel_PromCorreo="",""," ("&sel_PromCorreo&")")&"."', 26, Font(name=FONT, size=9, italic=True))
ws['B48'] = 'Cordialmente'
ws['E48'] = 'Propuesta aceptada por:'
for c in ('B48', 'E48'):
    ws[c].font = Font(name=FONT, size=10, bold=True, italic=True)
for col in (2, 5):
    ws.cell(row=51, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    ws.cell(row=51, column=col + 1).border = Border(bottom=Side(style='thin', color='000000'))
ws['B52'] = '=sel_PromNombre'
ws['B53'] = '=IF(sel_PromPuesto="","Promoción",sel_PromPuesto)'
ws['B54'] = '=par_Arrendador'
ws['E52'] = '=IF(inp_Contacto="","Representante legal",inp_Contacto)'
ws['E53'] = '=inp_Cliente'
for c in ('B52', 'E52'):
    ws[c].font = f_bold
for c in ('B53', 'B54', 'E53'):
    ws[c].font = Font(name=FONT, size=9)
doc_sheet(ws, [2, 34, 17, 3, 34, 17, 2], 55)

# ----------------------------------------------------------------- PAGO INICIAL
ws = ws_pin
membrete(ws, 4, 6)
ws['B8'] = 'INSTRUCCIONES PARA EL PAGO INICIAL'
ws['B8'].font = f_big
for i, (lab, f) in enumerate([('Cliente', '=inp_Cliente'), ('Atención', '=inp_Contacto&""'), ('Equipo', '=inp_Equipo'),
                               ('Folio / plazo', '=inp_Folio&"  ·  "&%s&" meses"' % CH('$C$5')),
                               ('El anticipo se factura como', '=inp_UsoAnticipo')]):
    rr = 10 + i
    a = ws.cell(row=rr, column=2, value=lab)
    a.font = f_bold
    a.fill = fill_box
    c = ws.cell(row=rr, column=3, value=f)
    c.font = f_base
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
box_title(ws, 16, 2, 3, 'CONCEPTOS DEL PAGO INICIAL')
kv(ws, 17, 2, 3, '=Propuesta!E28', '=Propuesta!F28')
kv(ws, 18, 2, 3, 'Comisión por apertura', '=Propuesta!F29')
kv(ws, 19, 2, 3, 'Renta proporcional', '=Propuesta!F30')
kv(ws, 20, 2, 3, 'Seguro, GPS y otros de contado', '=Propuesta!F31')
kv(ws, 21, 2, 3, 'Gastos de investigación / validación', '=Propuesta!F32')
kv(ws, 22, 2, 3, 'Subtotal antes de IVA', '=Propuesta!F33', bold=True, top=True)
kv(ws, 23, 2, 3, 'IVA', '=Propuesta!F34')
kv(ws, 24, 2, 3, 'Depósito en garantía', '=Propuesta!F35')
kv(ws, 25, 2, 3, 'TOTAL PAGO INICIAL', '=Propuesta!F36', bold=True, top=True, fill=fill_box)
box_title(ws, 16, 5, 6, 'DATOS PARA EL DEPÓSITO / TRANSFERENCIA')
for i, (lab, f) in enumerate([('Banco', '=IF(par_Banco="","(por definir)",par_Banco)'), ('Beneficiario', '=par_Beneficiario'),
                               ('Cuenta', '=IF(par_Cuenta="","(por definir)",par_Cuenta)'),
                               ('CLABE', '=IF(par_CLABE="","(por definir)",par_CLABE)'),
                               ('Convenio', '=IF(par_Convenio="","N/A",par_Convenio)'),
                               ('Referencia', '=inp_Folio'), ('Moneda', '=inp_Moneda'),
                               ('Importe exacto', '=Propuesta!F36')]):
    kv(ws, 17 + i, 5, 6, lab, f, MON if lab == 'Importe exacto' else None, bold=lab == 'Importe exacto')
parrafo(ws, 27, 2, 6, '="Envíe su comprobante de pago a su ejecutivo "&sel_PromNombre&IF(sel_PromCorreo="",""," al correo "&sel_PromCorreo)&", indicando la referencia "&inp_Folio&"."', 28, Font(name=FONT, size=10, bold=True))
parrafo(ws, 29, 2, 6, 'El pago inicial no incluye impuestos locales como tenencia, placas ni gastos de gestoría (en su caso). Estos conceptos, una vez determinados conforme a la reglamentación vigente de cada localidad, se cargarán a la cuenta domiciliada.', 30)
parrafo(ws, 30, 2, 6, 'Bajo la opción de pago de contado del seguro, las anualidades subsecuentes pueden incrementarse según la siniestralidad del equipo o las políticas de la aseguradora.', 26)
ws['B33'] = 'Cordialmente'
ws['E33'] = 'Recibido por:'
for col in (2, 5):
    ws.cell(row=36, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    ws.cell(row=36, column=col + 1).border = Border(bottom=Side(style='thin', color='000000'))
ws['B37'] = '=sel_PromNombre'
ws['B38'] = '=par_Arrendador'
ws['E37'] = '=IF(inp_Contacto="","Representante legal",inp_Contacto)'
ws['E38'] = '=inp_Cliente'
for c in ('B33', 'E33', 'B37', 'E37'):
    ws[c].font = f_bold
for c in ('B38', 'E38'):
    ws[c].font = Font(name=FONT, size=9)
doc_sheet(ws, [2, 34, 17, 3, 26, 26, 2], 39)

# ----------------------------------------------------------------- PROMESA DE VENTA
ws = ws_ven
membrete(ws, 4, 6)
ws['B8'] = '="Apreciable: "&IF(inp_Contacto="",inp_Cliente,inp_Contacto)'
ws['B8'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E8'] = '=par_Ciudad&", a "&' + FECHA_TXT.format(d='inp_Fecha')
ws['E8'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E8'].alignment = Alignment(horizontal='right')
ws.merge_cells('E8:F8')
ws['B9'] = 'P r e s e n t e'
ws['B9'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['B11'] = 'CARTA DE PROMESA DE VENTA AL TÉRMINO DEL ARRENDAMIENTO'
ws['B11'].font = f_big
parrafo(ws, 13, 2, 6, '=par_Arrendador&" (el Promitente Vendedor) venderá, al término del contrato de arrendamiento celebrado con "&inp_Cliente&", a la persona que éste designe (el Promitente Comprador), el equipo usado que a continuación se describe:"', 40, Font(name=FONT, size=10))
box_title(ws, 15, 2, 6, 'CONDICIONES DE LA VENTA')
kv(ws, 16, 2, 6, 'Equipo', '=inp_TipoActivo&" – "&inp_Equipo', None)
kv(ws, 17, 2, 6, 'Valor de venta (IVA incluido)', '=%s*%s' % (CH('$C$23'), IVA1), MON, bold=True)
kv(ws, 18, 2, 6, 'Forma de pago', '=IF(%s=0,"Un solo pago al término del plazo básico",%s&" pago(s) mensual(es) de "&TEXT(%s*%s,"$#,##0.00")&" IVA incluido (plazo sucesivo)")' % (CH('$C$6'), CH('$C$6'), CH('$C$25'), IVA1), None)
kv(ws, 19, 2, 6, 'Plazo para la venta', '="Al concluir "&(%s+%s)&" meses de arrendamiento"' % (CH('$C$5'), CH('$C$6')), None)
kv(ws, 20, 2, 6, 'Fecha estimada de venta', '=%s' % CH('$C$52'), FECHA)
for rr in range(16, 21):
    src = ws.cell(row=rr, column=6)
    v, nf, fo = src.value, src.number_format, copy.copy(src.font)
    src.value = None
    dst = ws.cell(row=rr, column=3)
    dst.value, dst.number_format, dst.font = v, nf, fo
    dst.alignment = Alignment(horizontal='left')
    dst.border = thin_b
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
box_title(ws, 22, 2, 6, 'TÉRMINOS Y CONDICIONES')
parrafo(ws, 23, 2, 6, 'El equipo será entregado por el Promitente Vendedor en el estado en que lo reciba del Arrendatario al término del contrato. El cambio de propietario, así como todos los derechos e impuestos que cause la compraventa, serán a cargo del Promitente Comprador, quien además deberá estar al corriente en todas las obligaciones del contrato de arrendamiento.', 44)
parrafo(ws, 24, 2, 6, 'Las partes celebrarán un contrato de promesa de compraventa bajo los términos anteriores. El precio podrá ajustarse si existen rentas, cargos moratorios u otros adeudos pendientes a la fecha de la venta.', 30)
parrafo(ws, 25, 2, 6, 'Esperamos que la presente propuesta cubra sus expectativas y podamos contar con una respuesta afirmativa en un futuro cercano.', 22)
ws['B28'] = 'Cordialmente'
ws['E28'] = 'Acepto:'
for col in (2, 5):
    ws.cell(row=31, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    ws.cell(row=31, column=col + 1).border = Border(bottom=Side(style='thin', color='000000'))
ws['B32'] = '=sel_PromNombre'
ws['B33'] = '=par_Arrendador'
ws['B33'].font = Font(name=FONT, size=9)
ws['B34'] = 'Promitente Vendedor'
ws['E32'] = '=IF(inp_Contacto="","Representante legal",inp_Contacto)'
ws['E33'] = '=inp_Cliente'
ws['E33'].font = Font(name=FONT, size=9)
ws['E34'] = 'Promitente Comprador'
for c in ('B28', 'E28', 'B32', 'E32', 'B34', 'E34'):
    ws[c].font = f_bold
doc_sheet(ws, [2, 34, 17, 3, 26, 26, 2], 35)

# ----------------------------------------------------------------- RESUMEN COMITÉ
ws = ws_com
logo(ws, 'B1', 50)
ws['D1'] = 'RESUMEN DE OPERACIÓN PARA COMITÉ DE CRÉDITO'
ws['D1'].font = f_title
ws.merge_cells('D1:H1')
ws['D2'] = '="Folio "&inp_Folio&"  ·  "&inp_Producto&"  ·  Escenario "&inp_EscTabla&"  ·  Elaborado: "&TEXT(inp_Fecha,"dd/mm/yyyy")'
ws['D2'].font = f_sub
ws.merge_cells('D2:H2')
ws['D3'] = 'Documento interno – no entregar al cliente'
ws['D3'].font = Font(name=FONT, size=9, bold=True, color='C00000')
box_title(ws, 5, 2, 4, 'DATOS GENERALES')
gen_c = [('Cliente', '=inp_Cliente', None), ('RFC', '=inp_RFC&""', None), ('Obligado solidario', '=inp_Obligado', None),
         ('Promotor', '=sel_PromNombre', None), ('Puesto / región', '=sel_PromPuesto&" · "&sel_PromRegion', None),
         ('Producto', '=inp_Producto', None), ('Tipo de operación', '=inp_Tipo&" · "&inp_Moneda', None),
         ('Tipo de activo', '=inp_TipoActivo', None), ('Equipo', '=inp_Equipo', None),
         ('Proveedor', '=inp_Proveedor&""', None), ('Fecha de firma', '=inp_FechaFirma', FECHA)]
for i, (lab, f, fmt) in enumerate(gen_c):
    kv(ws, 6 + i, 2, 3, lab, f, fmt)
    ws.merge_cells(start_row=6 + i, start_column=3, end_row=6 + i, end_column=4)
    ws.cell(row=6 + i, column=3).alignment = Alignment(horizontal='left')
box_title(ws, 5, 6, 8, 'CONDICIONES FINANCIERAS')
cond = [('Valor del equipo sin IVA', '=inp_Valor', MON), ('Anticipo ($ / %)', '=TEXT(%s,"$#,##0.00")&" / "&TEXT(%s/inp_Valor,"0.00%%")' % (CH('$C$16'), CH('$C$16')), None),
        ('Monto financiado', '=%s' % CH('$C$19'), MON), ('Plazo + sucesivo', '=%s&" + "&%s&" meses · "&inp_Modalidad' % (CH('$C$5'), CH('$C$6')), None),
        ('Tasa anual', '=%s' % CH('$C$8'), PCT), ('Modo de tasa', '=inp_ModoTasa&IF(inp_ModoTasa="TIIE + margen"," ("&TEXT(inp_TIIE,"0.00%%")&" + "&TEXT(%s-inp_TIIE,"0.00%%")&")","")' % CH('$C$8'), None),
        ('Tasa de fondeo / spread', '=TEXT(%s,"0.00%%")&" / "&TEXT(%s,"0.00%%")' % (CH('$C$12'), CH('$I$23')), None),
        ('Residual ($ / % vs. máximo)', '=TEXT(%s,"$#,##0")&" / "&TEXT(%s,"0.0%%")&" (máx. "&TEXT(%s,"0%%")&")"' % (CH('$C$21'), CH('$C$20'), CH('$C$47')), None),
        ('Comisión por apertura', '=%s' % CH('$C$28'), MON), ('Depósito en garantía', '=%s' % CH('$C$30'), MON),
        ('Seguro / GPS / otros financiados', '=%s+%s+%s' % (CH('$C$17'), CH('$C$18'), CH('$C$42')), MON)]
for i, (lab, f, fmt) in enumerate(cond):
    kv(ws, 6 + i, 6, 8, lab, f, fmt)
    ws.merge_cells(start_row=6 + i, start_column=6, end_row=6 + i, end_column=7)
box_title(ws, 18, 2, 4, 'PAGOS DEL CLIENTE')
pag = [('Renta mensual sin IVA', '=%s' % CH('$C$24')), ('Renta mensual con IVA', '=%s+%s' % (CH('$C$24'), CH('$I$21'))),
       ('Pago inicial total', '=%s' % CH('$I$20')), ('Pago final / sucesivo', '=%s+%s' % (CH('$C$25'), CH('$C$26'))),
       ('Total de pagos con IVA', '=%s' % CH('$I$22')), ('CAT informativo', '=%s' % CH('$I$10')),
       ('Deducibilidad estimada', '=%s' % CH('$C$51'))]
for i, (lab, f) in enumerate(pag):
    kv(ws, 19 + i, 2, 3, lab, f, PCT if lab in ('CAT informativo', 'Deducibilidad estimada') else MON)
    ws.merge_cells(start_row=19 + i, start_column=3, end_row=19 + i, end_column=4)
box_title(ws, 18, 6, 8, 'RENTABILIDAD vs. RATE CARD')
rent = [('TIR anual efectiva', '=%s' % CH('$I$8'), PCT), ('TIR mínima (rate card)', '=%s' % CH('$C$46'), PCT),
        ('Cash margin C/R', '=TEXT(%s,"$#,##0")&" · "&TEXT(%s,"0.00%%")' % (CH('$I$4'), CH('$I$5')), None),
        ('Cash margin S/R', '=TEXT(%s,"$#,##0")&" · "&TEXT(%s,"0.00%%")' % (CH('$I$6'), CH('$I$7')), None),
        ('Comisión del promotor', '=%s' % CH('$C$44'), MON),
        ('Cash margin NETO', '=TEXT(%s,"$#,##0")&" · "&TEXT(%s,"0.00%%")' % (CH('$I$25'), CH('$I$26')), None),
        ('CM mínimo (rate card / manual)', '=%s' % CH('$C$45'), PCT)]
for i, (lab, f, fmt) in enumerate(rent):
    kv(ws, 19 + i, 6, 8, lab, f, fmt, bold=lab.startswith('Cash margin NETO'))
    ws.merge_cells(start_row=19 + i, start_column=6, end_row=19 + i, end_column=7)
ws['B27'] = 'DICTAMEN'
ws['B27'].font = f_big
ws['C27'] = '=%s' % CH('$I$27')
ws['C27'].font = Font(name=FONT, size=14, bold=True)
ws.merge_cells('C27:D27')
ws.conditional_formatting.add('C27', CellIsRule(operator='equal', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
ws.conditional_formatting.add('C27', CellIsRule(operator='notEqual', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(name=FONT, bold=True, color='9C0006')))
ws['F27'] = 'Tasa mínima para cumplir:'
ws['F27'].font = f_bold
ws['H27'] = '=%s' % CH('$I$15')
ws['H27'].number_format = PCT
ws['H27'].font = f_bold
ws['B29'] = 'Alertas:'
ws['B29'].font = f_bold
parrafo(ws, 29, 3, 8, '=INDEX(res_Alertas,1,inp_EscTabla)', 30, Font(name=FONT, size=9, color='C00000'), 'left')
box_title(ws, 31, 2, 8, 'COMENTARIOS DEL PROMOTOR / CRÉDITO')
for rr in range(32, 36):
    for col in range(2, 9):
        ws.cell(row=rr, column=col).border = thin_b
box_title(ws, 37, 2, 8, 'AUTORIZACIONES')
for j, (col, lab) in enumerate([(2, 'Promotor'), (4, 'Gerente de Promoción'), (6, 'Crédito'), (8, 'Dirección')]):
    ws.cell(row=41, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    c = ws.cell(row=42, column=col, value=lab)
    c.font = f_bold
    c.alignment = Alignment(horizontal='center')
ws['B43'] = '=sel_PromNombre'
ws['B43'].font = Font(name=FONT, size=9)
ws['B43'].alignment = Alignment(horizontal='center')
doc_sheet(ws, [2, 30, 18, 14, 3, 30, 14, 20], 44, landscape=True)

# ==================================================================== SENSIBILIDAD
ws = ws_sen
for col, w in zip('ABCDEFGHIJ', [2, 22, 15, 15, 15, 15, 15, 15, 15, 2]):
    ws.column_dimensions[col].width = w
ws['B1'] = 'ANÁLISIS DE SENSIBILIDAD – RENTA MENSUAL'
ws['B1'].font = f_title
ws['B2'] = 'Usa el valor del equipo, enganche, residual, comisión y financiados del escenario elegido. Cambie las tasas amarillas libremente.'
ws['B2'].font = f_sub
ws['B4'] = 'Escenario base:'
ws['B4'].font = f_bold
inp(ws['C4'], 1, '0')
dv_list(ws, '"1,2,3,4"', 'C4')
CHs = lambda ref: 'CHOOSE($C$4,%s)' % ','.join("'Corrida %d'!%s" % (k, ref) for k in range(1, 5))
ws['B5'] = 'Monto a financiar:'
ws['C5'] = '=%s' % CHs('$C$19')
ws['C5'].number_format = MON
ws['B6'] = 'Valor residual:'
ws['C6'] = '=%s' % CHs('$C$21')
ws['C6'].number_format = MON
ws['B7'] = 'Tipo de pago:'
ws['C7'] = '=%s' % CHs('$C$7')
ws['D7'] = '=IF(C7=1,"anticipado","vencido")'
for c in ('B5', 'B6', 'B7'):
    ws[c].font = f_bold
for c in ('C5', 'C6', 'C7', 'D7'):
    ws[c].font = Font(name=FONT, size=10, color='008000')
header_bar(ws, 9, 2, 9, 'RENTA MENSUAL SIN IVA  (filas: plazo en meses · columnas: tasa anual)')
ws['B10'] = 'Plazo \\ Tasa'
ws['B10'].font = f_bold
steps = [-0.03, -0.02, -0.01, 0, 0.01, 0.02, 0.03]
for j, s in enumerate(steps):
    inp(ws.cell(row=10, column=3 + j), '=ROUND(%s+(%s),4)' % (CHs('$C$8'), s), PCT)
plazos = [12, 18, 24, 30, 36, 42, 48, 54, 60, 72, 84]
for i, n in enumerate(plazos):
    rr = 11 + i
    inp(ws.cell(row=rr, column=2), n, '0" meses"')
    for j in range(7):
        col = L(3 + j)
        c = ws.cell(row=rr, column=3 + j, value='=IFERROR(PMT(%s$10/12,$B%d,-$C$5,$C$6,$C$7),"")' % (col, rr))
        c.number_format = MON
        c.font = f_base
        c.border = box
R2 = 11 + len(plazos) + 1
header_bar(ws, R2, 2, 9, 'RENTA MENSUAL CON IVA')
for i, n in enumerate(plazos):
    rr = R2 + 1 + i
    c = ws.cell(row=rr, column=2, value='=B%d' % (11 + i))
    c.number_format = '0" meses"'
    c.font = f_bold
    for j in range(7):
        col = L(3 + j)
        c = ws.cell(row=rr, column=3 + j, value='=IFERROR(%s%d*(1+par_IVA),"")' % (col, 11 + i))
        c.number_format = MON
        c.font = f_base
        c.border = box
ws.conditional_formatting.add('C11:I%d' % (10 + len(plazos)), FormulaRule(formula=['AND($B11=%s,C$10=%s)' % (CHs('$C$5'), CHs('$C$8'))], fill=PatternFill('solid', fgColor='C6EFCE')))

# ==================================================================== HISTORIAL
ws = ws_his
hist_cols = ['Folio', 'Fecha', 'Cliente', 'Equipo', 'Promotor', 'Moneda', 'Valor equipo', 'Plazos',
             'Renta c/IVA Esc.1', 'Renta c/IVA Esc.2', 'Renta c/IVA Esc.3', 'Renta c/IVA Esc.4',
             'CM % Esc.1', 'CM % Esc.2', 'CM % Esc.3', 'CM % Esc.4', 'Estatus', 'Comentarios', 'Guardado por', 'Guardado el']
widths = [12, 11, 30, 26, 22, 14, 15, 14, 15, 15, 15, 15, 10, 10, 10, 10, 15, 30, 16, 17]
ws['A1'] = 'HISTORIAL DE COTIZACIONES'
ws['A1'].font = f_title
ws['A2'] = 'La macro "Guardar en historial" agrega una fila por folio. Para reabrir una cotización seleccione su fila y use "Cargar del historial". Las columnas a la derecha guardan todos los datos capturados.'
ws['A2'].font = f_sub
for j, (h, w) in enumerate(zip(hist_cols, widths), start=1):
    c = ws.cell(row=4, column=j, value=h)
    c.font = f_hdr
    c.fill = fill_hdr
    c.alignment = center
    ws.column_dimensions[L(j)].width = w
ws.row_dimensions[4].height = 30
ws.freeze_panes = 'B5'
ws.auto_filter.ref = 'A4:T4'
dv_list(ws, '=lst_Estatus', 'Q5:Q2000')
name('hist_Header', absref(ws.title, 'A4'))

# ==================================================================== INICIO
ws = ws_ini
for col, w in zip('ABCDEF', [2, 34, 70, 3, 3, 3]):
    ws.column_dimensions[col].width = w
logo(ws, 'B1', 70)
ws['C2'] = 'COTIZADOR DE ARRENDAMIENTO'
ws['C2'].font = Font(name=FONT, size=22, bold=True, color=MORADO)
ws['C3'] = 'Arrendamiento puro y financiero · 4 escenarios · cash margin · TIR · CAT · carta y tabla de pagos'
ws['C3'].font = f_sub
ws.row_dimensions[2].height = 30
r = 6
header_bar(ws, r, 2, 3, 'PASOS PARA COTIZAR')
pasos = [
    ('1. Nueva cotización', 'Cinta "Cotizador Arrendamiento" → Nueva cotización (o Ctrl+Shift+N). Limpia la captura y asigna folio.'),
    ('2. Capturar datos', 'En la hoja Cotizador llene solo las celdas AMARILLAS (texto azul): cliente, equipo, valor, fechas y los 4 escenarios.'),
    ('3. Revisar rentabilidad', 'Sección 5 del Cotizador: cash margin, TIR y dictamen (verde = cumple). Si no cumple use "Aplicar tasa mínima".'),
    ('4. Revisar documentos', 'Propuesta (1 escenario, estilo ejecutivo), Carta comparativa (hasta 4 plazos), Pago inicial con datos bancarios, Tabla de pagos y Promesa de venta.'),
    ('5. Entregar', 'Paquete completo en PDF (Ctrl+Shift+P) o Enviar por correo (Outlook). Para crédito: Resumen para comité.'),
    ('6. Guardar', 'Guardar en historial (Ctrl+Shift+G). Desde Historial puede volver a cargar cualquier folio.'),
]
r += 1
for a, b in pasos:
    ws.cell(row=r, column=2, value=a).font = f_bold
    ws.cell(row=r, column=3, value=b).font = f_base
    ws.cell(row=r, column=3).alignment = wrap
    ws.row_dimensions[r].height = 30
    r += 1
r += 1
header_bar(ws, r, 2, 3, 'IR A…')
r += 1
for sh, desc in [('Cotizador', 'Captura y resultados'), ('Carta Cotizacion', 'Carta para el cliente'),
                 ('Tabla de Pagos', 'Calendario de pagos con IVA'), ('Sensibilidad', 'Matriz de rentas por plazo y tasa'),
                 ('Propuesta', 'Propuesta ejecutiva del escenario elegido'),
                 ('Pago Inicial', 'Conceptos e instrucciones de depósito'),
                 ('Promesa de Venta', 'Carta de venta del equipo al término'),
                 ('Resumen Comite', 'Hoja interna para comité de crédito'),
                 ('Historial', 'Registro de cotizaciones'), ('Corrida 1', 'Corrida financiera detallada (también 2, 3 y 4)'),
                 ('Configuracion', 'Empresa, IVA, fondeo, promotores, textos (administrador)')]:
    c = ws.cell(row=r, column=2, value=sh)
    c.hyperlink = "#'%s'!A1" % sh
    c.font = Font(name=FONT, size=10, color='0563C1', underline='single')
    ws.cell(row=r, column=3, value=desc).font = f_base
    r += 1
r += 1
header_bar(ws, r, 2, 3, 'CÓDIGO DE COLORES')
r += 1
leyenda = [('Celda amarilla, texto azul', 'Dato que usted captura', fill_in, f_input),
           ('Texto verde', 'Resultado tomado de otra hoja (no editar)', None, f_link),
           ('Texto negro', 'Fórmula (no editar)', None, f_base)]
for a, b, fl, fo in leyenda:
    c = ws.cell(row=r, column=2, value=a)
    c.font = fo
    if fl:
        c.fill = fl
    ws.cell(row=r, column=3, value=b).font = f_base
    r += 1
r += 1
header_bar(ws, r, 2, 3, 'MACROS DISPONIBLES (cinta "Cotizador Arrendamiento" o Alt+F8)')
r += 1
macros = [
    ('NuevaCotizacion  (Ctrl+Shift+N)', 'Limpia la captura, pone valores por defecto y el siguiente folio.'),
    ('GuardarCotizacion  (Ctrl+Shift+G)', 'Agrega/actualiza el folio en Historial con todos los datos capturados.'),
    ('CargarCotizacion', 'Recupera una cotización del Historial (fila seleccionada o folio).'),
    ('ExportarPaqueteCliente  (Ctrl+Shift+P)', 'Un PDF con Propuesta + Carta comparativa + Pago inicial + Tabla de pagos.'),
    ('ExportarPropuestaPDF / ExportarCartaPDF', 'Propuesta ejecutiva o carta comparativa en PDF.'),
    ('ExportarTablaPDF / ExportarPromesaVentaPDF', 'Tabla de pagos o promesa de venta del escenario elegido.'),
    ('ExportarResumenComitePDF', 'Hoja interna para comité de crédito.'),
    ('EnviarPorCorreo', 'Genera el paquete PDF y abre un correo de Outlook al cliente con el PDF adjunto.'),
    ('AplicarTasaMinima', 'Pone la tasa (o el margen sobre TIIE) mínima que cumple el rate card.'),
    ('TasaParaRentaDeseada', 'Calcula y aplica la tasa que da la renta que pide el cliente.'),
    ('CopiarEscenario1', 'Copia tasas, enganche, residual, comisión, seguros, etc. del escenario 1 a los demás.'),
    ('ProtegerHojas / DesprotegerHojas', 'Protege (sin contraseña) las fórmulas para evitar borrados accidentales.'),
]
for a, b in macros:
    ws.cell(row=r, column=2, value=a).font = f_bold
    ws.cell(row=r, column=3, value=b).font = f_base
    ws.cell(row=r, column=3).alignment = wrap
    r += 1
r += 1
header_bar(ws, r, 2, 3, 'MODELO MATEMÁTICO')
r += 1
modelo = [
    ('Renta básica', 'R = PMT(i, n, −M, VR, tipo), con i = tasa anual/12, M = valor − enganche + seguro/GPS financiados, VR = residual × valor, tipo = 1 anticipado / 0 vencido.'),
    ('Renta sucesiva', 'Rs = PMT(i2, m, −PF, 0, tipo). Si m = 0 el pago final se cobra como opción de compra única.'),
    ('Amortización', 'Interés_t = (Saldo_t − tipo × Pago_t) × i ; Capital = Pago − Interés. El saldo al final del plazo básico es exactamente VR y termina en 0.'),
    ('Cash margin', 'CM = flujo en firma + Σ flujo_t × (1+fondeo/360)^(−días) × (1+fondeo/12)^(−t). Flujo en firma = −(valor + financiados) + enganche + comisión + renta proporcional + depósito − comisión banco. El depósito se devuelve en el último pago.'),
    ('Tasa mínima', 'CM es lineal en R: CM(R) = A + B·R. Renta mínima = (CM objetivo × M + comisión promotor − A)/B ; tasa mínima = RATE(n, Rmín, −M, VR, tipo) × 12. Sin tanteos ni Buscar objetivo.'),
    ('Tasa', 'Modo tasa fija, o TIIE + margen por escenario. El rate card (Configuracion) fija CM neto mínimo, TIR mínima y residual máximo por plazo.'),
    ('Desglose de renta', 'Por linealidad de PMT: R = PMT(equipo) + PMT(seguro financiado) + PMT(GPS y otros). Así la propuesta muestra cada parte y la suma cuadra exacta.'),
    ('TIR y CAT', 'TIR = XIRR de los flujos del arrendador con fechas reales. CAT informativo = XIRR de los flujos del cliente sin IVA (no oficial).'),
]
for a, b in modelo:
    ws.cell(row=r, column=2, value=a).font = f_bold
    ws.cell(row=r, column=3, value=b).font = f_base
    ws.cell(row=r, column=3).alignment = wrap
    ws.row_dimensions[r].height = 42
    r += 1
r += 1
ws.cell(row=r, column=2, value='Versión 3.0 – octubre 2026').font = f_note

# ==================================================================== IMPRESIÓN
for w_, orient in ((ws_his, 'landscape'), (ws_sen, 'landscape'), (ws_ini, 'portrait'), (ws_cfg, 'landscape')):
    w_.page_setup.orientation = orient
    w_.page_setup.paperSize = w_.PAPERSIZE_LETTER
    w_.page_setup.fitToWidth = 1
    w_.page_setup.fitToHeight = 0
    w_.sheet_properties.pageSetUpPr.fitToPage = True
ws_his.print_title_rows = '4:4'
ws_sen.print_area = 'A1:J%d' % (R2 + len(plazos))

# ==================================================================== PROTECCIÓN
for w in [ws_cot, ws_car, ws_tab, ws_sen, ws_pro, ws_pin, ws_ven, ws_com] + ws_cr:
    w.protection.sheet = True
    w.protection.formatColumns = False
    w.protection.formatRows = False
    w.protection.selectLockedCells = False
    w.protection.selectUnlockedCells = False

wb.active = 1
wb.save(OUT)
import json
json.dump({'P': P, 'names': names, 'FIRST': FIRST, 'LAST': LAST}, open('layout.json', 'w'), indent=1, ensure_ascii=False)
print('ok', OUT)
