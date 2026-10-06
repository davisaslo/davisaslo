import openpyxl, datetime as dt, subprocess
from dateutil.relativedelta import relativedelta
S='/root/.claude/skills/synced/b646bf01-e131-40dd-b005-f47363650fca_4f0b063a-135a-45c1-a143-d0849feced22/xlsx/scripts/recalc.py'
wb=openpyxl.load_workbook('Cotizador_Arrendamiento.xlsx'); c=wb['Cotizador']
def cell(nm,k=None):
    ref=wb.defined_names[nm].attr_text.split('!')[1].replace('$','').split(':')[0]
    col=ord(ref[0])-64; row=int(ref[1:]); return c.cell(row=row,column=col+(k-1 if k else 0))
PRECIO=1084136; TIIE=0.087
cell('inp_Precio').value=PRECIO; cell('inp_PrecioIVA').value='Sí'
cell('inp_ModoTasa').value='TIIE + margen'; cell('inp_TIIE').value=TIIE
cell('inp_GastosInv').value=5000; cell('inp_ComProm').value=0.01
cell('inp_Modalidad').value='Vencido'; cell('inp_TipoActivo').value='Automóvil'
cell('inp_FechaFirma').value=dt.datetime(2026,9,20); cell('inp_FechaPrimera').value=dt.datetime(2026,11,1)
margins=[0.2047,0.1892,0.168,0.173]; plazos=[12,24,36,48]; res=[0.30,0.20,0.15,0.15]
for k in range(1,5):
    cell('esc_Margen',k).value=margins[k-1]; cell('esc_Plazo',k).value=plazos[k-1]; cell('esc_Residual',k).value=res[k-1]
    cell('esc_Enganche',k).value=0.10; cell('esc_Comision',k).value=0.02; cell('esc_Deposito',k).value=0; cell('esc_Sucesivo',k).value=0
cell('esc_Otros',1).value=30000; cell('esc_Seguro',2).value=40000; cell('esc_Otros',3).value=12000; cell('esc_OtrosForma',3).value='Contado'
wb.save('t3.xlsx')
print(subprocess.run(['python3',S,'t3.xlsx','150'],capture_output=True,text=True).stdout.replace('\n',' '))
wv=openpyxl.load_workbook('t3.xlsx',data_only=True)
def pmt(i,n,pv,fv,t):
    if i==0: return -(pv+fv)/n
    return -(pv*(1+i)**n+fv)*i/((1+i*t)*((1+i)**n-1))
def xirr(cfs):
    d0=cfs[0][0]; f=lambda r: sum(v/(1+r)**((d-d0).days/365) for d,v in cfs); lo,hi=-0.99,10
    for _ in range(300):
        mid=(lo+hi)/2
        if f(lo)*f(mid)<=0: hi=mid
        else: lo=mid
    return mid
V=PRECIO/1.16; ff=dt.datetime(2026,9,20); fp=dt.datetime(2026,11,1); tau=0; fa=.205; f=fa/12
fi=fp-relativedelta(months=1); days=(fi-ff).days; DF=(1+fa/360)**(-days); gi=5000
ok=True
for k in range(1,5):
    ws=wv['Corrida %d'%k]; C=lambda r: ws['C%d'%r].value; I=lambda r: ws['I%d'%r].value
    n=plazos[k-1]; ia=TIIE+margins[k-1]; i=ia/12
    eng=.10*V; segf=40000 if k==2 else 0; otrf=30000 if k==1 else 0; otc=12000 if k==3 else 0
    M=V-eng+segf+otrf; VR=res[k-1]*V
    R=pmt(i,n,-M,VR,tau); Req=pmt(i,n,-(V-eng),VR,tau); Rseg=pmt(i,n,-segf,0,tau); Rot=pmt(i,n,-otrf,0,tau)
    rows=[(t,R) for t in range(1,n+1)]+[(n,VR)]
    com=.02*V; rp=R/30*days
    net0=-(V+segf+otrf)+eng+com+rp+gi
    cm=net0+sum(v*DF*(1+f)**(-t) for t,v in rows)
    cmn=cm-0.01*M
    tir=xirr([(ff,net0)]+[(fi+relativedelta(months=t),v) for t,v in rows])
    cat=xirr([(ff,M-com-rp-gi)]+[(fi+relativedelta(months=t),-v) for t,v in rows])
    # tasa minima por biseccion sobre CM neto = 8%
    def cmn_at(ia_):
        ii=ia_/12; RR=pmt(ii,n,-M,VR,tau); rpp=RR/30*days
        n0=-(V+segf+otrf)+eng+com+rpp+gi
        fl=[(t,RR) for t in range(1,n+1)]+[(n,VR)]
        return (n0+sum(v*DF*(1+f)**(-t) for t,v in fl)-0.01*M)/M
    lo,hi=0,2
    for _ in range(200):
        mid=(lo+hi)/2
        if cmn_at(mid)<0.08: lo=mid
        else: hi=mid
    pib=eng+com+rp+gi+otc
    ded=min(1,200*30/R)
    chk=[('R',C(24),R),('Req',C(48),Req),('Rseg',C(49),Rseg),('Rotr',C(50),Rot),('ia',C(8),ia),('M',C(19),M),
         ('CMn',I(25),cmn),('TIR',I(8),tir),('CAT',I(10),cat),('TasaMin',I(15),mid),('PIb',I(17),pib),('ded',C(51),ded),('Saldo',I(16),0)]
    bad=[(n_,a,b) for n_,a,b in chk if abs(a-b)>max(1e-6,abs(b)*1e-7)]
    ok&=not bad
    print('Esc%d n=%d R=%.2f CMneto=%.2f TIR=%.4f%% TasaMin=%.4f%% ded=%.2f%% -> %s'%(k,n,R,cmn,tir*100,mid*100,ded*100,'OK' if not bad else bad))
print('TODO OK' if ok else 'HAY DIFERENCIAS')
