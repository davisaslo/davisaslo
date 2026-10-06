"""Ensambla el .xlsm final: valores en caché + proyecto VBA + pestaña de cinta."""
import os, re, zipfile, datetime as dt, shutil, subprocess, sys, html
import openpyxl
from vbabuild import build_vba_project

SRC = 'Cotizador_FASTPLUS.xlsx'
CALC = 'calcf.xlsx'
OUT = sys.argv[1] if len(sys.argv) > 1 else 'Cotizador_FASTPLUS.xlsm'
RECALC = os.environ.get('RECALC_SCRIPT', 'recalc.py')  # script de recálculo con LibreOffice

shutil.copy(SRC, CALC)
res = subprocess.run(['python3', RECALC, CALC, '180'], capture_output=True, text=True)
print(res.stdout.strip())
assert '"total_errors": 0' in res.stdout

vals = openpyxl.load_workbook(CALC, data_only=True)
src_wb = openpyxl.load_workbook(SRC)
EPOCH = dt.datetime(1899, 12, 30)


def enc(v):
    if isinstance(v, bool):
        return 'b', '1' if v else '0'
    if isinstance(v, (int, float)):
        return None, repr(float(v)) if isinstance(v, float) else str(v)
    if isinstance(v, dt.datetime):
        d = v - EPOCH
        return None, repr(d.days + d.seconds / 86400)
    if isinstance(v, dt.date):
        return None, str((dt.datetime(v.year, v.month, v.day) - EPOCH).days)
    if v is None:
        return 'str', ''
    s = str(v)
    if s.startswith('#'):
        return 'e', s
    return 'str', html.escape(s, quote=False)


zin = zipfile.ZipFile(SRC)
files = {n: zin.read(n) for n in zin.namelist()}
wbxml = files['xl/workbook.xml'].decode('utf-8')
rels = files['xl/_rels/workbook.xml.rels'].decode('utf-8')
sheets = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbxml)
relmap = dict((i, t) for t, i in re.findall(r'Target="([^"]+)" Id="([^"]+)"', rels))
n_inj = 0
for sname, rid in sheets:
    sname = html.unescape(sname)
    path = relmap[rid].lstrip('/')
    xml = files[path].decode('utf-8')
    ws = vals[sname]

    def repl(m):
        global n_inj
        ref, attrs, f = m.group(1), m.group(2), m.group(3)
        t, v = enc(ws[ref].value)
        attrs = re.sub(r'\s+t="[^"]*"', '', attrs)
        n_inj += 1
        tt = ' t="%s"' % t if t else ''
        return '<c r="%s"%s%s><f>%s</f><v>%s</v></c>' % (ref, attrs, tt, f, v)
    xml = re.sub(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>(.*?)</f><v ?/></c>', repl, xml)
    files[path] = xml.encode('utf-8')
print('valores inyectados:', n_inj)

# ---- botones dentro de las hojas (formas con macro asignada, no se imprimen)
MORADO_BTN, ORO_BTN, GRIS_BTN = '4B2A7B', 'C9A227', '6B6B6B'
BUTTONS = {
    'Cotizador': [('Nueva cotización', 'NuevaCotizacion', 'B', 4, 'B', 4, MORADO_BTN),
                  ('Guardar', 'GuardarCotizacion', 'C', 4, 'C', 4, MORADO_BTN),
                  ('Cargar folio', 'CargarCotizacion', 'D', 4, 'D', 4, MORADO_BTN),
                  ('Propuesta completa PDF', 'ExportarPaqueteCliente', 'E', 4, 'F', 4, ORO_BTN),
                  ('Generar e-Propuesta', 'GenerarEPropuesta', 'G', 4, 'H', 4, ORO_BTN),
                  ('Enviar por correo', 'EnviarPorCorreo', 'I', 4, 'J', 4, ORO_BTN),
                  ('Ajustar renta deducible', 'AjustarRentaDeducible', 'K', 4, 'L', 4, MORADO_BTN),
                  ('Acceso gerencia', 'AccesoGerencia', 'M', 4, 'N', 4, GRIS_BTN)],
    'Propuesta': [('Generar e-Propuesta', 'GenerarEPropuesta', 'I', 2, 'J', 3, ORO_BTN),
                  ('Propuesta en PDF', 'ExportarPropuestaPDF', 'I', 5, 'J', 6, MORADO_BTN),
                  ('Regresar al Cotizador', 'IrCotizador', 'I', 8, 'J', 9, GRIS_BTN)],
    'Pago Inicial': [('Generar e-Pago Inicial', 'GenerarEPagoInicial', 'J', 2, 'K', 3, ORO_BTN),
                     ('Propuesta completa PDF', 'ExportarPaqueteCliente', 'J', 5, 'K', 6, MORADO_BTN),
                     ('Regresar al Cotizador', 'IrCotizador', 'J', 8, 'K', 9, GRIS_BTN)],
    'Venta': [('Generar e-Venta', 'GenerarEVenta', 'I', 2, 'J', 3, ORO_BTN),
              ('Carta de venta PDF', 'ExportarPromesaVentaPDF', 'I', 5, 'J', 6, MORADO_BTN),
              ('Regresar al Cotizador', 'IrCotizador', 'I', 8, 'J', 9, GRIS_BTN)],
    'Factores': [('Ocultar hoja', 'OcultarGerencia', 'B', 16, 'B', 17, GRIS_BTN),
                 ('Ajustar renta deducible', 'AjustarRentaDeducible', 'C', 16, 'D', 17, MORADO_BTN),
                 ('Margen mínimo (rate card)', 'AplicarTasaMinima', 'E', 16, 'F', 17, MORADO_BTN),
                 ('Expediente interno PDF', 'ExportarCreditoPDF', 'G', 16, 'H', 17, ORO_BTN),
                 ('Regresar al Cotizador', 'IrCotizador', 'I', 16, 'J', 17, GRIS_BTN)],
    'Bonos': [('Mostrar bonos', 'MostrarBonos', 'G', 2, 'H', 3, MORADO_BTN),
              ('Ocultar bonos', 'OcultarBonos', 'G', 5, 'H', 6, MORADO_BTN),
              ('Ocultar hoja', 'OcultarGerencia', 'G', 8, 'H', 9, GRIS_BTN)],
    'Riesgo': [('Hoja de riesgo PDF', 'ExportarResumenComitePDF', 'H', 2, 'I', 3, ORO_BTN),
               ('Ocultar hoja', 'OcultarGerencia', 'H', 5, 'I', 6, GRIS_BTN),
               ('Regresar al Cotizador', 'IrCotizador', 'H', 8, 'I', 9, GRIS_BTN)],
    'Tabla de Pagos': [('Tabla de pagos PDF', 'ExportarTablaPDF', 'K', 2, 'L', 3, MORADO_BTN),
                       ('Regresar al Cotizador', 'IrCotizador', 'K', 5, 'L', 6, GRIS_BTN)],
    'Historial': [('Cargar folio seleccionado', 'CargarCotizacion', 'J', 1, 'K', 2, MORADO_BTN),
                  ('Regresar al Cotizador', 'IrCotizador', 'L', 1, 'M', 2, GRIS_BTN)],
    'Catalogos': [('Ocultar hojas restringidas', 'OcultarRestringidas', 'H', 1, 'I', 2, GRIS_BTN),
                  ('Proteger libro y hojas', 'ProtegerHojas', 'J', 1, 'K', 2, MORADO_BTN)],
}
XDR = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def col_idx(letter):
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return n - 1


def col_emu(ws, letter):
    d = ws.column_dimensions.get(letter)
    w = d.width if (d is not None and d.customWidth and d.width) else 8.43
    return int((w * 7 + 5) * 9525)


def row_emu(ws, r):
    d = ws.row_dimensions.get(r)
    h = d.height if (d is not None and d.height) else 15
    return int(h * 12700)


def boton(ws, sid, text, macro, c1, r1, c2, r2, color, pfx):
    pad = 38100
    t = html.escape(text, quote=False)
    a = '' if pfx == '' else 'xdr:'
    return ('<{a}twoCellAnchor editAs="absolute">'
            '<{a}from><{a}col>{c1}</{a}col><{a}colOff>{p}</{a}colOff><{a}row>{r1}</{a}row><{a}rowOff>{p}</{a}rowOff></{a}from>'
            '<{a}to><{a}col>{c2}</{a}col><{a}colOff>{c2o}</{a}colOff><{a}row>{r2}</{a}row><{a}rowOff>{r2o}</{a}rowOff></{a}to>'
            '<{a}sp macro="[0]!{m}" textlink=""><{a}nvSpPr><{a}cNvPr id="{sid}" name="btn_{m}_{sid}"/><{a}cNvSpPr/></{a}nvSpPr>'
            '<{a}spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
            '<a:solidFill><a:srgbClr val="{col}"/></a:solidFill><a:ln w="12700"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln>'
            '<a:effectLst><a:outerShdw blurRad="38100" dist="19050" dir="5400000" algn="t" rotWithShape="0"><a:prstClr val="black"><a:alpha val="35000"/></a:prstClr></a:outerShdw></a:effectLst></{a}spPr>'
            '<{a}txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" wrap="square" lIns="36000" tIns="18000" rIns="36000" bIns="18000" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
            '<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="es-MX" sz="900" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
            '<a:latin typeface="Arial"/><a:cs typeface="Arial"/></a:rPr><a:t>{t}</a:t></a:r></a:p></{a}txBody></{a}sp>'
            '<{a}clientData fPrintsWithSheet="0"/></{a}twoCellAnchor>').format(
        a=a, c1=col_idx(c1), r1=r1 - 1, c2=col_idx(c2), r2=r2 - 1, p=pad,
        c2o=col_emu(ws, c2) - pad, r2o=row_emu(ws, r2) - pad, m=macro, sid=sid, col=color, t=t)


ct = files['[Content_Types].xml'].decode('utf-8')
nuevo = 0
for sname, rid in sheets:
    sname = html.unescape(sname)
    if sname not in BUTTONS:
        continue
    spath = relmap[rid].lstrip('/')
    rpath = spath.replace('worksheets/', 'worksheets/_rels/') + '.rels'
    sxml = files[spath].decode('utf-8')
    rels_s = files[rpath].decode('utf-8') if rpath in files else '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
    m = re.search(r'Type="[^"]*/drawing" Target="([^"]+)"', rels_s)
    if m:
        dpath = m.group(1).lstrip('/')
        dxml = files[dpath].decode('utf-8')
    else:
        nuevo += 1
        dpath = 'xl/drawings/drawingBtn%d.xml' % nuevo
        dxml = '<xdr:wsDr xmlns:xdr="%s" xmlns:a="%s" xmlns:r="%s"></xdr:wsDr>' % (XDR, A_NS, R_NS)
        rels_s = rels_s.replace('</Relationships>', '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="/%s" Id="rIdBtnDr" /></Relationships>' % dpath)
        tag = '<drawing xmlns:r="%s" r:id="rIdBtnDr" />' % R_NS
        if '<legacyDrawing' in sxml:
            sxml = sxml.replace('<legacyDrawing', tag + '<legacyDrawing', 1)
        elif '<tableParts' in sxml:
            sxml = sxml.replace('<tableParts', tag + '<tableParts', 1)
        else:
            sxml = sxml.replace('</worksheet>', tag + '</worksheet>', 1)
        ct = ct.replace('</Types>', '<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml" /></Types>' % dpath)
        files[rpath] = rels_s.encode('utf-8')
        files[spath] = sxml.encode('utf-8')
    pfx = 'xdr' if dxml.startswith('<xdr:') else ''
    close = '</xdr:wsDr>' if pfx else '</wsDr>'
    ws_src = src_wb[sname]
    shapes = ''.join(boton(ws_src, 900 + i, *b, pfx) for i, b in enumerate(BUTTONS[sname]))
    dxml = dxml.replace(close, shapes + close)
    files[dpath] = dxml.encode('utf-8')
files['[Content_Types].xml'] = ct.encode('utf-8')
print('botones agregados en', len(BUTTONS), 'hojas')

# ---- libro con macros
wbxml = wbxml.replace('<workbookPr />', '<workbookPr codeName="ThisWorkbook" />')
assert 'codeName="ThisWorkbook"' in wbxml
files['xl/workbook.xml'] = wbxml.encode('utf-8')
ct = files['[Content_Types].xml'].decode('utf-8')
ct = ct.replace('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
ct = ct.replace('<Default Extension="xml"', '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject" /><Default Extension="xml"', 1)
assert 'macroEnabled' in ct and 'vbaProject' in ct
files['[Content_Types].xml'] = ct.encode('utf-8')
rels = rels.replace('</Relationships>', '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="/xl/vbaProject.bin" Id="rIdVBA1" /></Relationships>')
files['xl/_rels/workbook.xml.rels'] = rels.encode('utf-8')

# ---- módulos VBA
mods = [dict(name='ThisWorkbook', document=True, base='0{00020819-0000-0000-C000-000000000046}',
             code=open('../macros/ThisWorkbook.cls', encoding='utf-8').read())]
for ws in src_wb.worksheets:
    mods.append(dict(name=ws.sheet_properties.codeName, document=True,
                     base='0{00020820-0000-0000-C000-000000000046}', code=''))
mods.append(dict(name='modCotizador', document=False, code=open('../macros/modCotizador.bas', encoding='utf-8').read()))
import os
mods.append(dict(name='modRibbon', document=False, code=open('../macros/modRibbon.bas', encoding='utf-8').read()))
files['xl/vbaProject.bin'] = build_vba_project('orig_vbaProject.bin', mods, 'vbaProject.bin')

# ---- cinta personalizada
def btn(id_, label, action, img=None, size='large', tip=''):
    a = ' imageMso="%s"' % img if img else ''
    s = ' size="%s"' % size if size else ''
    return '<button id="%s" label="%s"%s%s onAction="%s" screentip="%s"/>' % (id_, html.escape(label), s, a, action, html.escape(tip or label))

ui = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
      '<customUI xmlns="http://schemas.microsoft.com/office/2006/01/customui"><ribbon><tabs>'
      '<tab id="tabFastplus" label="FASTPLUS" insertAfterMso="TabHome">'
      '<group id="gCot" label="Cotización">'
      + btn('bNueva', 'Nueva cotización', 'rbNueva', 'FileNew', tip='Limpia la captura y asigna folio (Ctrl+Shift+N)')
      + btn('bGuardar', 'Guardar', 'rbGuardar', 'FileSave', tip='Guarda la cotización en Historial (Ctrl+Shift+G)')
      + btn('bCargar', 'Cargar folio', 'rbCargar', 'FileOpen')
      + '</group><group id="gEnt" label="Cliente">'
      + btn('bPaquete', 'Propuesta completa PDF', 'rbPaquete', 'FileSaveAsPdfOrXps', tip='Propuesta + Pago inicial + Tabla de pagos (Ctrl+Shift+P)')
      + btn('bCorreo', 'Enviar por correo', 'rbCorreo', 'FileSendAsAttachment')
      + btn('bProp', 'Propuesta', 'rbPropuesta', None, 'normal')
      + btn('bTabla', 'Tabla de pagos', 'rbTablaPDF', None, 'normal')
      + btn('bVenta', 'Carta de venta', 'rbVenta', None, 'normal')
      + '</group><group id="gEdoc" label="e-Documentos (Excel)">'
      + btn('bEProp', 'e-Propuesta', 'rbEPropuesta', 'FileSaveAs', tip='Copia la propuesta a un libro nuevo solo con valores')
      + btn('bEPago', 'e-Pago Inicial', 'rbEPago', None, 'normal')
      + btn('bEVenta', 'e-Venta', 'rbEVenta', None, 'normal')
      + '</group><group id="gCre" label="Crédito">'
      + btn('bCredito', 'Expediente interno PDF', 'rbCredito', 'FileSaveAsPdfOrXps', tip='Factores + Bonos + Riesgo')
      + btn('bRiesgo', 'Hoja de riesgo', 'rbRiesgo', None, 'normal')
      + '</group><group id="gFin" label="Ajustes">'
      + btn('bTasaMin', 'Margen mínimo (rate card)', 'rbTasaMin', 'AutoSum', tip='Pone el margen sobre TIIE que cumple el rate card')
      + btn('bTasaRenta', 'Tasa para renta deseada', 'rbTasaRenta', 'PercentStyle')
      + btn('bRentaDed', 'Renta deducible', 'rbRentaDed', 'AutoSum', tip='Calcula la renta extraordinaria para que la renta sea 100% deducible')
      + '</group><group id="gIr" label="Ir a">'
      + btn('bIrCot', 'Cotizador', 'rbIrCotizador', None, 'normal')
      + btn('bIrFac', 'Factores', 'rbIrFactores', None, 'normal')
      + btn('bIrProp', 'Propuesta', 'rbIrPropuesta', None, 'normal')
      + btn('bIrPI', 'Pago inicial', 'rbIrPagoInicial', None, 'normal')
      + btn('bIrVen', 'Venta', 'rbIrVenta', None, 'normal')
      + btn('bIrBon', 'Bonos', 'rbIrBonos', None, 'normal')
      + btn('bIrRie', 'Riesgo', 'rbIrRiesgo', None, 'normal')
      + btn('bIrTab', 'Tabla de pagos', 'rbIrTabla', None, 'normal')
      + btn('bIrHis', 'Historial', 'rbIrHistorial', None, 'normal')
      + btn('bIrCor', 'Corrida', 'rbIrCorrida', None, 'normal')
      + btn('bIrCat', 'Catálogos', 'rbIrCatalogos', None, 'normal')
      + '</group><group id="gAdm" label="Acceso">'
      + btn('bGer', 'Acceso gerencia', 'rbGerencia', 'Lock')
      + btn('bOculGer', 'Ocultar gerencia', 'rbOcultarGer', None, 'normal')
      + btn('bAdmin', 'Administrador', 'rbAdmin', None, 'normal')
      + btn('bProt', 'Proteger hojas', 'rbProteger', None, 'normal')
      + btn('bDesp', 'Desproteger hojas', 'rbDesproteger', None, 'normal')
      + '</group></tab></tabs></ribbon></customUI>')
files['customUI/customUI.xml'] = ui.encode('utf-8')
rr = files['_rels/.rels'].decode('utf-8')
rr = rr.replace('</Relationships>', '<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/ui/extensibility" Target="customUI/customUI.xml" Id="rIdUI1" /></Relationships>')
files['_rels/.rels'] = rr.encode('utf-8')

order = ['[Content_Types].xml', '_rels/.rels'] + [n for n in files if n not in ('[Content_Types].xml', '_rels/.rels')]
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for n in order:
        z.writestr(n, files[n])
print('escrito', OUT)
