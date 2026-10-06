import datetime as dt
import copy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
import openpyxl.worksheet.page as page

OUT = 'Cotizador_FASTPLUS.xlsx'
import os as _os
CLAVE_ADMIN = _os.environ.get('CLAVE_ADMIN', 'FP-Admin2026')        # protege libro y hojas; abre Catálogos
CLAVE_GERENCIA = _os.environ.get('CLAVE_GERENCIA', 'FP-Gerencia2026')  # abre Factores, Bonos y Riesgo
import os
# Promotores: si existe la variable PROMOTORES (ruta al Excel "Relación de colaboradores"), se cargan
# los datos reales; si no, se usan datos de ejemplo (versión pública sin datos personales).
PROMOTORES = []
if os.environ.get('PROMOTORES'):
    import openpyxl as _ox
    _ws = _ox.load_workbook(os.environ['PROMOTORES'], data_only=True).worksheets[0]
    for _row in _ws.iter_rows(min_row=3, values_only=True):
        if _row[1]:
            _nom = ' '.join(str(x).strip() for x in _row[1:4] if x)
            _p = ' '.join(w if w in ('de', 'del', 'la', 'y') else w.capitalize() for w in str(_row[6] or '').lower().split())
            PROMOTORES.append((str(_row[0]), _nom.upper(), _p, _row[5] or '', _row[4] or ''))
    PROM_DEFAULT = next((p[1] for p in PROMOTORES if 'Fastplus' in p[2]), PROMOTORES[0][1])
else:
    PROMOTORES = [('001', 'NOMBRE DEL PROMOTOR', 'Promotor Fastplus', 'CDMX', 'promotor@empresa.com.mx'),
                  ('002', 'NOMBRE DEL GERENTE', 'Gerente De Promoción', 'CDMX', 'gerente@empresa.com.mx')]
    PROM_DEFAULT = 'NOMBRE DEL PROMOTOR'
LOGO = 'logo.jpeg'

# ------------------------------------------------------------------ estilo
MORADO = '4B2A7B'
MORADO_CL = 'E9E2F3'
GRIS = 'F2F2F2'
AMARILLO = 'FFF8D6'
FONT = 'Arial'

f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_input = Font(name=FONT, size=10, color='0000FF')
f_link = Font(name=FONT, size=10, color='008000')
f_hdr = Font(name=FONT, size=11, bold=True, color='FFFFFF')
f_title = Font(name=FONT, size=18, bold=True, color=MORADO)
f_sub = Font(name=FONT, size=10, italic=True, color='666666')
f_note = Font(name=FONT, size=9, italic=True, color='666666')
fill_hdr = PatternFill('solid', fgColor=MORADO)
fill_sub = PatternFill('solid', fgColor=MORADO_CL)
fill_in = PatternFill('solid', fgColor=AMARILLO)
fill_gris = PatternFill('solid', fgColor=GRIS)
thin = Side(style='thin', color='BFBFBF')
med = Side(style='medium', color=MORADO)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
topline = Border(top=med)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
right = Alignment(horizontal='right', vertical='center')
left_c = Alignment(horizontal='left', vertical='center', wrap_text=True)

MON = '$#,##0.00;[Red]-$#,##0.00;"-"'
MON0 = '$#,##0;[Red]-$#,##0;"-"'
PCT = '0.00%;[Red]-0.00%;"-"'
PCT4 = '0.0000%'
FECHA = 'dd/mm/yyyy'
NUM = '#,##0.00'

wb = Workbook()
wb.remove(wb.active)
names = {}


def name(nm, ref):
    names[nm] = ref
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def absref(sheet, rng):
    parts = rng.split(':')
    def a(c):
        col = ''.join(ch for ch in c if ch.isalpha())
        row = ''.join(ch for ch in c if ch.isdigit())
        return '$%s$%s' % (col, row)
    q = "'%s'" % sheet if ' ' in sheet else sheet
    return q + '!' + ':'.join(a(p) for p in parts)


def style_range(ws, rng, font=None, fill=None, border=None, align=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = border
            if align: c.alignment = align
            if fmt: c.number_format = fmt


def header_bar(ws, row, c1, c2, text):
    ws.cell(row=row, column=c1, value=text)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_hdr
        cell.font = f_hdr
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[row].height = 20


def inp(cell, value, fmt=None, note=None):
    cell.value = value
    cell.font = f_input
    cell.fill = fill_in
    cell.border = box
    cell.protection = Protection(locked=False)
    if fmt:
        cell.number_format = fmt
    if note:
        cell.comment = Comment(note, 'Cotizador')


def calc(cell, formula, fmt=None, bold=False):
    cell.value = formula
    cell.font = f_bold if bold else f_base
    cell.border = box
    if fmt:
        cell.number_format = fmt


def logo(ws, anchor, h=60):
    img = XLImage(LOGO)
    ratio = img.width / img.height
    img.height = h
    img.width = int(h * ratio)
    ws.add_image(img, anchor)


def base_sheet(title, code, tab=None):
    ws = wb.create_sheet(title)
    ws.sheet_properties.codeName = code
    if tab:
        ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    return ws


def dv_list(ws, formula, cells, prompt=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.error = 'Seleccione un valor de la lista.'
    dv.errorTitle = 'Valor no válido'
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cells)


def dv_num(ws, cells, lo, hi, kind='decimal', msg=''):
    dv = DataValidation(type=kind, operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.error = msg or ('Capture un valor entre %s y %s.' % (lo, hi))
    dv.errorTitle = 'Valor fuera de rango'
    ws.add_data_validation(dv)
    dv.add(cells)


# ==================================================================== HOJAS (formato tipo ABC Leasing, marca FASTPLUS)
AMAR = PatternFill('solid', fgColor='FFFBD6')      # fondo de captura (amarillo pálido)
BLANCO = PatternFill('solid', fgColor='FFFFFF')
f_bar = Font(name=FONT, size=10, bold=True, italic=True, color='FFFFFF')
f_lab = Font(name=FONT, size=10, color='4B2A7B')
f_inbox = Font(name=FONT, size=10, color='0000FF')
f_outbox = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_big = Font(name=FONT, size=12, bold=True, color=MORADO)
fill_out = PatternFill('solid', fgColor=MORADO)
box_in = Border(left=Side(style='thin', color='4B2A7B'), right=Side(style='thin', color='4B2A7B'),
                top=Side(style='thin', color='4B2A7B'), bottom=Side(style='thin', color='4B2A7B'))

ws_cot = base_sheet('Cotizador', 'Hoja1', 'FFC000')
ws_fac = base_sheet('Factores', 'Hoja2', 'FFC000')
ws_pro = base_sheet('Propuesta', 'Hoja3', '70AD47')
ws_pin = base_sheet('Pago Inicial', 'Hoja4', '70AD47')
ws_ven = base_sheet('Venta', 'Hoja5', '70AD47')
ws_bon = base_sheet('Bonos', 'Hoja6', '5B9BD5')
ws_rie = base_sheet('Riesgo', 'Hoja7', 'C00000')
ws_tab = base_sheet('Tabla de Pagos', 'Hoja8', '70AD47')
ws_his = base_sheet('Historial', 'Hoja9', '5B9BD5')
ws_cr = [base_sheet('Corrida %d' % k, 'Hoja%d' % (9 + k), '808080') for k in range(1, 5)]
ws_cfg = base_sheet('Catalogos', 'Hoja14', '808080')
for w in ws_cr:
    w.sheet_state = 'hidden'


def pinta(ws, rng, fill):
    for row in ws[rng]:
        for c in row:
            c.fill = fill


def barra(ws, row, c1, c2, text):
    ws.cell(row=row, column=c1, value=text)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_hdr
        cell.font = f_bar
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.row_dimensions[row].height = 17


def etiqueta(ws, coord, text, align='right'):
    c = ws[coord]
    c.value = text
    c.font = f_lab
    c.alignment = Alignment(horizontal=align, vertical='center')


def caja(ws, rng, value, nm=None, fmt=None, lst=None, note=None, salida=False):
    first = rng.split(':')[0]
    c = ws[first]
    c.value = value
    if ':' in rng:
        ws.merge_cells(rng)
    rr_ = rng if ':' in rng else '%s:%s' % (rng, rng)
    for row in ws[rr_]:
        for x in row:
            x.border = box_in
            x.fill = fill_out if salida else BLANCO
    c.font = f_outbox if salida else f_inbox
    c.alignment = Alignment(horizontal='center', vertical='center')
    if not salida:
        c.protection = Protection(locked=False)
    if fmt:
        c.number_format = fmt
    if nm:
        name(nm, absref(ws.title, first))
    if lst:
        dv_list(ws, lst, first)
    if note:
        c.comment = Comment(note, 'Cotizador')
    return c


# ==================================================================== CATÁLOGOS
ws = ws_cfg
for col, w in zip('ABCDEF', [2, 50, 22, 14, 14, 14]):
    ws.column_dimensions[col].width = w
ws['B1'] = 'CATÁLOGOS Y PARÁMETROS – COTIZADOR FASTPLUS'
ws['B1'].font = f_title
ws['B2'] = 'Solo el administrador debe modificar esta hoja. Celdas amarillas editables.'
ws['B2'].font = f_sub
header_bar(ws, 3, 2, 6, 'DATOS DE LA EMPRESA Y PARÁMETROS')
cfg = [
    ('Razón social del arrendador', 'SOFOPLUS, S.A.P.I. DE C.V., E.R.', 'par_Arrendador', None),
    ('Nombre comercial del producto', 'FASTPLUS', 'par_Comercial', None),
    ('Domicilio corporativo (línea 1)', 'Paseo de los Tamarindos No. 90 Piso 24 Torre 1', 'par_Dir1', None),
    ('Domicilio corporativo (línea 2)', 'Col. Bosques de las Lomas', 'par_Dir2', None),
    ('Domicilio corporativo (línea 3)', 'CDMX, C.P. 05120', 'par_Dir3', None),
    ('Ciudad para la fecha de los documentos', 'Ciudad de México', 'par_Ciudad', None),
    ('Sitio web / teléfono general', '', 'par_Contacto', None),
    ('Tasa de IVA rentas', 0.16, 'par_IVA', PCT),
    ('Tasa de fondeo anual (valor por defecto)', 0.205, 'par_Fondeo', PCT),
    ('TIIE 28 días (valor por defecto)', 0.086, 'par_TIIE', PCT4),
    ('Vigencia de la propuesta (días naturales)', 15, 'par_Vigencia', '0'),
    ('Base de días para costo de fondeo diario', 360, 'par_BaseDias', '0'),
    ('IVA en arrendamiento financiero', 'Sobre renta', 'par_BaseIVAFin', None),
    ('Prefijo del folio', 'FP-', 'par_Prefijo', None),
    ('Último consecutivo utilizado (lo actualiza la macro)', 0, 'par_Consecutivo', '0'),
    ('Carpeta para PDF (vacío = carpeta del archivo)', '', 'par_CarpetaPDF', None),
    ('Correo con copia (CC) al enviar cotizaciones', '', 'par_CorreoCC', None),
    ('Validación de cuenta / gastos de investigación (sin IVA)', 0, 'par_GastosInv', MON),
    ('Límite deducible renta automóvil ($ diarios, LISR art. 28-XIII)', 200, 'par_LimAuto', MON),
    ('Límite deducible renta auto eléctrico / híbrido ($ diarios)', 285, 'par_LimAutoEV', MON),
    ('Transferencia: banco', '', 'par_Banco', None),
    ('Transferencia: CLABE', '', 'par_CLABE', None),
    ('Depósito en sucursal: banco', '', 'par_BancoSuc', None),
    ('Depósito en sucursal: cuenta', '', 'par_Cuenta', None),
    ('Depósito en sucursal: convenio', '', 'par_Convenio', None),
    ('Beneficiario', 'SOFOPLUS, S.A.P.I. DE C.V., E.R.', 'par_Beneficiario', None),
]
r = 4
for label, val, nm, fmt in cfg:
    ws.cell(row=r, column=2, value=label).font = f_base
    inp(ws.cell(row=r, column=3), val, fmt)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    name(nm, absref(ws.title, 'C%d' % r))
    if nm == 'par_BaseIVAFin':
        dv_list(ws, '=lst_BaseIVA', 'C%d' % r)
    r += 1
ws['C12'].comment = Comment('Dato del archivo original de SOFOPLUS (20.5%). Actualice con su costo de fondeo.', 'Cotizador')
ws['C13'].comment = Comment('Actualice con la TIIE 28 días publicada por Banxico.', 'Cotizador')
ws['C22'].comment = Comment('Ley del ISR art. 28 fr. XIII: renta de automóviles deducible hasta $200 diarios ($285 eléctricos/híbridos). Verifique con su área fiscal.', 'Cotizador')

# ---- valores fijos del modelo (no editar)
r += 1
header_bar(ws, r, 2, 6, 'VALORES INTERNOS DEL MODELO (no editar)')
r += 1
fijos = [
    ('El precio se captura con IVA', 'Sí', 'inp_PrecioIVA'),
    ('Modo de tasa', 'TIIE + margen', 'inp_ModoTasa'),
    ('Fecha de firma / pago al proveedor', '=inp_Fecha', 'inp_FechaFirma'),
    ('Fecha de la primera renta', '=IF(inp_Modalidad="Anticipado",inp_Fecha+inp_DiasRP,EDATE(inp_Fecha+inp_DiasRP,1))', 'inp_FechaPrimera'),
    ('Escenario del plazo solicitado', '=IFERROR(MATCH(inp_PlazoSol,esc_Plazo,0),1)', 'inp_EscTabla'),
]
for lab, f, nm in fijos:
    ws.cell(row=r, column=2, value=lab).font = f_base
    c = ws.cell(row=r, column=3, value=f)
    c.font = f_base
    c.border = box
    if nm in ('inp_FechaFirma', 'inp_FechaPrimera'):
        c.number_format = FECHA
    name(nm, absref(ws.title, 'C%d' % r))
    r += 1

# ---- listas
header_bar(ws, 3, 8, 22, 'LISTAS DESPLEGABLES')
listas = [
    ('H', 'Tipo de arrendamiento', ['Arrendamiento Puro', 'Arrendamiento Financiero'], 'lst_Tipo'),
    ('I', 'Moneda', ['Moneda Nacional', 'Dólares Americanos'], 'lst_Moneda'),
    ('J', 'Modalidad', ['Vencido', 'Anticipado'], 'lst_Modalidad'),
    ('K', 'Sí / No', ['Sí', 'No'], 'lst_SiNo'),
    ('L', 'Seguro', ['No financiado', 'Financiado'], 'lst_Seguro'),
    ('M', 'Base IVA financiero', ['Sobre renta', 'Sobre intereses'], 'lst_BaseIVA'),
    ('N', 'Estatus', ['Enviada', 'En seguimiento', 'Aceptada', 'Rechazada', 'Vencida'], 'lst_Estatus'),
    ('O', 'Producto', ['FASTPLUS', 'Arrendamiento SOFOPLUS'], 'lst_Producto'),
    ('P', 'Tipo de activo', ['Equipo médico estándar', 'Maquinaria y equipo', 'Equipo de cómputo / tecnología',
                             'Automóvil', 'Automóvil eléctrico / híbrido', 'Vehículo de carga / utilitario'], 'lst_TipoActivo'),
    ('Q', 'Uso del anticipo', ['Como renta extraordinaria', 'Como enganche'], 'lst_UsoAnticipo'),
    ('R', 'Estado del bien', ['Nuevo', 'Seminuevo'], 'lst_Estado'),
    ('S', 'Originador', ['Promotor FASTPLUS', 'Referenciador', 'Agencia / distribuidor', 'Empleado SOFOPLUS'], 'lst_Originador'),
    ('T', 'Aseguradora por parte de', ['Cliente', 'FASTPLUS'], 'lst_AsegPor'),
    ('U', 'Plazos', [12, 24, 36, 48], 'lst_Plazos'),
]
for col, title, vals, nm in listas:
    c = ws['%s4' % col]
    c.value = title
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, v in enumerate(vals):
        cell = ws['%s%d' % (col, 5 + i)]
        cell.value = v
        cell.font = f_base
        cell.border = box
    name(nm, absref(ws.title, '%s5:%s%d' % (col, col, 4 + len(vals))))
    ws.column_dimensions[col].width = 18
ws.row_dimensions[4].height = 30

# ---- oficinas
OF0 = 13
header_bar(ws, OF0, 8, 12, 'OFICINAS POR REGIÓN (encabezado de la propuesta)')
for j, h in enumerate(['Región', 'Domicilio (línea 1)', 'Domicilio (línea 2)', 'Domicilio (línea 3)', 'Teléfono']):
    c = ws.cell(row=OF0 + 1, column=8 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
oficinas = [('CDMX', 'Paseo de los Tamarindos No. 90 Piso 24 Torre 1', 'Col. Bosques de las Lomas', 'CDMX, C.P. 05120', ''),
            ('GUADALAJARA', '', '', '', ''), ('CANCÚN', '', '', '', ''), ('QUERÉTARO', '', '', '', ''),
            ('', '', '', '', ''), ('', '', '', '', '')]
for i, row in enumerate(oficinas):
    for j, v in enumerate(row):
        inp(ws.cell(row=OF0 + 2 + i, column=8 + j), v if v else None)
name('ofi_Tabla', absref(ws.title, 'H%d:L%d' % (OF0 + 2, OF0 + 1 + len(oficinas))))
name('ofi_Region', absref(ws.title, 'H%d:H%d' % (OF0 + 2, OF0 + 1 + len(oficinas))))
ws.cell(row=OF0 + 2 + len(oficinas), column=8, value='Si una región no tiene domicilio se usa el corporativo.').font = f_note

# ---- promotores
PR0 = OF0 + 10
NPROM = 60
header_bar(ws, PR0, 8, 12, 'CATÁLOGO DE PROMOTORES')
for j, h in enumerate(['Clave', 'Nombre completo', 'Puesto', 'Región (oficina)', 'Correo electrónico']):
    c = ws.cell(row=PR0 + 1, column=8 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
for i in range(NPROM):
    row = PROMOTORES[i] if i < len(PROMOTORES) else (None,) * 5
    for j, v in enumerate(row):
        inp(ws.cell(row=PR0 + 2 + i, column=8 + j), v)
name('prom_Tabla', absref(ws.title, 'H%d:L%d' % (PR0 + 2, PR0 + 1 + NPROM)))
name('lst_Promotores', absref(ws.title, 'I%d:I%d' % (PR0 + 2, PR0 + 1 + NPROM)))
ws.column_dimensions['I'].width = 34
ws.column_dimensions['J'].width = 28
ws.column_dimensions['K'].width = 18
ws.column_dimensions['L'].width = 28

# ---- promotor seleccionado
r += 1
SEL0 = r
header_bar(ws, SEL0, 2, 6, 'PROMOTOR SELECCIONADO (automático)')
sel = [
    ('Nombre (formato carta)', '=PROPER(inp_Promotor)', 'sel_PromNombre'),
    ('Puesto', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),3)&"","")', 'sel_PromPuesto'),
    ('Región', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),4)&"","")', 'sel_PromRegion'),
    ('Correo', '=IFERROR(INDEX(prom_Tabla,MATCH(inp_Promotor,lst_Promotores,0),5)&"","")', 'sel_PromCorreo'),
    ('Domicilio línea 1', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir1,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""),par_Dir1)', 'sel_Dir1'),
    ('Domicilio línea 2', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir2,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),3)&""),par_Dir2)', 'sel_Dir2'),
    ('Domicilio línea 3', '=IFERROR(IF(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),2)&""="",par_Dir3,INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),4)&""),par_Dir3)', 'sel_Dir3'),
    ('Teléfono', '=IFERROR(INDEX(ofi_Tabla,MATCH(sel_PromRegion,ofi_Region,0),5)&"","")', 'sel_Tel'),
]
for i, (lab, f, nm) in enumerate(sel):
    rr = SEL0 + 1 + i
    ws.cell(row=rr, column=2, value=lab).font = f_base
    calc(ws.cell(row=rr, column=3), f)
    ws.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
    name(nm, absref(ws.title, 'C%d' % rr))
ws.freeze_panes = 'A4'

# ==================================================================== COTIZADOR (captura, formato ABC)
ws = ws_cot
ws.column_dimensions['A'].width = 2
for col in 'BCDEFGHIJKLMN':
    ws.column_dimensions[col].width = 13.5
ws.column_dimensions['B'].width = 24
ws.column_dimensions['O'].width = 2
pinta(ws, 'B1:N66', AMAR)
logo(ws, 'B1', 48)
ws['D2'] = 'Cotizador Master FASTPLUS versión 1.0'
ws['D2'].font = Font(name=FONT, size=12, bold=True, color=MORADO)
ws.merge_cells('D2:G2')
etiqueta(ws, 'H2', 'Promotor:')
caja(ws, 'I2:K2', PROM_DEFAULT, 'inp_Promotor', lst='=lst_Promotores')
etiqueta(ws, 'L2', 'Oficina:')
caja(ws, 'M2:N2', '=sel_PromRegion', salida=True)
etiqueta(ws, 'H3', 'Folio:')
caja(ws, 'I3', 'FP-0001', 'inp_Folio')
etiqueta(ws, 'L3', 'Tipo:')
caja(ws, 'M3:N3', '=sel_PromPuesto', salida=True)
ws.row_dimensions[2].height = 20
ws.row_dimensions[4].height = 30

barra(ws, 5, 2, 14, 'Introducir datos generales')
etiqueta(ws, 'B7', 'Cliente:')
caja(ws, 'C7:G7', 'CLIENTE DE EJEMPLO, S.A. DE C.V.', 'inp_Cliente')
etiqueta(ws, 'H7', 'Atención:')
caja(ws, 'I7:K7', '', 'inp_Contacto')
etiqueta(ws, 'L7', 'RFC:')
caja(ws, 'M7:N7', '', 'inp_RFC')
etiqueta(ws, 'B9', 'Fecha:')
caja(ws, 'C9', dt.datetime(2026, 10, 6), 'inp_Fecha', FECHA)
etiqueta(ws, 'E9', 'Originador de operación:')
ws.merge_cells('E9:F9')
caja(ws, 'G9:H9', 'Promotor FASTPLUS', 'inp_Originador', lst='=lst_Originador')
etiqueta(ws, 'I9', 'Comisión para originador:')
ws.merge_cells('I9:K9')
caja(ws, 'L9', 0, 'inp_ComProm', PCT, note='% del monto financiado. Se descuenta del margen de caja (margen neto).')
etiqueta(ws, 'B11', 'Tipo de activo:')
caja(ws, 'C11:E11', 'Automóvil eléctrico / híbrido', 'inp_TipoActivo', lst='=lst_TipoActivo')
etiqueta(ws, 'F11', 'Descripción de activo:')
ws.merge_cells('F11:G11')
caja(ws, 'H11:N11', 'Auto BYD M9 2026', 'inp_Equipo')
etiqueta(ws, 'B13', 'Estado del bien:')
caja(ws, 'C13', 'Nuevo', 'inp_Estado', lst='=lst_Estado')
etiqueta(ws, 'D13', 'Proveedor:')
caja(ws, 'E13:H13', '', 'inp_Proveedor')
etiqueta(ws, 'I13', 'Producto:')
caja(ws, 'J13:K13', 'FASTPLUS', 'inp_Producto', lst='=lst_Producto')
etiqueta(ws, 'L13', 'Ventas anuales:')
caja(ws, 'M13:N13', 'Menos de $50 millones', 'inp_Tier', lst='=rc_Tiers')
etiqueta(ws, 'B15', 'Correo del cliente:')
caja(ws, 'C15:E15', '', 'inp_Correo')
etiqueta(ws, 'F15', 'Obligado solidario:')
ws.merge_cells('F15:G15')
caja(ws, 'H15:K15', 'Por definir', 'inp_Obligado')
etiqueta(ws, 'L15', 'Tipo:')
caja(ws, 'M15:N15', 'Arrendamiento Puro', 'inp_Tipo', lst='=lst_Tipo')

barra(ws, 17, 2, 14, 'Introducir datos para el cálculo (valores IVA incluido, excepto anticipo)')
etiqueta(ws, 'B19', 'Precio equipo:')
caja(ws, 'C19:D19', 879802, 'inp_Precio', MON, note='Precio del proveedor CON IVA (ejemplo: BYD M9 de $758,450 + IVA).')
etiqueta(ws, 'E19', 'Descuento para FASTPLUS (con IVA):')
ws.merge_cells('E19:G19')
caja(ws, 'H19:I19', 0, 'inp_Descuento', MON, note='Descuento que el proveedor otorga a FASTPLUS. No cambia la renta del cliente; mejora el margen.')
etiqueta(ws, 'J19', 'Renta proporcional:')
ws.merge_cells('J19:K19')
caja(ws, 'L19', 0, 'inp_DiasRP', '0" días"', note='Días entre el pago al proveedor y el inicio del plazo. Se cobra renta/30 × días.')
etiqueta(ws, 'B21', 'Valor sin IVA:')
caja(ws, 'C21:D21', '=inp_Precio/(1+inp_IVAEquipo)', 'inp_Valor', MON, salida=True)
etiqueta(ws, 'E21', 'Tarifa aplicable de IVA equipo:')
ws.merge_cells('E21:G21')
caja(ws, 'H21', 0.16, 'inp_IVAEquipo', PCT)
etiqueta(ws, 'J21', 'Uso del anticipo:')
ws.merge_cells('J21:K21')
caja(ws, 'L21:N21', 'Como renta extraordinaria', 'inp_UsoAnticipo', lst='=lst_UsoAnticipo')
etiqueta(ws, 'B23', 'Renta extraordinaria (sin IVA):')
caja(ws, 'C23:D23', 75845, 'inp_Anticipo', MON, note='Anticipo sin IVA (ejemplo: 10% del valor).')
etiqueta(ws, 'E23', 'Tasa IVA rentas:')
ws.merge_cells('E23:G23')
caja(ws, 'H23', '=par_IVA', None, PCT, salida=True)
etiqueta(ws, 'J23', 'Seguro financiado:')
ws.merge_cells('J23:K23')
caja(ws, 'L23:M23', 'No financiado', 'inp_SeguroFin', lst='=lst_Seguro')
etiqueta(ws, 'B25', 'Depósito en garantía:')
caja(ws, 'C25:D25', None, None, MON, salida=True)
etiqueta(ws, 'E25', 'TIIE:')
ws.merge_cells('E25:G25')
caja(ws, 'H25', 0.086, 'inp_TIIE', PCT4)
etiqueta(ws, 'J25', 'Plazo solicitado:')
ws.merge_cells('J25:K25')
caja(ws, 'L25', 24, 'inp_PlazoSol', '0" meses"', lst='=esc_Plazo')
ws['C26'] = None
etiqueta(ws, 'B27', 'Total:')
caja(ws, 'C27:D27', None, None, MON, salida=True)
etiqueta(ws, 'E27', 'Modalidad de rentas:')
ws.merge_cells('E27:G27')
caja(ws, 'H27', 'Vencido', 'inp_Modalidad', lst='=lst_Modalidad')
etiqueta(ws, 'J27', 'Fondeo (interno):')
ws.merge_cells('J27:K27')
caja(ws, 'L27', 0.205, 'inp_Fondeo', PCT)
etiqueta(ws, 'B29', 'Otros gastos (financiados):')
caja(ws, 'C29:D29', 0, 'inp_OtrosMonto', MON, note='IVA incluido. Mantenimiento preventivo, garantía extendida u otros; se financian en la renta.')
etiqueta(ws, 'E29', 'Descripción otros gastos:')
ws.merge_cells('E29:G29')
caja(ws, 'H29:I29', '', 'inp_OtrosDesc')
etiqueta(ws, 'J29', 'GPS (financiado):')
ws.merge_cells('J29:K29')
caja(ws, 'L29:M29', 0, 'inp_GPSMonto', MON, note='IVA incluido. Se financia en la renta.')
etiqueta(ws, 'B31', 'Validación / investigación:')
caja(ws, 'C31:D31', 0, 'inp_GastosInv', MON, note='Sin IVA. Se cobra en el pago inicial.')
etiqueta(ws, 'E31', 'Comisión banco / alianza:')
ws.merge_cells('E31:G31')
caja(ws, 'H31:I31', 0, 'inp_ComBanco', MON)
etiqueta(ws, 'J31', 'Margen objetivo manual:')
ws.merge_cells('J31:K31')
caja(ws, 'L31', None, 'inp_CMobj', PCT, note='Vacío = usa el rate card de la hoja Factores.')

barra(ws, 33, 2, 14, 'Introducir datos para el seguro y tipo de cambio')
ws['B35'] = '=IF(inp_SeguroFin="No financiado","Valor del seguro pago de contado (primer año) IVA incluido:","Seguro por los "&inp_PlazoSol&" meses de financiamiento IVA incluido:")'
ws['B35'].font = f_lab
ws['B35'].alignment = Alignment(horizontal='right')
ws.merge_cells('B35:F35')
caja(ws, 'G35:H35', 0, 'inp_SeguroMonto', MON)
etiqueta(ws, 'I35', 'Nombre aseguradora:')
ws.merge_cells('I35:J35')
caja(ws, 'K35:N35', '', 'inp_Aseguradora')
etiqueta(ws, 'B37', 'Aseguradora por parte de:')
caja(ws, 'C37:D37', 'Cliente', 'inp_AseguradoraPor', lst='=lst_AsegPor')
etiqueta(ws, 'E37', 'Moneda:')
caja(ws, 'F37:G37', 'Moneda Nacional', 'inp_Moneda', lst='=lst_Moneda')
etiqueta(ws, 'H37', 'Tipo de cambio:')
caja(ws, 'I37', 1, 'inp_TC', '#,##0.0000')

barra(ws, 39, 2, 14, 'Condiciones de la operación: resultados del cálculo por plazo')
RES0 = 41

# ==================================================================== FACTORES (parámetros por plazo, formato ABC)
ws = ws_fac
ws.column_dimensions['A'].width = 2
ws.column_dimensions['B'].width = 30
for col in 'CDEFGHIJK':
    ws.column_dimensions[col].width = 15
ws.column_dimensions['L'].width = 2
ws['B2'] = 'Tabla de Factores Cotizador Master FASTPLUS'
ws['B2'].font = Font(name=FONT, size=14, bold=True, italic=True, color=MORADO)
ws.merge_cells('B2:H2')
for c in range(2, 12):
    ws.cell(row=2, column=c).fill = fill_sub

FAC0 = 20
barra(ws, FAC0 - 1, 2, 10, 'Factores por plazo (editar celdas blancas)')
hdr = ['Plazo (meses)', 'Margen renta sobre TIIE', 'Tasa de cálculo renta', '% Residual', 'Residual máximo (rate card)',
       '% Comisión por apertura', '% Depósito (del total financiado)', 'Depósito (# rentas con IVA)', 'Incluir']
for j, h in enumerate(hdr):
    c = ws.cell(row=FAC0, column=2 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
    c.border = box
ws.row_dimensions[FAC0].height = 42
fac_vals = [(12, 0.160, 0.30, 0.02, 0.0, 0), (24, 0.155, 0.20, 0.02, 0.0, 0),
            (36, 0.150, 0.15, 0.02, 0.0, 0), (48, 0.145, 0.10, 0.02, 0.0, 0)]
for i, (pl, mg, res, com, depp, depn) in enumerate(fac_vals):
    rr = FAC0 + 1 + i
    vals = [(pl, '0', True), (mg, PCT, True), ('=inp_TIIE+C%d' % rr, PCT, False), (res, PCT, True),
            ('=INDEX(rc_ResMax,1,IFERROR(MATCH(B%d,rc_Plazos,1),1))' % rr, PCT, False), (com, PCT, True),
            (depp, PCT, True), (depn, '0.00', True), ('Sí', None, True)]
    for j, (v, fmt, editable) in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v)
        c.border = box_in
        c.alignment = Alignment(horizontal='center')
        if fmt:
            c.number_format = fmt
        if editable:
            c.font = f_inbox
            c.protection = Protection(locked=False)
        else:
            c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
            c.fill = fill_out
R1, R4 = FAC0 + 1, FAC0 + 4
for col, nm in zip('BCEGHIJ', ['esc_Plazo', 'esc_Margen', 'esc_Residual', 'esc_Comision', 'esc_DepPct', 'esc_Deposito', 'esc_Incluir']):
    name(nm, absref(ws.title, '%s%d:%s%d' % (col, R1, col, R4)))
dv_list(ws, '=lst_SiNo', 'J%d:J%d' % (R1, R4))
dv_num(ws, 'B%d:B%d' % (R1, R4), 1, 84, 'whole')
dv_num(ws, 'E%d:E%d' % (R1, R4), 0, 0.9)
# columnas internas (ocultas) que alimentan el motor
helpers = [
    ('M', 'esc_Enganche', '=IF(inp_Valor=0,0,inp_Anticipo/inp_Valor)'),
    ('N', 'esc_Sucesivo', 0),
    ('O', 'esc_Tasa', '=D{r}'),
    ('P', 'esc_TasaSuc', '=D{r}'),
    ('Q', 'esc_PagoFinal', None),
    ('R', 'esc_Seguro', '=inp_SeguroMonto/(1+par_IVA)'),
    ('S', 'esc_SeguroForma', '=IF(inp_SeguroFin="Financiado","Financiado","Contado")'),
    ('T', 'esc_GPS', '=inp_GPSMonto/(1+par_IVA)'),
    ('U', 'esc_GPSForma', 'Financiado'),
    ('V', 'esc_Otros', '=inp_OtrosMonto/(1+par_IVA)'),
    ('W', 'esc_OtrosForma', 'Financiado'),
]
for col, nm, f in helpers:
    ws['%s%d' % (col, FAC0)] = nm
    ws['%s%d' % (col, FAC0)].font = f_note
    for rr in range(R1, R4 + 1):
        ws['%s%d' % (col, rr)] = f.format(r=rr) if isinstance(f, str) else f
        ws['%s%d' % (col, rr)].font = f_note
    name(nm, absref(ws.title, '%s%d:%s%d' % (col, R1, col, R4)))
    ws.column_dimensions[col].hidden = True

# ---- rate card
RC0 = 58
barra(ws, RC0 - 1, 2, 10, 'Rate Card FASTPLUS (mínimos de rentabilidad – administrador)')
for j, h in enumerate(['TIER', 'Ventas anuales del solicitante', 'TIR mínima']):
    c = ws.cell(row=RC0, column=2 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
for j, pl in enumerate([12, 24, 36, 48]):
    c = ws.cell(row=RC0, column=5 + j, value=pl)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
    c.number_format = '"CM mín. "0" m"'
tiers = [('1', 'Más de $500 millones', 0.20, [0.020, 0.030, 0.040, 0.050]),
         ('2', '$300 a $500 millones', 0.21, [0.025, 0.035, 0.045, 0.055]),
         ('3', '$100 a $300 millones', 0.22, [0.030, 0.040, 0.050, 0.060]),
         ('4', '$50 a $100 millones', 0.23, [0.035, 0.045, 0.055, 0.065]),
         ('5', 'Menos de $50 millones', 0.24, [0.040, 0.050, 0.060, 0.070])]
for i, (t, v, tir, cms) in enumerate(tiers):
    rr = RC0 + 1 + i
    ws.cell(row=rr, column=2, value=t).alignment = center
    inp(ws.cell(row=rr, column=3), v)
    inp(ws.cell(row=rr, column=4), tir, PCT)
    for j, cm in enumerate(cms):
        inp(ws.cell(row=rr, column=5 + j), cm, PCT)
rr = RC0 + 6
ws.cell(row=rr, column=3, value='Residual máximo por plazo').font = f_bold
for j, rm in enumerate([0.40, 0.38, 0.35, 0.30]):
    inp(ws.cell(row=rr, column=5 + j), rm, PCT)
name('rc_Tiers', absref(ws.title, 'C%d:C%d' % (RC0 + 1, RC0 + 5)))
name('rc_TIR', absref(ws.title, 'D%d:D%d' % (RC0 + 1, RC0 + 5)))
name('rc_CM', absref(ws.title, 'E%d:H%d' % (RC0 + 1, RC0 + 5)))
name('rc_Plazos', absref(ws.title, 'E%d:H%d' % (RC0, RC0)))
name('rc_ResMax', absref(ws.title, 'E%d:H%d' % (rr, rr)))
ws.cell(row=rr + 1, column=2, value='¹ Margen de caja mínimo después de la comisión del originador. Valores sugeridos: ajústelos a la política de FASTPLUS.').font = f_note
ws.cell(row=rr + 4, column=2, value='Aprobado por:').font = f_bold
for col in range(3, 6):
    ws.cell(row=rr + 4, column=col).border = Border(bottom=Side(style='thin', color='000000'))

# ==================================================================== CORRIDAS
FIRST = 60          # primera fila de pagos
NROWS = 98          # filas de pagos (84 + 12 + holgura)
LAST = FIRST + NROWS - 1
T0 = FIRST - 1

P = {}  # etiquetas -> celda de parámetro (columna C)


def build_corrida(ws, k):
    for col, w in zip('ABCDEFGHIJKLMNOPQRS', [7, 17, 7, 12, 16, 15, 14, 15, 16, 13, 16, 16, 11, 16, 16, 12, 16, 16, 16]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['B'].width = 44
    ws['A1'] = '="CORRIDA FINANCIERA – PLAZO "&$C$5&" MESES (escenario "&$C$3&")"'
    ws['A1'].font = f_title
    ws['A2'] = 'Hoja de cálculo interna (protegida). Todos los datos vienen de la hoja Cotizador.'
    ws['A2'].font = f_sub
    params = [
        ('Escenario', k, None, 'esc'),
        ('Incluido en la carta', '=INDEX(esc_Incluir,$C$3)', None, 'incl'),
        ('Plazo básico (n, meses)', '=INDEX(esc_Plazo,$C$3)', '0', 'n'),
        ('Plazo sucesivo (m, meses)', '=INDEX(esc_Sucesivo,$C$3)', '0', 'm'),
        ('Tipo de pago (1 = anticipado, 0 = vencido)', '=IF(inp_Modalidad="Anticipado",1,0)', '0', 'tau'),
        ('Tasa anual plazo básico', '=IF(inp_ModoTasa="TIIE + margen",inp_TIIE+INDEX(esc_Margen,$C$3),INDEX(esc_Tasa,$C$3))', PCT4, 'ia'),
        ('Tasa mensual plazo básico (i)', '=C8/12', PCT4, 'i'),
        ('Tasa anual plazo sucesivo', '=IF(inp_ModoTasa="TIIE + margen",C8,INDEX(esc_TasaSuc,$C$3))', PCT4, 'i2a'),
        ('Tasa mensual plazo sucesivo (i2)', '=C10/12', PCT4, 'i2'),
        ('Tasa de fondeo anual', '=inp_Fondeo', PCT4, 'fa'),
        ('Tasa de fondeo mensual (f)', '=C12/12', PCT4, 'f'),
        ('Tasa de IVA', '=par_IVA', PCT, 'iva'),
        ('Valor del equipo sin IVA (V)', '=inp_Valor', MON, 'V'),
        ('Enganche', '=C15*INDEX(esc_Enganche,$C$3)', MON, 'eng'),
        ('Seguro financiado', '=IF(INDEX(esc_SeguroForma,$C$3)="Financiado",INDEX(esc_Seguro,$C$3),0)', MON, 'segf'),
        ('GPS financiado', '=IF(INDEX(esc_GPSForma,$C$3)="Financiado",INDEX(esc_GPS,$C$3),0)', MON, 'gpsf'),
        ('Monto a financiar (M = V − anticipo + financiados)', '=C15-C16+C17+C18+C42', MON, 'M'),
        ('Valor residual (% del valor)', '=INDEX(esc_Residual,$C$3)', PCT, 'resp'),
        ('Valor residual (VR)', '=C15*C20', MON, 'VR'),
        ('Pago final (% del valor)', '=IF(INDEX(esc_PagoFinal,$C$3)="",C20,INDEX(esc_PagoFinal,$C$3))', PCT, 'pfp'),
        ('Pago final ($)', '=C15*C22', MON, 'PF'),
        ('RENTA BÁSICA  R = PMT(i, n, −M, VR, tipo)', '=PMT(C9,C5,-C19,C21,C7)', MON, 'R'),
        ('Renta plazo sucesivo  Rs = PMT(i2, m, −PF, 0, tipo)', '=IF(C6>0,PMT(C11,C6,-C23,0,C7),0)', MON, 'Rs'),
        ('Opción de compra (si m = 0)', '=IF(C6=0,C23,0)', MON, 'OC'),
        ('Comisión por apertura (%)', '=INDEX(esc_Comision,$C$3)', PCT, 'comp'),
        ('Comisión por apertura ($, sobre el monto financiado)', '=C19*C27', MON, 'com'),
        ('Depósito en garantía (# rentas con IVA)', '=INDEX(esc_Deposito,$C$3)', '0.00', 'depn'),
        ('Depósito en garantía ($ = rentas + % del total financiado)', '=C29*C24*(1+C14)+INDEX(esc_DepPct,$C$3)*C19*(1+C14)', MON, 'dep'),
        ('Fecha de firma / desembolso', '=inp_FechaFirma', FECHA, 'ff'),
        ('Fecha de la primera renta', '=inp_FechaPrimera', FECHA, 'fp'),
        ('Inicio del plazo (t = 0)', '=IF(C7=1,C32,EDATE(C32,-1))', FECHA, 'fi'),
        ('Días entre firma e inicio del plazo', '=MAX(0,C33-C31)', '0', 'd'),
        ('Renta proporcional (R/30 × días)', '=C24/30*C34', MON, 'rp'),
        ('Comisión banco / alianza', '=inp_ComBanco', MON, 'cb'),
        ('Factor de fondeo por días  (1+fa/base)^(−días)', '=(1+C12/par_BaseDias)^(-C34)', '0.000000', 'DF'),
        ('Número de pagos en la tabla', '=C5+IF(C6>0,C6,IF(C23>0,1,0))', '0', 'N'),
        ('Desembolso del arrendador en la firma', '=-(C15-inp_Descuento/(1+inp_IVAEquipo)+C17+C18+C42)', MON, 'out0'),
        ('Cobros en la firma (anticipo+comisión+renta prop.+depósito+gastos−com. banco)', '=C16+C28+C35+C30-C36+C43', MON, 'in0'),
        ('Flujo neto del arrendador en la firma', '=C39+C40', MON, 'net0'),
        ('Mantenimiento / garantía / otros financiados', '=IF(INDEX(esc_OtrosForma,$C$3)="Financiado",INDEX(esc_Otros,$C$3),0)', MON, 'otrf'),
        ('Gastos de investigación / validación', '=inp_GastosInv', MON, 'gi'),
        ('Comisión del promotor ($)', '=inp_ComProm*C19', MON, 'cp'),
        ('Cash margin objetivo (manual o rate card)', '=IF(inp_CMobj="",INDEX(rc_CM,IFERROR(MATCH(inp_Tier,rc_Tiers,0),5),IFERROR(MATCH(C5,rc_Plazos,1),1)),inp_CMobj)', PCT, 'obj'),
        ('TIR mínima (rate card)', '=INDEX(rc_TIR,IFERROR(MATCH(inp_Tier,rc_Tiers,0),5))', PCT, 'tirmin'),
        ('Residual máximo (rate card)', '=INDEX(rc_ResMax,1,IFERROR(MATCH(C5,rc_Plazos,1),1))', PCT, 'resmax'),
        ('Renta: parte del equipo  PMT(i, n, −(V−anticipo), VR, tipo)', '=PMT(C9,C5,-(C15-C16),C21,C7)', MON, 'Req'),
        ('Renta: parte del seguro financiado', '=PMT(C9,C5,-C17,0,C7)', MON, 'Rseg'),
        ('Renta: parte de otros gastos financiados', '=PMT(C9,C5,-C42,0,C7)', MON, 'Rotr'),
        ('Deducibilidad estimada de la renta (ISR)', '=IF(inp_TipoActivo="Automóvil",MIN(1,par_LimAuto*30/C24),IF(inp_TipoActivo="Automóvil eléctrico / híbrido",MIN(1,par_LimAutoEV*30/C24),1))', PCT, 'ded'),
        ('Fecha del último pago', '=IFERROR(INDEX(D%d:D%d,C38),"")' % (FIRST, LAST), FECHA, 'fult'),
        ('Renta: parte del GPS financiado', '=PMT(C9,C5,-C18,0,C7)', MON, 'Rgps'),
    ]
    rr = 3
    for label, f, fmt, key in params:
        ws.cell(row=rr, column=2, value=label).font = f_base
        c = ws.cell(row=rr, column=3)
        calc(c, f, fmt, bold=key in ('R', 'M'))
        if key == 'esc':
            c.font = f_input
        P[key] = 'C%d' % rr
        rr += 1
    assert P['net0'] == 'C41', P['net0']
    assert P['otrf'] == 'C42' and P['gi'] == 'C43' and P['cp'] == 'C44' and P['obj'] == 'C45' and P['fult'] == 'C52', P
    assert P['R'] == 'C24' and P['N'] == 'C38' and P['DF'] == 'C37'

    # ---- resultados
    ws['E3'] = 'RESULTADOS DEL ESCENARIO'
    for c in range(5, 10):
        ws.cell(row=3, column=c).fill = fill_hdr
        ws.cell(row=3, column=c).font = f_hdr
    rows = [
        ('Cash margin C/R ($) = VP al fondeo de todos los flujos', '=C41+SUM(N%d:N%d)' % (FIRST, LAST), MON, 'CM'),
        ('Cash margin C/R (% del monto financiado)', '=IFERROR(I4/C19,0)', PCT, 'CMp'),
        ('Cash margin S/R ($) = sin pago final / residual', '=C41+SUM(S%d:S%d)' % (FIRST, LAST), MON, 'CMs'),
        ('Cash margin S/R (%)', '=IFERROR(I6/C19,0)', PCT, 'CMsp'),
        ('TIR anual efectiva del arrendador (XIRR)', '=IFERROR(XIRR(Q%d:Q%d,P%d:P%d),"n/d")' % (T0, LAST, T0, LAST), PCT, 'TIR'),
        ('TIR nominal anual (capitalizable mensual)', '=IFERROR(12*((1+I8)^(1/12)-1),"n/d")', PCT, 'TIRn'),
        ('CAT informativo sin IVA (flujos del cliente)', '=IFERROR(XIRR(R%d:R%d,P%d:P%d),"n/d")' % (T0, LAST, T0, LAST), PCT, 'CAT'),
        ('Σ factores de descuento de rentas básicas', '=SUMIF(B%d:B%d,"Básico",M%d:M%d)' % (FIRST, LAST, FIRST, LAST), '0.000000', 'SB'),
        ('Factor de descuento del último pago', '=IFERROR(INDEX(M%d:M%d,C38),0)' % (FIRST, LAST), '0.000000', 'ML'),
        ('Sensibilidad del CM a la renta (∂CM/∂R)', '=I11+C29*(1+C14)*(1-I12)+C34/30', '0.000000', 'B'),
        ('Renta mínima para el CM objetivo (neto de comisión)', '=IFERROR((C45*C19+C44-(I4-C24*I13))/I13,0)', MON, 'Rmin'),
        ('Tasa anual mínima para el CM objetivo', '=IFERROR(RATE(C5,I14,-C19,C21,C7)*12,"n/d")', PCT, 'TasaMin'),
        ('Saldo final de la tabla (debe ser 0)', '=IFERROR(INDEX(I%d:I%d,C38),0)' % (FIRST, LAST), MON, 'Saldo'),
        ('Pago inicial: base sin IVA (anticipo+comisión+renta prop.+contado+gastos)', '=C16+C28+C35+C43+I29', MON, 'PIb'),
        ('Pago inicial: IVA', '=I17*C14', MON, 'PIi'),
        ('Pago inicial: depósito en garantía', '=C30', MON, 'PId'),
        ('PAGO INICIAL TOTAL (con IVA)', '=I17+I18+I19', MON, 'PI'),
        ('IVA de la primera renta', '=IFERROR(J%d,0)' % FIRST, MON, 'IVA1'),
        ('Total de pagos con IVA (rentas + pago final)', '=SUM(K%d:K%d)' % (FIRST, LAST), MON, 'TotK'),
        ('Spread (tasa cliente − tasa fondeo)', '=C8-C12', PCT, 'Spread'),
        ('Tasa anual aplicada', '=C8', PCT, 'Tasa'),
        ('Cash margin neto de comisión del promotor ($)', '=I4-C44', MON, 'CMn'),
        ('Cash margin neto (% del monto financiado)', '=IFERROR(I25/C19,0)', PCT, 'CMnp'),
        ('Dictamen vs. rate card', '=IF(AND(I26>=C45,N(I8)>=C46),"CUMPLE",IF(I26>=C45,"CM OK / TIR BAJA","NO CUMPLE"))', None, 'Dict'),
        ('Margen mínimo sobre TIIE (si la tasa mínima aplica)', '=IFERROR(I15-inp_TIIE,"n/d")', PCT, 'MargMin'),
        ('Seguro, GPS y otros de contado (sin IVA)', '=IF(INDEX(esc_SeguroForma,$C$3)="Contado",INDEX(esc_Seguro,$C$3),0)+IF(INDEX(esc_GPSForma,$C$3)="Contado",INDEX(esc_GPS,$C$3),0)+IF(INDEX(esc_OtrosForma,$C$3)="Contado",INDEX(esc_Otros,$C$3),0)', MON, 'Cont'),
    ]
    rr = 4
    for label, f, fmt, key in rows:
        ws.cell(row=rr, column=5, value=label).font = f_base
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=8)
        calc(ws.cell(row=rr, column=9), f, fmt, bold=key in ('CM', 'CMp', 'TIR', 'PI'))
        P[key] = 'I%d' % rr
        rr += 1
    assert P['CM'] == 'I4' and P['CMp'] == 'I5' and P['CMs'] == 'I6' and P['TIR'] == 'I8'
    assert P['SB'] == 'I11' and P['ML'] == 'I12' and P['B'] == 'I13' and P['Rmin'] == 'I14' and P['PIb'] == 'I17'
    assert P['CMn'] == 'I25' and P['CMnp'] == 'I26' and P['Cont'] == 'I29', P

    ws['E34'] = ('Método: CM = flujo en firma + Σ flujo_t × (1+fa/base)^(−días) × (1+f)^(−t). '
                 'Como el CM es lineal en la renta, la renta mínima se despeja en forma cerrada y la tasa mínima '
                 'se obtiene con RATE. La tabla usa interés = (saldo − tipo × pago) × i, de modo que el saldo al término '
                 'del plazo básico es exactamente el valor residual.')
    ws['E34'].font = f_note
    ws['E34'].alignment = wrap
    ws.merge_cells('E34:I40')

    # ---- tabla
    hdr = ['#', 'Concepto', 't', 'Fecha', 'Saldo inicial', 'Pago s/IVA', 'Interés', 'Capital', 'Saldo final',
           'IVA', 'Total c/IVA', 'Flujo arrendador', 'Factor desc.', 'VP flujo', 'Flujo S/R', 'Fecha (XIRR)',
           'Flujo XIRR', 'Flujo cliente (CAT)', 'VP flujo S/R']
    for j, h in enumerate(hdr):
        c = ws.cell(row=T0 - 1, column=1 + j, value=h)
        c.font = f_hdr
        c.fill = fill_hdr
        c.alignment = center
    ws.row_dimensions[T0 - 1].height = 30
    # fila 0 (firma)
    t0 = {1: 0, 2: 'Firma / desembolso', 3: 0, 4: '=C31', 12: '=C41', 16: '=C31', 17: '=C41',
          18: '=C19-C28-C35-C43'}
    for col, v in t0.items():
        c = ws.cell(row=T0, column=col, value=v)
        c.font = f_bold
        c.fill = fill_gris
    for col in range(1, 20):
        ws.cell(row=T0, column=col).fill = fill_gris
    for kk in range(1, NROWS + 1):
        r_ = FIRST + kk - 1
        prev = r_ - 1
        f = {
            'A': kk,
            'B': '=IF(A{r}>$C$38,"",IF(A{r}<=$C$5,"Básico",IF($C$6>0,"Sucesivo","Opción de compra")))',
            'C': '=IF(B{r}="","",IF(B{r}="Opción de compra",$C$5,A{r}-$C$7))',
            'D': '=IF(B{r}="","",EDATE($C$33,C{r}))',
            'E': ('=IF(B{r}="","",$C$19)' if kk == 1 else '=IF(B{r}="","",I{p})'),
            'F': '=IF(B{r}="","",IF(B{r}="Básico",$C$24,IF(B{r}="Sucesivo",$C$25,$C$26)))',
            'G': '=IF(B{r}="","",IF(B{r}="Opción de compra",0,(E{r}-$C$7*F{r})*IF(B{r}="Sucesivo",$C$11,$C$9)))',
            'H': '=IF(B{r}="","",F{r}-G{r})',
            'I': '=IF(B{r}="","",E{r}-H{r})',
            'J': '=IF(B{r}="","",IF(OR(inp_Tipo="Arrendamiento Puro",par_BaseIVAFin="Sobre renta",B{r}="Opción de compra"),F{r},G{r})*$C$14)',
            'K': '=IF(B{r}="","",F{r}+J{r})',
            'L': '=IF(B{r}="","",F{r}-IF(A{r}=$C$38,$C$30,0))',
            'M': '=IF(B{r}="","",$C$37*(1+$C$13)^(-C{r}))',
            'N': '=IF(B{r}="","",L{r}*M{r})',
            'O': '=IF(B{r}="","",IF(B{r}="Básico",F{r},0)-IF(A{r}=$C$38,$C$30,0))',
            'P': '=IF(B{r}="",$C$31,D{r})',
            'Q': '=IF(B{r}="",0,L{r})',
            'R': '=IF(B{r}="",0,-F{r})',
            'S': '=IF(B{r}="","",O{r}*M{r})',
        }
        for col, form in f.items():
            c = ws['%s%d' % (col, r_)]
            c.value = form.format(r=r_, p=prev) if isinstance(form, str) else form
            c.font = f_base
            if col in 'EFGHIJKLNOQRS':
                c.number_format = MON
            elif col in 'DP':
                c.number_format = FECHA
            elif col == 'M':
                c.number_format = '0.000000'
    ws.freeze_panes = 'C%d' % T0
    ws.conditional_formatting.add('A%d:S%d' % (FIRST, LAST),
                                  FormulaRule(formula=['$B%d="Sucesivo"' % FIRST], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.conditional_formatting.add('A%d:S%d' % (FIRST, LAST),
                                  FormulaRule(formula=['$B%d="Opción de compra"' % FIRST], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '%d:%d' % (T0 - 1, T0 - 1)


for k, w in enumerate(ws_cr, start=1):
    build_corrida(w, k)


def cr(key, k):
    return "'Corrida %d'!%s" % (k, P[key].replace('C', '$C$', 1) if P[key][0] == 'C' else P[key].replace('I', '$I$', 1))


# ==================================================================== COTIZADOR: tabla de resultados por plazo
ws = ws_cot
CH = lambda ref: 'CHOOSE(inp_EscTabla,%s)' % ','.join("'Corrida %d'!%s" % (k, ref) for k in range(1, 5))
caja(ws, 'C25:D25', '=%s' % CH('$C$30'), None, MON, salida=True)
caja(ws, 'C27:D27', '=%s' % CH('$C$19'), None, MON, salida=True)
ws['C26'] = '="Anticipo total: "&TEXT(inp_Anticipo*(1+par_IVA)+C25,"$#,##0.00")&" ("&TEXT(IF(inp_Precio=0,0,(inp_Anticipo*(1+par_IVA)+C25)/inp_Precio),"0.00%")&")"'
ws['C26'].font = Font(name=FONT, size=9, bold=True, color=MORADO)
ws.merge_cells('C26:G26')
BLK = ['C', 'E', 'G', 'I']          # cada plazo ocupa 2 columnas (C:D, E:F, G:H, I:J)
hdr_r = RES0
c = ws.cell(row=hdr_r, column=2, value='Conceptos')
for k in range(4):
    a = BLK[k]
    b = L(ord(a) - 64 + 1)
    ws['%s%d' % (a, hdr_r)] = '=IF(inp_EscTabla=%d,"Plazo elegido","")' % (k + 1)
    ws.merge_cells('%s%d:%s%d' % (a, hdr_r, b, hdr_r))
for col in range(2, 12):
    cc = ws.cell(row=hdr_r, column=col)
    cc.fill = fill_hdr
    cc.font = f_bar
    cc.alignment = center
filas = [
    ('Plazo', lambda k: '=INDEX(esc_Plazo,%d)&" meses"' % k, None, None),
    ('Anticipo (sin IVA)', lambda k: '=%s' % cr('eng', k), MON, None),
    ('Depósito en garantía', lambda k: '=%s' % cr('dep', k), MON, None),
    ('Renta', lambda k: '=%s+%s' % (cr('Req', k), cr('Rgps', k)), MON, None),
    ('Seguro', lambda k: '=%s' % cr('Rseg', k), MON, None),
    ('Otros gastos', lambda k: '=%s' % cr('Rotr', k), MON, None),
    ('Total pago mensual (sin IVA)', lambda k: '=%s' % cr('R', k), MON, None),
    ('Total pago mensual (IVA incluido)', lambda k: '=%s+%s' % (cr('R', k), cr('IVA1', k)), MON, None),
    ('Comisión', lambda k: '=%s' % cr('com', k), MON, None),
    ('Valor residual', lambda k: '=%s' % cr('VR', k), MON, None),
    ('Pago inicial total (IVA incluido)', lambda k: '=%s' % cr('PI', k), MON, None),
    ('Dictamen rate card', lambda k: '=%s' % cr('Dict', k), None, 'res_DictCot'),
]
r = hdr_r + 1
for label, fn, fmt, nm in filas:
    lc = ws.cell(row=r, column=2, value=label)
    lc.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    lc.fill = fill_hdr
    for k in range(4):
        a = BLK[k]
        b = L(ord(a) - 64 + 1)
        cell = ws['%s%d' % (a, r)]
        cell.value = fn(k + 1)
        cell.number_format = fmt or 'General'
        cell.alignment = Alignment(horizontal='center')
        cell.font = Font(name=FONT, size=10, bold=label.startswith('Total'))
        ws.merge_cells('%s%d:%s%d' % (a, r, b, r))
        for cc in (a, b):
            ws['%s%d' % (cc, r)].border = box_in
            ws['%s%d' % (cc, r)].fill = PatternFill('solid', fgColor='E9E2F3')
    if nm:
        name(nm, absref(ws.title, 'C%d:J%d' % (r, r)))
    r += 1
RES_LAST = r - 1
for k in range(4):
    a = BLK[k]
    b = L(ord(a) - 64 + 1)
    ws.conditional_formatting.add('%s%d:%s%d' % (a, hdr_r + 1, b, RES_LAST),
                                  FormulaRule(formula=['inp_EscTabla=%d' % (k + 1)], fill=PatternFill('solid', fgColor='FFE699'),
                                              font=Font(name=FONT, bold=True, color='000000')))
ws.conditional_formatting.add('C%d:J%d' % (RES_LAST, RES_LAST), CellIsRule(operator='equal', formula=['"NO CUMPLE"'], font=Font(name=FONT, bold=True, color='C00000')))
ws['B%d' % (RES_LAST + 2)] = 'La columna resaltada es el plazo solicitado; con ese plazo se generan Propuesta, Pago Inicial, Venta, Bonos y Riesgo.'
ws['B%d' % (RES_LAST + 2)].font = f_note
ws.freeze_panes = 'A5'
ws.print_area = 'A1:O%d' % (RES_LAST + 2)
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ==================================================================== FACTORES: pagos, residual y rentabilidad
ws = ws_fac
for i, (lab, f, fmt) in enumerate([('Cliente:', '=inp_Cliente', None), ('Equipo:', '=inp_TipoActivo&" "&inp_Equipo', None),
                                    ('Ventas anuales solicitante:', '=inp_Tier', None)]):
    ws.cell(row=4 + i, column=2, value=lab).font = f_bold
    c = ws.cell(row=4 + i, column=3, value=f)
    c.font = f_base
    ws.merge_cells(start_row=4 + i, start_column=3, end_row=4 + i, end_column=6)
    for col in range(2, 7):
        ws.cell(row=4 + i, column=col).border = box
for i, (lab, f, fmt) in enumerate([('Plazo:', '=inp_PlazoSol&" meses"', None), ('Anticipo:', '=inp_Anticipo', MON),
                                    ('Depósito en garantía:', '=%s' % CH('$C$30'), MON)]):
    ws.cell(row=4 + i, column=8, value=lab).font = f_bold
    c = ws.cell(row=4 + i, column=10, value=f)
    c.font = f_base
    if fmt:
        c.number_format = fmt
    for col in range(8, 11):
        ws.cell(row=4 + i, column=col).border = box


def tabla_kv(ws, r0, c0, titulo, filas_):
    barra(ws, r0, c0, c0 + 2, titulo)
    for i, (lab, f, fmt, bold) in enumerate(filas_):
        rr = r0 + 1 + i
        a = ws.cell(row=rr, column=c0, value=lab)
        a.font = Font(name=FONT, size=10, bold=bold, italic=not bold)
        ws.merge_cells(start_row=rr, start_column=c0, end_row=rr, end_column=c0 + 1)
        b = ws.cell(row=rr, column=c0 + 2, value=f)
        b.font = Font(name=FONT, size=10, bold=bold)
        b.number_format = fmt or 'General'
        b.alignment = Alignment(horizontal='right')
        for col in range(c0, c0 + 3):
            ws.cell(row=rr, column=col).border = box
            if bold:
                ws.cell(row=rr, column=col).fill = fill_sub


tabla_kv(ws, 8, 2, 'Pagos mensuales (no incluyen IVA)', [
    ('Renta básica', '=%s+%s' % (CH('$C$48'), CH('$C$53')), MON, False),
    ('Pago por seguro', '=%s' % CH('$C$49'), MON, False),
    ('Otros gastos', '=%s' % CH('$C$50'), MON, False),
    ('Total pago mensual', '=%s' % CH('$C$24'), MON, True),
    ('IVA', '=%s' % CH('$I$21'), MON, False),
    ('Pago mensual IVA incluido', '=%s+%s' % (CH('$C$24'), CH('$I$21')), MON, True)])
tabla_kv(ws, 8, 7, 'Residual', [
    ('Valor residual nominal', '=%s' % CH('$C$21'), MON, False),
    ('Porcentaje', '=%s' % CH('$C$20'), PCT, False),
    ('IVA', '=%s*par_IVA' % CH('$C$21'), MON, False),
    ('Residual neto con IVA', '=%s*(1+par_IVA)' % CH('$C$21'), MON, True),
    ('Deducibilidad de la renta', '=%s' % CH('$C$51'), PCT, True)])

# ---- análisis de rentabilidad del plazo solicitado (Operación vs Rate card)
AR0 = 26
barra(ws, AR0, 2, 10, 'Análisis de rentabilidad (plazo solicitado)')
for j, h in enumerate(['Concepto', 'Operación', 'Rate card', 'Observaciones']):
    c = ws.cell(row=AR0 + 1, column=[2, 3, 4, 5][j], value=h)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
ws.merge_cells(start_row=AR0 + 1, start_column=5, end_row=AR0 + 1, end_column=10)
an = [
    ('TIR real (XIRR)', CH('$I$8'), CH('$C$46'), 'TIR menor a la mínima'),
    ('TIR nominal (cap. mensual)', CH('$I$9'), CH('$C$46'), None),
    ('Margen antes de comisiones', CH('$I$5'), CH('$C$45'), 'Margen antes de comisiones menor al mínimo'),
    ('Margen de caja neto', CH('$I$26'), CH('$C$45'), 'Margen de caja neto menor al mínimo'),
    ('Tasa de cálculo de la renta', CH('$C$8'), CH('$I$15'), 'Tasa menor a la mínima para cumplir el rate card'),
    ('Spread contra fondeo', CH('$I$23'), None, None),
    ('Residual', CH('$C$20'), CH('$C$47'), '__RES__'),
]
for i, (lab, op, rc, obs) in enumerate(an):
    rr = AR0 + 2 + i
    ws.cell(row=rr, column=2, value=lab).font = f_base
    a = ws.cell(row=rr, column=3, value='=' + op)
    a.number_format = PCT
    a.alignment = Alignment(horizontal='center')
    if rc:
        b = ws.cell(row=rr, column=4, value='=' + rc)
        b.number_format = PCT
        b.alignment = Alignment(horizontal='center')
    if obs == '__RES__':
        f = '=IF(C{r}>D{r},"Residual mayor al máximo permitido en "&TEXT(C{r}-D{r},"0.00%"),"")'.format(r=rr)
    elif obs:
        f = '=IF(N(C{r})<D{r},"{t} en "&TEXT(N(C{r})-D{r},"0.00%"),"")'.format(r=rr, t=obs)
    else:
        f = None
    o = ws.cell(row=rr, column=5, value=f)
    o.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=10)
    ws.conditional_formatting.add('E%d:J%d' % (rr, rr), FormulaRule(formula=['$E%d<>""' % rr], fill=PatternFill('solid', fgColor='FF0000')))
    for col in range(2, 11):
        ws.cell(row=rr, column=col).border = box
DICT_R = AR0 + 2 + len(an)
ws.cell(row=DICT_R, column=2, value='Dictamen').font = f_big
d = ws.cell(row=DICT_R, column=3, value='=' + CH('$I$27'))
d.font = Font(name=FONT, size=12, bold=True)
ws.merge_cells(start_row=DICT_R, start_column=3, end_row=DICT_R, end_column=4)
ws.conditional_formatting.add('C%d' % DICT_R, CellIsRule(operator='equal', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(name=FONT, bold=True, color='006100')))
ws.conditional_formatting.add('C%d' % DICT_R, CellIsRule(operator='notEqual', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(name=FONT, bold=True, color='9C0006')))

# ---- análisis por plazo (los 4 plazos)
AP0 = 38
barra(ws, AP0, 2, 6, 'Análisis por plazo')
ws.cell(row=AP0 + 1, column=2, value='Concepto').font = f_bold
for k in range(4):
    c = ws.cell(row=AP0 + 1, column=3 + k, value='=INDEX(esc_Plazo,%d)&" meses"' % (k + 1))
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
ws.cell(row=AP0 + 1, column=2).fill = fill_sub
grid = [
    ('Renta mensual sin IVA', 'R', MON, 'res_Renta'),
    ('Renta mensual IVA incluido', 'RIVA', MON, 'res_RentaIVA'),
    ('Tasa de cálculo renta', 'Tasa', PCT, 'res_Tasa'),
    ('TIR real', 'TIR', PCT, 'res_TIR'),
    ('TIR mínima (rate card)', 'tirmin', PCT, None),
    ('Margen de caja antes de comisiones (%)', 'CMp', PCT, 'res_CMpct'),
    ('Comisión originador ($)', 'cp', MON, None),
    ('Margen de caja neto ($)', 'CMn', MON, 'res_CMneto'),
    ('Margen de caja neto (%)', 'CMnp', PCT, 'res_CMnetoPct'),
    ('Margen mínimo (rate card)', 'obj', PCT, None),
    ('Dictamen', 'Dict', None, 'res_Dictamen'),
    ('Tasa mínima para cumplir', 'TasaMin', PCT, 'res_TasaMin'),
    ('Margen mínimo sobre TIIE', 'MargMin', PCT, None),
    ('CAT informativo sin IVA', 'CAT', PCT, None),
    ('Cuadre de la corrida', 'Saldo', None, None),
]
for i, (lab, key, fmt, nm) in enumerate(grid):
    rr = AP0 + 2 + i
    ws.cell(row=rr, column=2, value=lab).font = f_base
    for k in range(4):
        if key == 'RIVA':
            f = '=%s+%s' % (cr('R', k + 1), cr('IVA1', k + 1))
        elif key == 'Saldo':
            f = '=IF(ABS(%s)<0.01,"OK","Dif. "&TEXT(%s,"#,##0.00"))' % (cr('Saldo', k + 1), cr('Saldo', k + 1))
        else:
            f = '=%s' % cr(key, k + 1)
        c = ws.cell(row=rr, column=3 + k, value=f)
        c.number_format = fmt or 'General'
        c.alignment = Alignment(horizontal='center')
        c.font = f_link
        c.border = box
    if nm:
        name(nm, absref(ws.title, 'C%d:F%d' % (rr, rr)))
    if key == 'Dict':
        ws.conditional_formatting.add('C%d:F%d' % (rr, rr), CellIsRule(operator='equal', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='C6EFCE')))
        ws.conditional_formatting.add('C%d:F%d' % (rr, rr), CellIsRule(operator='notEqual', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='FFC7CE')))
rr = AP0 + 2 + len(grid)
ws.cell(row=rr, column=2, value='Alertas').font = f_bold
for k in range(1, 5):
    E = lambda nm: 'INDEX(%s,%d)' % (nm, k)
    f = ('IF(inp_Valor<=0,"Capture el precio del equipo. ","")'
         '&IF({tasa}<inp_Fondeo,"Spread negativo (tasa < fondeo). ","")'
         '&IF({cmp}<{obj},"Margen neto por debajo del mínimo. ","")'
         '&IF(N({tir})<{tirmin},"TIR por debajo de la mínima. ","")'
         '&IF({res}>{resmax},"Residual mayor al máximo. ","")'
         '&IF(inp_Anticipo>=inp_Valor,"Anticipo mayor o igual al valor. ","")'
         '&IF(ABS({saldo})>=0.01,"La corrida no cuadra. ","")').format(
        tasa=cr('ia', k), cmp=cr('CMnp', k), obj=cr('obj', k), tir=cr('TIR', k), tirmin=cr('tirmin', k),
        res=E('esc_Residual'), resmax=cr('resmax', k), saldo=cr('Saldo', k))
    c = ws.cell(row=rr, column=2 + k, value='=IF(%s="","Sin alertas",%s)' % (f, f))
    c.font = Font(name=FONT, size=8, color='C00000')
    c.alignment = wrap
    c.border = box
ws.row_dimensions[rr].height = 60
name('res_Alertas', absref(ws.title, 'C%d:F%d' % (rr, rr)))
assert rr < RC0 - 2, (rr, RC0)
ws.print_area = 'A1:K%d' % (RC0 + 11)
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ==================================================================== TABLA DE PAGOS
ws = ws_tab
for col, w in zip('ABCDEFGHI', [2, 7, 13, 22, 17, 15, 17, 18, 2]):
    ws.column_dimensions[col].width = w
logo(ws, 'B1', 55)
ws['D1'] = 'TABLA DE PAGOS'
ws['D1'].font = f_title
ws.merge_cells('D1:H1')
ws['D2'] = '=par_Arrendador'
ws['D2'].font = f_sub
ws.merge_cells('D2:H2')
lab = [('Cliente:', '=inp_Cliente'), ('Equipo:', '=inp_Equipo'), ('Folio:', '=inp_Folio')]
for i, (a, b) in enumerate(lab):
    ws.cell(row=4 + i, column=2, value=a).font = f_bold
    ws.cell(row=4 + i, column=4, value=b).font = f_base
    ws.merge_cells(start_row=4 + i, start_column=4, end_row=4 + i, end_column=8)
ws['B7'] = 'Plazo mostrado:'
ws['B7'].font = f_bold
ws['F7'] = 'Se elige en el Cotizador (plazo solicitado)'
ws['F7'].font = f_note
ws.merge_cells('B7:C7')
ws['D7'] = '=inp_PlazoSol&" meses"'
ws['D7'].font = f_link
ws['D7'].alignment = Alignment(horizontal='left')
CH = lambda ref: 'CHOOSE(inp_EscTabla,%s)' % ','.join("'Corrida %d'!%s" % (k, ref) for k in range(1, 5))
ws['B8'] = '="Plazo: "&%s&" meses + "&%s&" sucesivo  |  Tasa anual: "&TEXT(%s,"0.00%%")&"  |  Pago "&LOWER(inp_Modalidad)&"  |  "&inp_Moneda' % (CH('$C$5'), CH('$C$6'), CH('$C$8'))
ws['B8'].font = f_base
ws.merge_cells('B8:H8')
hdr = ['No.', 'Fecha', 'Concepto', 'Importe sin IVA', 'IVA', 'Total con IVA', 'Acumulado con IVA']
for j, h in enumerate(hdr):
    c = ws.cell(row=10, column=2 + j, value=h)
    c.font = f_hdr
    c.fill = fill_hdr
    c.alignment = center
# pago inicial
r = 11
ws.cell(row=r, column=2, value=0)
ws.cell(row=r, column=3, value='=inp_FechaFirma').number_format = FECHA
ws.cell(row=r, column=4, value='Pago inicial')
ws.cell(row=r, column=5, value='=%s+%s' % (CH('$I$17'), CH('$I$19'))).number_format = MON
ws.cell(row=r, column=6, value='=%s' % CH('$I$18')).number_format = MON
ws.cell(row=r, column=7, value='=E11+F11').number_format = MON
ws.cell(row=r, column=8, value='=G11').number_format = MON
for col in range(2, 9):
    ws.cell(row=r, column=col).fill = fill_gris
    ws.cell(row=r, column=col).font = f_base
for kk in range(1, NROWS + 1):
    r = 11 + kk
    src = FIRST + kk - 1
    ws.cell(row=r, column=2, value='=IF(%s="","",%d)' % (CH('$B$%d' % src), kk))
    ws.cell(row=r, column=3, value='=IF(B%d="","",%s)' % (r, CH('$D$%d' % src))).number_format = FECHA
    ws.cell(row=r, column=4, value='=IF(B%d="","",IF(%s="Básico","Renta mensual",IF(%s="Sucesivo","Renta plazo sucesivo","Opción de compra")))' % (r, CH('$B$%d' % src), CH('$B$%d' % src)))
    ws.cell(row=r, column=5, value='=IF(B%d="","",%s)' % (r, CH('$F$%d' % src))).number_format = MON
    ws.cell(row=r, column=6, value='=IF(B%d="","",%s)' % (r, CH('$J$%d' % src))).number_format = MON
    ws.cell(row=r, column=7, value='=IF(B%d="","",E%d+F%d)' % (r, r, r)).number_format = MON
    ws.cell(row=r, column=8, value='=IF(B%d="","",H%d+G%d)' % (r, r - 1, r)).number_format = MON
    for col in range(2, 9):
        ws.cell(row=r, column=col).font = f_base
        ws.cell(row=r, column=col).alignment = Alignment(horizontal='center') if col <= 4 else Alignment()
TOTR = 11 + NROWS + 1
ws.cell(row=TOTR, column=4, value='TOTAL').font = f_bold
for col, L_ in ((5, 'E'), (6, 'F'), (7, 'G')):
    c = ws.cell(row=TOTR, column=col, value='=SUM(%s11:%s%d)' % (L_, L_, TOTR - 1))
    c.number_format = MON
    c.font = f_bold
    c.border = Border(top=Side(style='thin', color='000000'))
ws.cell(row=TOTR + 1, column=2, value='El depósito en garantía se incluye en el pago inicial y se aplica en el último pago. Cifras sujetas a aprobación de crédito.').font = f_note
ws.merge_cells(start_row=TOTR + 1, start_column=2, end_row=TOTR + 1, end_column=8)
ws.conditional_formatting.add('B12:H%d' % (TOTR - 1), FormulaRule(formula=['AND($B12<>"",MOD($B12,2)=0)'], fill=PatternFill('solid', fgColor='F7F4FB')))
ws.freeze_panes = 'B11'
ws.print_area = 'A1:I%d' % (TOTR + 1)
ws.print_title_rows = '10:10'
ws.page_setup.orientation = 'portrait'
ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddFooter.center.text = 'Página &P de &N'

# ==================================================================== DOCUMENTOS (formato tipo ABC)
MESES = '"enero","febrero","marzo","abril","mayo","junio","julio","agosto","septiembre","octubre","noviembre","diciembre"'
FECHA_TXT = 'DAY({d})&" de "&CHOOSE(MONTH({d}),%s)&" de "&YEAR({d})' % MESES
IVA1 = '(1+par_IVA)'
fill_box = PatternFill('solid', fgColor='E9E2F3')
fill_band = PatternFill('solid', fgColor='D9CCEB')
f_head = Font(name=FONT, size=9, bold=True, italic=True, color=MORADO)
thin_b = Border(bottom=Side(style='hair', color='BFBFBF'))
total_b = Border(top=Side(style='thin', color='000000'))


def doc_sheet(ws, widths, last_row, landscape=False):
    for i, w in enumerate(widths):
        ws.column_dimensions[L(i + 1)].width = w
    ws.print_area = 'A1:%s%d' % (L(len(widths)), last_row)
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = page.PageMargins(left=0.45, right=0.45, top=0.45, bottom=0.5)


def membrete_abc(ws, last_col):
    """Datos de la empresa a la izquierda (cursiva) y logo a la derecha, como el formato ABC."""
    lines = ['=par_Arrendador', '=sel_Dir1', '=sel_Dir2&" "&sel_Dir3',
             '=IF(sel_Tel="","","Teléfono: "&sel_Tel)', '=sel_PromCorreo',
             '=IF(par_Contacto="","","Visite "&par_Contacto)']
    for i, f in enumerate(lines, start=1):
        c = ws.cell(row=i, column=2, value=f)
        c.font = f_head
    ws.cell(row=7, column=2, value='=par_Comercial').font = Font(name=FONT, size=11, bold=True, italic=True, color=MORADO)
    logo(ws, '%s1' % L(last_col - 1), 62)


def banda(ws, row, c1, c2, text, align='center'):
    c = ws.cell(row=row, column=c1, value=text)
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = fill_band
        ws.cell(row=row, column=col).font = Font(name=FONT, size=10, bold=True, italic=True)
    c.alignment = Alignment(horizontal=align, vertical='center')
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)


def kv(ws, row, cl, cv, label, formula, fmt=MON, bold=False, fill=None):
    a = ws.cell(row=row, column=cl, value=label)
    b = ws.cell(row=row, column=cv, value=formula)
    a.font = Font(name=FONT, size=10, bold=bold, italic=True)
    b.font = Font(name=FONT, size=10, bold=bold, italic=True)
    b.alignment = Alignment(horizontal='right')
    if fmt:
        b.number_format = fmt
    if fill:
        for col in range(cl, cv + 1):
            ws.cell(row=row, column=col).fill = fill


def parrafo(ws, row, c1, c2, formula, height, font=None, align='justify'):
    c = ws.cell(row=row, column=c1, value=formula)
    c.font = font or Font(name=FONT, size=9, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical='top', horizontal=align)
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.row_dimensions[row].height = height


def firmas(ws, row, izq, der, sub_izq, sub_der, c_izq=2, c_der=5):
    ws.cell(row=row, column=c_izq, value=izq).font = Font(name=FONT, size=10, bold=True, italic=True)
    ws.cell(row=row, column=c_der, value=der).font = Font(name=FONT, size=10, bold=True, italic=True)
    for i, (a, b) in enumerate(zip(sub_izq, sub_der)):
        x = ws.cell(row=row + 3 + i, column=c_izq, value=a)
        y = ws.cell(row=row + 3 + i, column=c_der, value=b)
        for c in (x, y):
            c.font = Font(name=FONT, size=10, bold=True, italic=True)
            c.alignment = Alignment(horizontal='center')
        ws.merge_cells(start_row=row + 3 + i, start_column=c_izq, end_row=row + 3 + i, end_column=c_izq + 1)
        ws.merge_cells(start_row=row + 3 + i, start_column=c_der, end_row=row + 3 + i, end_column=c_der + 1)


SEGURO_CONTADO = '%s*%s' % (CH('$I$29'), IVA1)   # seguro/GPS/otros de contado con IVA

# ----------------------------------------------------------------- PROPUESTA
ws = ws_pro
membrete_abc(ws, 6)
ws['B9'] = '="Apreciable:   "&IF(inp_Contacto="",inp_Cliente,inp_Contacto)'
ws['B9'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E9'] = '=' + FECHA_TXT.format(d='inp_Fecha')
ws['E9'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['E9'].alignment = Alignment(horizontal='right')
ws.merge_cells('E9:F9')
ws['B11'] = 'P r e s e n t e'
ws['B11'].font = Font(name=FONT, size=10, bold=True, italic=True)
ws['F11'] = '="Folio: "&inp_Folio'
ws['F11'].font = Font(name=FONT, size=9, italic=True)
ws['F11'].alignment = Alignment(horizontal='right')
parrafo(ws, 13, 2, 6, '=par_Arrendador&", a través de su producto "&par_Comercial&", se complace en poner a su amable consideración la presente propuesta para celebrar una operación de "&LOWER(inp_Tipo)&" sobre el equipo que a continuación se describe:"', 30, Font(name=FONT, size=10, italic=True))
banda(ws, 15, 2, 6, 'Descripción de equipo')
for i, (lab, f) in enumerate([('Tipo de activo', '=inp_TipoActivo'), ('Descripción de activo', '=inp_Equipo')]):
    ws.cell(row=16 + i, column=2, value=lab).font = Font(name=FONT, size=10, italic=True)
    c = ws.cell(row=16 + i, column=3, value=f)
    c.font = Font(name=FONT, size=10, italic=True)
    ws.merge_cells(start_row=16 + i, start_column=3, end_row=16 + i, end_column=6)
banda(ws, 19, 2, 3, 'Valor equipos y conceptos financiados (IVA incluido)')
kv(ws, 20, 2, 3, 'Valor del activo', '=inp_Precio')
kv(ws, 21, 2, 3, 'Seguro financiado¹', '=%s*%s' % (CH('$C$17'), IVA1))
kv(ws, 22, 2, 3, '="Otros gastos¹"&IF(inp_OtrosDesc="",""," ("&inp_OtrosDesc&")")', '=(%s+%s)*%s' % (CH('$C$18'), CH('$C$42'), IVA1))
kv(ws, 23, 2, 3, '=IF(inp_UsoAnticipo="Como enganche","(-menos) Enganche","(-menos) Renta extraordinaria")', '=-%s*%s' % (CH('$C$16'), IVA1))
kv(ws, 24, 2, 3, 'Monto Total Financiado', '=SUM(C20:C23)', bold=True, fill=fill_band)
ws['E19'] = 'Plazo del arrendamiento:'
ws['F19'] = '=inp_PlazoSol&" meses"'
for c in ('E19', 'F19'):
    ws[c].font = Font(name=FONT, size=10, bold=True, italic=True)
    ws[c].fill = fill_band
ws['F19'].alignment = Alignment(horizontal='right')
banda(ws, 20, 5, 6, 'Pago inicial')
kv(ws, 21, 5, 6, 'Depósito en garantía', '=%s' % CH('$C$30'))
kv(ws, 22, 5, 6, 'Renta proporcional', '=%s' % CH('$C$35'))
kv(ws, 23, 5, 6, '=IF(inp_SeguroFin="Financiado","Seguro primer año (financiado)","Seguro primer año")', '=%s' % CH('$I$29'))
kv(ws, 24, 5, 6, '=IF(inp_UsoAnticipo="Como enganche","Enganche","Renta extraordinaria")', '=%s' % CH('$C$16'))
kv(ws, 25, 5, 6, 'Comisión por apertura', '=%s' % CH('$C$28'))
kv(ws, 26, 5, 6, 'Validación de cuenta²', '=%s' % CH('$C$43'))
kv(ws, 27, 5, 6, 'Sub-total pago inicial', '=SUM(F21:F26)', bold=True)
kv(ws, 28, 5, 6, 'Impuesto al Valor Agregado', '=SUM(F22:F26)*par_IVA')
kv(ws, 29, 5, 6, 'Total Pago Inicial', '=F27+F28', bold=True, fill=fill_band)
banda(ws, 31, 2, 3, 'Pagos mensuales (no incluyen IVA)')
kv(ws, 32, 2, 3, 'Renta básica', '=%s+%s' % (CH('$C$48'), CH('$C$53')))
kv(ws, 33, 2, 3, 'Pago por seguro', '=%s' % CH('$C$49'))
kv(ws, 34, 2, 3, 'Otros gastos', '=%s' % CH('$C$50'))
kv(ws, 35, 2, 3, 'Total Pago Mensual', '=SUM(C32:C34)', bold=True, fill=fill_band)
banda(ws, 33, 5, 6, 'Valor residual (no incluye IVA)')
kv(ws, 34, 5, 6, 'Valor residual', '=%s' % CH('$C$21'))
kv(ws, 35, 5, 6, 'Vigencia propuesta', '=inp_Fecha+par_Vigencia', FECHA)
parrafo(ws, 38, 2, 6, '="¹/  Montos financiados por el plazo total del arrendamiento."', 14, Font(name=FONT, size=9, bold=True, italic=True), 'left')
parrafo(ws, 39, 2, 6, '="²/  La validación de la cuenta se cobra por separado, a través de la cuenta domiciliada para el cobro de las rentas."', 14, Font(name=FONT, size=9, bold=True, italic=True), 'left')
parrafo(ws, 41, 2, 6, 'Cualquier cambio respecto a las condiciones establecidas dentro de la aprobación de crédito correspondiente invalida la presente propuesta, misma que únicamente será posible ejercer durante su vigencia.', 26, Font(name=FONT, size=10, italic=True))
parrafo(ws, 42, 2, 6, '="Para formalizar la presente operación será necesario: 1) Realizar el pago inicial por la cantidad de "&TEXT(F29,"$#,##0.00")&" y 2) Firmar el contrato de arrendamiento correspondiente."', 26, Font(name=FONT, size=10, italic=True))
parrafo(ws, 43, 2, 6, '="En "&par_Comercial&" nos esforzamos por otorgarle un buen servicio; si tiene cualquier duda o comentario respecto a la presente propuesta, no dude en contactar a su ejecutivo "&sel_PromNombre&IF(sel_PromCorreo="",""," ("&sel_PromCorreo&")")&"."', 26, Font(name=FONT, size=10, italic=True))
firmas(ws, 46, 'Cordialmente', 'Propuesta aceptada por:', ['=sel_PromNombre', '=par_Arrendador'],
       ['=IF(inp_Contacto="","Representante legal",inp_Contacto)', '=inp_Cliente'])
parrafo(ws, 52, 2, 6, '=IF(inp_SeguroFin="No financiado","Nota importante: bajo la opción de pago de la prima de seguro anual de contado, es posible que las anualidades siguientes sufran incrementos, dependiendo de la siniestralidad del bien o de las políticas de la aseguradora; bajo esta modalidad se acepta que dicha prima pudiera incrementarse en las anualidades subsecuentes.","")', 40, Font(name=FONT, size=9, bold=True, italic=True), 'center')
doc_sheet(ws, [2, 33, 16, 3, 33, 18, 2], 53)

# ----------------------------------------------------------------- PAGO INICIAL
ws = ws_pin
logo(ws, 'B1', 60)
datos_pi = [('Nombre del cliente', '=inp_Cliente'), ('Representante', '=inp_Contacto&""'),
            ('Bien arrendado', '=inp_TipoActivo&" "&inp_Equipo'), ('Opción tomada para el anticipo', '=inp_UsoAnticipo')]
for i, (lab, f) in enumerate(datos_pi):
    a = ws.cell(row=2 + i, column=4, value=lab)
    a.font = f_bold
    a.fill = fill_band
    b = ws.cell(row=2 + i, column=5, value=f)
    b.font = f_base
    ws.merge_cells(start_row=2 + i, start_column=5, end_row=2 + i, end_column=7)
    for col in range(4, 8):
        ws.cell(row=2 + i, column=col).border = box
banda(ws, 8, 2, 3, 'Conceptos de Pago Inicial')
pi = [('Depósito en garantía', '=Propuesta!F21'), ('Renta proporcional', '=Propuesta!F22'), ('=Propuesta!E23', '=Propuesta!F23'),
      ('Comisión por apertura', '=Propuesta!F25'), ('Validación de cuenta', '=Propuesta!F26'), ('=Propuesta!E24', '=Propuesta!F24'),
      ('Subtotal antes de IVA', '=Propuesta!F27'), ('IVA conceptos anteriores', '=Propuesta!F28'), ('Total Pago Inicial', '=Propuesta!F29')]
for i, (lab, f) in enumerate(pi):
    rr = 9 + i
    bold = lab in ('Total Pago Inicial',)
    kv(ws, rr, 2, 3, lab, f, bold=bold, fill=fill_band if bold else None)
    for col in (2, 3):
        ws.cell(row=rr, column=col).border = box
banda(ws, 8, 5, 7, 'Datos para Depósito Pago Inicial')
dep = [('__T', 'Transferencia electrónica:'), ('Banco', '=IF(par_Banco="","(por definir)",par_Banco)'),
       ('CLABE', '=IF(par_CLABE="","(por definir)",par_CLABE)'),
       ('__T', 'Depósito en sucursal bancaria:'), ('Banco', '=IF(par_BancoSuc="","(por definir)",par_BancoSuc)'),
       ('Cuenta', '=IF(par_Cuenta="","(por definir)",par_Cuenta)'), ('Convenio', '=IF(par_Convenio="","N/A",par_Convenio)'),
       ('Referencia', '=inp_Folio'),
       ('__T', 'En ambos casos:'), ('Beneficiario', '=par_Beneficiario'), ('Moneda', '=inp_Moneda'),
       ('Pago seguro', '=IF(inp_SeguroFin="Financiado","Seguro financiado en la renta mensual.","Seguro no financiado: primer pago anual incluido en el pago inicial; anualidades siguientes pendientes de pago.")')]
for i, (lab, f) in enumerate(dep):
    rr = 9 + i
    if lab == '__T':
        c = ws.cell(row=rr, column=5, value=f)
        c.font = f_bold
        c.alignment = center
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=7)
        for col in range(5, 8):
            ws.cell(row=rr, column=col).fill = fill_box
    else:
        ws.cell(row=rr, column=5, value=lab).font = f_base
        c = ws.cell(row=rr, column=6, value=f)
        c.font = f_base
        c.alignment = Alignment(horizontal='center', wrap_text=True, vertical='top')
        ws.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=7)
    for col in range(5, 8):
        ws.cell(row=rr, column=col).border = box
ws.row_dimensions[20].height = 40
parrafo(ws, 23, 2, 7, 'Bajo la opción de pago de la prima de seguro anual de contado, las anualidades subsecuentes pueden sufrir incrementos, dependiendo de la siniestralidad del bien o de cambios en las políticas de la aseguradora; bajo esta modalidad se acepta que dicha prima pudiera incrementarse en las anualidades subsecuentes.', 44, Font(name=FONT, size=10, bold=True))
parrafo(ws, 25, 2, 7, 'El pago inicial no incluye los gastos generados por concepto de placas, tenencia (en su caso) y gestoría. Estos conceptos, una vez determinados conforme a la reglamentación vigente en cada localidad, serán cargados a la cuenta a domiciliar.', 44, Font(name=FONT, size=10, bold=True))
firmas(ws, 28, 'Cordialmente', 'Propuesta aceptada por', ['=sel_PromNombre', '=par_Arrendador'],
       ['=IF(inp_Contacto="","Representante legal",inp_Contacto)', '=inp_Cliente'], 2, 5)
doc_sheet(ws, [2, 28, 16, 22, 16, 22, 16, 2], 33)

# ----------------------------------------------------------------- VENTA (promesa de venta del residual)
ws = ws_ven
membrete_abc(ws, 6)
ws['B9'] = '="Apreciable:   "&IF(inp_Contacto="",inp_Cliente,inp_Contacto)'
ws['E9'] = '=' + FECHA_TXT.format(d='inp_Fecha')
ws['E9'].alignment = Alignment(horizontal='right')
ws.merge_cells('E9:F9')
ws['B11'] = 'P r e s e n t e .'
for c in ('B9', 'E9', 'B11'):
    ws[c].font = Font(name=FONT, size=10, bold=True, italic=True)
parrafo(ws, 13, 2, 6, '=par_Arrendador&", (el promitente vendedor) venderá al término del contrato de arrendamiento celebrado con "&inp_Cliente&" a la persona que éste designe y que se conocerá como (el promitente comprador), el bien usado que a continuación se describe."', 44, Font(name=FONT, size=10, italic=True))
venta = [('Tipo de bien:', '=inp_TipoActivo&" "&inp_Equipo', None),
         ('Valor para venta de la unidad IVA incluido:', '=%s*%s' % (CH('$C$23'), IVA1), MON),
         ('Plazo para la venta después de:', '=inp_PlazoSol&" meses"', None),
         ('Fecha de pago de esta propuesta:', '=%s' % CH('$C$52'), FECHA)]
for i, (lab, f, fmt) in enumerate(venta):
    rr = 15 + 2 * i
    a = ws.cell(row=rr, column=2, value=lab)
    a.font = Font(name=FONT, size=10, italic=True, underline='single')
    b = ws.cell(row=rr, column=5, value=f)
    b.font = Font(name=FONT, size=10, bold=True, italic=True)
    b.alignment = Alignment(horizontal='right')
    if fmt:
        b.number_format = fmt
    ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=6)
banda(ws, 24, 2, 6, 'Términos y condiciones')
parrafo(ws, 26, 2, 6, 'El bien será entregado por "el promitente vendedor" en el estado en que lo reciba del arrendatario al término del contrato. El cambio de propietario, así como todos los derechos e impuestos que cause la compraventa, serán a cargo de "el promitente comprador", quien deberá estar al corriente en todos los pagos del contrato de arrendamiento.', 44, Font(name=FONT, size=10, italic=True))
parrafo(ws, 27, 2, 6, 'Las partes celebrarán un contrato de promesa de compraventa bajo los términos anteriores. El precio podrá ajustarse si existen rentas u otros adeudos pendientes a la fecha de la venta.', 30, Font(name=FONT, size=10, italic=True))
parrafo(ws, 28, 2, 6, 'Esperamos que la presente propuesta cubra sus expectativas y podamos contar con una respuesta afirmativa. Si tiene cualquier pregunta o aclaración, no dude en comunicarse con nosotros.', 30, Font(name=FONT, size=10, italic=True))
firmas(ws, 30, 'Cordialmente:', 'Acepto:', ['=sel_PromNombre', '=par_Arrendador', 'Promitente Vendedor'],
       ['=IF(inp_Contacto="","Representante legal",inp_Contacto)', '=inp_Cliente', 'Promitente Comprador'])
doc_sheet(ws, [2, 33, 16, 3, 33, 18, 2], 37)

# ----------------------------------------------------------------- BONOS
ws = ws_bon
for col, w in zip('ABCDEF', [2, 32, 32, 18, 14, 2]):
    ws.column_dimensions[col].width = w
info_b = [('Cliente', '=inp_Cliente'), ('Promotor', '=sel_PromNombre'), ('Puesto / oficina', '=sel_PromPuesto&" · "&sel_PromRegion'),
          ('Originador', '=inp_Originador'), ('Fecha', '=inp_Fecha')]
for i, (lab, f) in enumerate(info_b):
    a = ws.cell(row=2 + i, column=2, value=lab)
    a.font = f_outbox
    a.fill = fill_out
    b = ws.cell(row=2 + i, column=3, value=f)
    b.font = f_base
    if lab == 'Fecha':
        b.number_format = FECHA
        b.alignment = Alignment(horizontal='left')
    ws.merge_cells(start_row=2 + i, start_column=3, end_row=2 + i, end_column=4)
    for col in range(2, 5):
        ws.cell(row=2 + i, column=col).border = box
barra(ws, 8, 2, 4, 'Margen generado por la operación')
mg = [('Valor de la inversión (monto financiado)', '=%s' % CH('$C$19'), MON),
      ('Margen de caja antes de comisiones ($)', '=%s' % CH('$I$4'), MON),
      ('Margen de caja antes de comisiones (%)', '=%s' % CH('$I$5'), PCT),
      ('Comisión para originador', '=%s' % CH('$C$44'), MON),
      ('Margen de caja neto ($)', '=%s' % CH('$I$25'), MON)]
for i, (lab, f, fmt) in enumerate(mg):
    rr = 9 + i
    ws.cell(row=rr, column=2, value=lab).font = f_bold
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
    c = ws.cell(row=rr, column=4, value=f)
    c.number_format = fmt
    c.font = f_bold
    for col in range(2, 5):
        ws.cell(row=rr, column=col).border = box
barra(ws, 15, 2, 5, 'Distribución del margen de caja neto')
for j, h in enumerate(['Centro de costos', 'Nombre', 'Bono ($)', '%']):
    c = ws.cell(row=16, column=2 + j, value=h)
    c.font = f_bold
    c.fill = fill_sub
    c.alignment = center
dist = [('Promotor', '=sel_PromNombre', 0.08), ('Gerencia regional', '', 0.005), ('Reserva bonos', '', 0.004),
        ('Gastos de operación', '', 0.0), ('Referenciador / agencia', '', 0.0)]
for i, (cc, nm_, pct) in enumerate(dist):
    rr = 17 + i
    ws.cell(row=rr, column=2, value=cc).font = f_base
    inp(ws.cell(row=rr, column=3), nm_ if nm_ else None)
    inp(ws.cell(row=rr, column=5), pct, PCT)
    c = ws.cell(row=rr, column=4, value='=MAX(0,$D$13)*E%d' % rr)
    c.number_format = MON
    c.font = f_base
    for col in range(2, 6):
        ws.cell(row=rr, column=col).border = box
rr = 17 + len(dist)
ws.cell(row=rr, column=2, value='Total bonos').font = f_bold
ws.cell(row=rr, column=4, value='=SUM(D17:D%d)' % (rr - 1)).number_format = MON
ws.cell(row=rr, column=5, value='=SUM(E17:E%d)' % (rr - 1)).number_format = PCT
ws.cell(row=rr + 1, column=2, value='Margen neto para FASTPLUS').font = f_bold
ws.cell(row=rr + 1, column=4, value='=D13-D%d' % rr).number_format = MON
ws.cell(row=rr + 1, column=5, value='=1-E%d' % rr).number_format = PCT
for x in (rr, rr + 1):
    for col in range(2, 6):
        ws.cell(row=x, column=col).border = box
        ws.cell(row=x, column=col).fill = fill_sub
name('bon_Distribucion', absref(ws.title, 'A15:F%d' % (rr + 1)))
ws.cell(row=rr + 3, column=2, value='Observaciones bonos:').font = f_bold
for col in range(3, 6):
    ws.cell(row=rr + 3, column=col).border = box
ws.cell(row=rr + 5, column=2, value='Porcentajes sobre el margen de caja neto (sugeridos; editables). Si el margen es negativo no se generan bonos.').font = f_note
doc_sheet(ws, [2, 32, 32, 18, 14, 2], rr + 6)

# ----------------------------------------------------------------- RIESGO
ws = ws_rie
for col, w in zip('ABCDEFG', [2, 28, 26, 14, 30, 18, 2]):
    ws.column_dimensions[col].width = w
ws['B2'] = 'Hoja de Datos Riesgo'
ws['B2'].font = Font(name=FONT, size=14, bold=True)
ws['B2'].alignment = center
ws.merge_cells('B2:F2')
for col in range(2, 7):
    ws.cell(row=2, column=col).border = box
izq = [('Nombre evaluado', '=inp_Cliente', None), ('Acreditado', '=inp_Cliente', None), ('RFC', '=inp_RFC&""', None),
       ('Obligado solidario', '=inp_Obligado', None), ('Unidades', 1, '0" unidad(es)"'),
       ('Plazo', '=inp_PlazoSol&" meses"', None), ('Valor equipo (sin IVA)', '=inp_Valor', MON),
       ('Otros financiados', '=%s+%s+%s' % (CH('$C$17'), CH('$C$18'), CH('$C$42')), MON),
       ('Renta básica', '=%s' % CH('$C$24'), MON), ('Renta total IVA incluido', '=%s+%s' % (CH('$C$24'), CH('$I$21')), MON),
       ('Anticipo', '=%s' % CH('$C$16'), MON), ('Depósito en garantía', '=%s' % CH('$C$30'), MON),
       ('Suma anticipo y depósito', '=%s+%s' % (CH('$C$16'), CH('$C$30')), MON),
       ('Valor residual neto', '=%s' % CH('$C$21'), MON), ('Descripción del bien', '=inp_TipoActivo&", "&inp_Equipo', None)]
for i, (lab, f, fmt) in enumerate(izq):
    rr = 3 + i
    ws.cell(row=rr, column=2, value=lab).font = f_bold
    c = ws.cell(row=rr, column=3, value=f)
    c.font = f_base
    c.alignment = Alignment(horizontal='left' if lab.startswith('Descripción') else 'right', shrink_to_fit=lab.startswith('Descripción'))
    if fmt:
        c.number_format = fmt
    if lab == 'Unidades':
        inp(c, 1, '0" unidad(es)"')
    for col in (2, 3):
        ws.cell(row=rr, column=col).border = box
pct_col = {'Anticipo': '=%s/inp_Valor' % CH('$C$16'), 'Depósito en garantía': '=%s/inp_Valor' % CH('$C$30'),
           'Suma anticipo y depósito': '=(%s+%s)/inp_Valor' % (CH('$C$16'), CH('$C$30')), 'Valor residual neto': '=%s' % CH('$C$20')}
for i, (lab, f, fmt) in enumerate(izq):
    if lab in pct_col:
        c = ws.cell(row=3 + i, column=4, value='=IFERROR(%s,0)' % pct_col[lab][1:])
        c.number_format = '0.00%'
        c.alignment = center
        c.border = box
der = [('TIR real', '=%s' % CH('$I$8'), PCT), ('TIR mínima (rate card)', '=%s' % CH('$C$46'), PCT),
       ('Margen de caja neto', '=%s' % CH('$I$26'), PCT), ('Margen rate card', '=%s' % CH('$C$45'), PCT),
       ('Margen de caja neto ($)', '=%s' % CH('$I$25'), MON), ('Tasa de la operación', '=%s' % CH('$C$8'), PCT),
       ('TIIE + margen', '=TEXT(inp_TIIE,"0.00%%")&" + "&TEXT(%s-inp_TIIE,"0.00%%")' % CH('$C$8'), None),
       ('Ventas anuales (tier)', '=inp_Tier', None), ('Producto', '=inp_Producto', None),
       ('Promotor', '=sel_PromNombre', None), ('Oficina', '=sel_PromRegion', None),
       ('CAT informativo', '=%s' % CH('$I$10'), PCT), ('Deducibilidad de la renta', '=%s' % CH('$C$51'), PCT)]
for i, (lab, f, fmt) in enumerate(der):
    rr = 3 + i
    ws.cell(row=rr, column=5, value=lab).font = f_bold
    c = ws.cell(row=rr, column=6, value=f)
    c.font = f_base
    c.alignment = Alignment(horizontal='right')
    if fmt:
        c.number_format = fmt
    for col in (5, 6):
        ws.cell(row=rr, column=col).border = box
RR = 3 + len(izq) + 1
ws.cell(row=RR, column=2, value='Dictamen').font = f_big
d = ws.cell(row=RR, column=3, value='=%s' % CH('$I$27'))
d.font = Font(name=FONT, size=12, bold=True)
ws.conditional_formatting.add('C%d' % RR, CellIsRule(operator='equal', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='C6EFCE')))
ws.conditional_formatting.add('C%d' % RR, CellIsRule(operator='notEqual', formula=['"CUMPLE"'], fill=PatternFill('solid', fgColor='FFC7CE')))
ws.cell(row=RR, column=5, value='Tasa mínima para cumplir').font = f_bold
x = ws.cell(row=RR, column=6, value='=%s' % CH('$I$15'))
x.number_format = PCT
ws.cell(row=RR + 1, column=2, value='Alertas').font = f_bold
parrafo(ws, RR + 1, 3, 6, '=INDEX(res_Alertas,1,inp_EscTabla)', 30, Font(name=FONT, size=9, color='C00000'), 'left')
barra(ws, RR + 3, 2, 6, 'Datos acta de aprobación')
for j, h in enumerate(['Unidades', 'Plazo', 'Equipo básico', 'Renta total IVA incl.', 'Anticipo + depósito']):
    c = ws.cell(row=RR + 4, column=2 + j, value=h)
    c.font = f_bold
    c.alignment = center
    c.border = box
for j, f in enumerate(['=C7', '=C8', '=C9', '=C12', '=C15']):
    c = ws.cell(row=RR + 5, column=2 + j, value=f)
    c.number_format = MON if j >= 2 else 'General'
    c.alignment = center
    c.border = box
barra(ws, RR + 7, 2, 6, 'Autorizaciones')
for j, (col, lab) in enumerate([(2, 'Promotor'), (3, 'Gerente de Promoción'), (5, 'Crédito'), (6, 'Dirección')]):
    ws.cell(row=RR + 10, column=col).border = Border(bottom=Side(style='thin', color='000000'))
    c = ws.cell(row=RR + 11, column=col, value=lab)
    c.font = f_bold
    c.alignment = center
doc_sheet(ws, [2, 28, 26, 14, 30, 18, 2], RR + 12)

# ==================================================================== HISTORIAL
ws = ws_his
hist_cols = ['Folio', 'Fecha', 'Cliente', 'Equipo', 'Promotor', 'Moneda', 'Valor equipo', 'Plazos',
             'Renta c/IVA Esc.1', 'Renta c/IVA Esc.2', 'Renta c/IVA Esc.3', 'Renta c/IVA Esc.4',
             'CM % Esc.1', 'CM % Esc.2', 'CM % Esc.3', 'CM % Esc.4', 'Estatus', 'Comentarios', 'Guardado por', 'Guardado el']
widths = [12, 11, 30, 26, 22, 14, 15, 14, 15, 15, 15, 15, 10, 10, 10, 10, 15, 30, 16, 17]
ws['A1'] = 'HISTORIAL DE COTIZACIONES'
ws['A1'].font = f_title
ws['A2'] = 'La macro "Guardar en historial" agrega una fila por folio. Para reabrir una cotización seleccione su fila y use "Cargar del historial". Las columnas a la derecha guardan todos los datos capturados.'
ws['A2'].font = f_sub
for j, (h, w) in enumerate(zip(hist_cols, widths), start=1):
    c = ws.cell(row=4, column=j, value=h)
    c.font = f_hdr
    c.fill = fill_hdr
    c.alignment = center
    ws.column_dimensions[L(j)].width = w
ws.row_dimensions[4].height = 30
ws.freeze_panes = 'B5'
ws.auto_filter.ref = 'A4:T4'
dv_list(ws, '=lst_Estatus', 'Q5:Q2000')
name('hist_Header', absref(ws.title, 'A4'))


# ==================================================================== IMPRESIÓN / PROTECCIÓN
ws_his.page_setup.orientation = 'landscape'
ws_his.page_setup.fitToWidth = 1
ws_his.page_setup.fitToHeight = 0
ws_his.sheet_properties.pageSetUpPr.fitToPage = True
from openpyxl.workbook.protection import WorkbookProtection
for w in [ws_fac, ws_bon, ws_rie, ws_cfg]:
    w.sheet_state = 'hidden'
wb.security = WorkbookProtection(workbookPassword=CLAVE_ADMIN, lockStructure=True)
for w in [ws_cot, ws_fac, ws_pro, ws_pin, ws_ven, ws_bon, ws_rie, ws_tab] + ws_cr:
    w.protection.sheet = True
    w.protection.password = CLAVE_ADMIN
    w.protection.formatColumns = False
    w.protection.formatRows = False
    w.protection.selectLockedCells = False
    w.protection.selectUnlockedCells = False
wb.active = 0
wb.save(OUT)
import json
json.dump({'P': P, 'names': names, 'FIRST': FIRST, 'LAST': LAST}, open('layout_fast.json', 'w'), indent=1, ensure_ascii=False)
print('ok', OUT)
