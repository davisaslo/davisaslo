import openpyxl, datetime as dt, subprocess, sys
from dateutil.relativedelta import relativedelta
S='/root/.claude/skills/synced/b646bf01-e131-40dd-b005-f47363650fca_4f0b063a-135a-45c1-a143-d0849feced22/xlsx/scripts/recalc.py'
def pmt(i,n,pv,fv,t):
    if i==0: return -(pv+fv)/n
    return -(pv*(1+i)**n+fv)*i/((1+i*t)*((1+i)**n-1))
def xirr(cfs):
    d0=cfs[0][0]; f=lambda r: sum(v/(1+r)**((d-d0).days/365) for d,v in cfs); lo,hi=-0.99,10
    for _ in range(300):
        mid=(lo+hi)/2
        if f(lo)*f(mid)<=0: hi=mid
        else: lo=mid
    return mid
def setv(wb,nm,v,k=None):
    sh,ref=wb.defined_names[nm].attr_text.split('!'); sh=sh.strip("'"); ref=ref.replace('$','').split(':')[0]
    ws=wb[sh]; c=ws[ref]
    if k: c=ws.cell(row=c.row+k-1, column=c.column) if ':' in wb.defined_names[nm].attr_text and wb.defined_names[nm].attr_text.split(':')[1].replace('$','')[0]==ref[0] else ws.cell(row=c.row, column=c.column+k-1)
    c.value=v
def run(case):
    wb=openpyxl.load_workbook('Cotizador_FASTPLUS.xlsx')
    for nm,v in case['inp'].items(): setv(wb,nm,v)
    for (nm,k),v in case.get('esc',{}).items(): setv(wb,nm,v,k)
    wb.save('tf.xlsx')
    out=subprocess.run(['python3',S,'tf.xlsx','150'],capture_output=True,text=True).stdout
    assert '"total_errors": 0' in out, out
    wv=openpyxl.load_workbook('tf.xlsx',data_only=True)
    p=case['p']; ok=True
    for k in range(1,5):
        ws=wv['Corrida %d'%k]; C=lambda r: ws['C%d'%r].value; I=lambda r: ws['I%d'%r].value
        n=p['plazos'][k-1]; ia=p['tiie']+p['marg'][k-1]; i=ia/12; tau=p['tau']
        V=p['precio']/1.16; eng=p['antic']; segf=p['seg']/1.16 if p['segfin'] else 0; gpsf=p['gps']/1.16; otrf=p['otros']/1.16
        M=V-eng+segf+gpsf+otrf; VR=p['res'][k-1]*V
        R=pmt(i,n,-M,VR,tau)
        ff=p['fecha']; dias=p['dias']; inicio=ff+dt.timedelta(days=dias)
        days=dias; fa=p['fondeo']; f=fa/12; DF=(1+fa/360)**(-days)
        com=p['com'][k-1]*M; dep=p['depp'][k-1]*M*1.16; rp=R/30*days; gi=p['gi']
        segc=0 if p['segfin'] else p['seg']/1.16
        rows=[(t+1-tau,R) for t in range(n)]+[(n,VR)] if VR>0 else [(t+1-tau,R) for t in range(n)]
        net0=-(V-p['desc']/1.16+segf+gpsf+otrf)+eng+com+rp+dep+gi
        N=len(rows)
        cm=net0+sum((v-(dep if ix==N-1 else 0))*DF*(1+f)**(-t) for ix,(t,v) in enumerate(rows))
        cmn=cm-p['cp']*M
        # fechas: inicio del plazo = fecha + dias (vencido: primera = EDATE(inicio,1))
        tir=xirr([(ff,net0)]+[(inicio+relativedelta(months=t),v-(dep if ix==N-1 else 0)) for ix,(t,v) in enumerate(rows)])
        def cmn_at(ia_):
            ii=ia_/12; RR=pmt(ii,n,-M,VR,tau); rpp=RR/30*days
            n0=-(V-p['desc']/1.16+segf+gpsf+otrf)+eng+com+rpp+dep+gi
            fl=[(t+1-tau,RR) for t in range(n)]+([(n,VR)] if VR>0 else [])
            return (n0+sum((v-(dep if ix==len(fl)-1 else 0))*DF*(1+f)**(-t) for ix,(t,v) in enumerate(fl))-p['cp']*M)/M
        lo,hi=0,2
        for _ in range(200):
            mid=(lo+hi)/2
            if cmn_at(mid)<p['obj'][k-1]: lo=mid
            else: hi=mid
        mid=max(mid,0.26)
        pi=(eng+com+rp+gi+segc)*1.16+dep
        chk=[('R',C(24),R),('M',C(19),M),('dep',C(30),dep),('com',C(28),com),('CMn',I(25),cmn),('TIR',I(8),tir),
             ('TasaMin',I(15),mid),('PI',I(20),pi),('Req+Rgps+Rseg+Rotr',C(48)+C(53)+C(49)+C(50),R),('obj',C(45),p['obj'][k-1]),('Saldo',I(16),0)]
        bad=[(a,x,y) for a,x,y in chk if abs(x-y)>max(1e-6,abs(y)*1e-7)]
        ok&=not bad
        print('  Plazo %d: renta %.2f  margen neto %.2f  TIR %.3f%%  tasa mín %.3f%%  pago ini %.2f -> %s'%(n,R,cmn,tir*100,mid*100,pi,'OK' if not bad else bad))
    return ok,wv
base=dict(precio=879802,antic=75845,tiie=0.086,marg=[0.214,0.204,0.194,0.184],res=[0.3,0.2,0.15,0.1],plazos=[12,24,36,48],
          tau=0,fecha=dt.datetime(2026,10,6),dias=0,fondeo=0.21,com=[0.02]*4,depp=[0]*4,gi=0,seg=0,segfin=False,gps=0,otros=0,
          desc=0,cp=0,obj=[0.04,0.05,0.06,0.07])
print('Caso 1: datos por defecto')
ok1,wv=run(dict(inp={},p=base))
alt=dict(base, antic=87980, tiie=0.087, dias=12, desc=40000, cp=0.01, gi=1500, seg=45000, segfin=True, gps=3132, otros=11600,
         depp=[0.05]*4, tau=1, obj=[0.10]*4)
print('Caso 2: anticipado, 12 días, descuento, seguro financiado, GPS, otros, depósito 5%, comisión originador 1%, objetivo manual 10%')
ok2,wv2=run(dict(inp={'inp_Anticipo':87980,'inp_TIIE':0.087,'inp_DiasRP':12,'inp_Descuento':40000,'inp_ComProm':0.01,
   'inp_GastosInv':1500,'inp_SeguroMonto':45000,'inp_SeguroFin':'Financiado','inp_GPSMonto':3132,'inp_OtrosMonto':11600,
   'inp_Modalidad':'Anticipado','inp_CMobj':0.10},
   esc={('esc_DepPct',k):0.05 for k in range(1,5)}, p=alt))
print('TODO OK' if ok1 and ok2 else 'HAY DIFERENCIAS')
print('Caso 3: tasa de fondeo manual 12.6%')
ok3,_=run(dict(inp={'inp_FondeoManual':0.126}, p=dict(base, fondeo=0.126)))
print('Caso 4: tasa de fondeo manual 15%')
ok4,_=run(dict(inp={'inp_FondeoManual':0.15}, p=dict(base, fondeo=0.15)))
print('FONDEO OK' if ok3 and ok4 else 'FONDEO CON DIFERENCIAS')
